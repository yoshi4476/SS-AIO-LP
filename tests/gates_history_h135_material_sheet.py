# -*- coding: utf-8 -*-
"""2026-10-21 運用者「採点で材料が足りないと出た記事の質問シートを、社ごとに機械で作る」「コンフレックス以外も」。
材料の依頼（material_requests.json）と、採点で何度も足りないと言われる材料の決まった質問から、社ごとの質問シートを作る。
返ってきたシートは intake/ に置けば intake_watch.py が拾い、答えに無い数字・既存と食い違う答え・掲載不可・お客様の声を
登録せずに保留にする（material_sheet.py）。AI は差し替えて呼ばない。
"""
import json
import sys
import tempfile
from pathlib import Path

from test_gates import check, ROOT

sys.path.insert(0, str(ROOT / "scripts"))


def _fill(path, answers):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["質問"]
    for row in ws.iter_rows(min_row=5):
        no = str(row[0].value or "")
        if no in answers:
            ans, ok = answers[no]
            row[4].value = ans
            row[5].value = "2024年1月〜2026年9月"
            row[7].value = ok
    wb.save(path)


def test_make_sheet_per_site_with_requests_and_template():
    import material_sheet as MS
    print("\n■ 質問シート: 依頼の質問を先に、決まった質問（事例カード・失敗の件数・判断の根拠…）を後に。隠しタブで見分ける")
    reqs = {"zz-slug": {"keyword": "庭木 剪定 業者", "since": "2026-10-20",
                        "questions": ["実際にあったことを2〜3件", "全体で何件か", "1件の事例"], "scorer": ["母数が無い"]}}
    old = MS.requests_of, MS.display_name, MS.fact_gaps
    MS.fact_gaps = lambda *a, **k: []          # 補いの行は test_fact_fill_rows_update_in_place で確かめる
    MS.requests_of, MS.display_name = (lambda s: reqs), (lambda s: "テストの社")
    try:
        with tempfile.TemporaryDirectory() as d:
            p, n = MS.make("ai-lab", d, today="2026-10-21")
            from openpyxl import load_workbook
            wb = load_workbook(p)
            meta = {str(r[0]): str(r[1]) for r in wb["材料の質問"].iter_rows(values_only=True)}
            items = [str(r[1]) for r in wb["質問"].iter_rows(min_row=5, values_only=True) if r[0]]
            check("依頼の3問＋決まった質問（自社は費用の実例を聞かない）",
                  (n, items[0].startswith("記事「庭木"), "事例カード" in items, "費用の実例" in items), (10, True, True, False))
            check("隠しタブに社と各問の印", (meta.get("site"), meta.get("1"), wb["材料の質問"].sheet_state), ("ai-lab", "req:zz-slug:0", "hidden"))
            import intake_watch as IW
            check("intake_watch が材料の質問シートと見分ける", IW.kind(p), "material")
    finally:
        MS.requests_of, MS.display_name, MS.fact_gaps = old


def test_apply_holds_invented_numbers_conflicts_and_voices():
    import material_sheet as MS
    print("\n■ 登録: 答えに無い数字・既存と食い違う答え・掲載不可・お客様の声は登録しない")
    old = MS.requests_of, MS.display_name, MS.existing, MS.fact_gaps
    MS.fact_gaps = lambda *a, **k: []
    MS.requests_of, MS.display_name = (lambda s: {}), (lambda s: "テストの社")
    MS.existing = lambda s: ([("zz-hiyou", "代表的な12件の費用は80万〜800万円でした")], [])
    try:
        with tempfile.TemporaryDirectory() as d:
            p, _ = MS.make("ai-lab", d, today="2026-10-21")
            # 1=事例カード（正しい）、2=事例カード（AI が数字を作る）、3=失敗（既存と食い違い）、4=不可、7=お客様の声（自社は費用の実例を聞かないので7問目）
            _fill(p, {"1": ("2025年3月に卸売の会社で入力が月40時間から8時間になった", "可"),
                      "2": ("2024年に飲食店で予約が増えた", "可"),
                      "3": ("代表的な12件の費用は80万〜650万円", "可"),
                      "4": ("社外秘の件", "不可"),
                      "7": ("「とても良かった」と言われた", "可")})
            reply = json.dumps({"facts": [
                {"no": "1", "claim": "2025年3月に卸売の会社で、入力が月40時間から8時間になりました", "topic": ["卸売"]},
                {"no": "2", "claim": "2024年に飲食店で予約が月20件から35件に増えました", "topic": ["飲食"]}],
                "experiences": [], "conflicts": [{"no": "3", "existing": "zz-hiyou", "why": "上限が800万円と650万円で違う"}]},
                ensure_ascii=False)
            site, answers = MS.read(p)
            f, e, held = MS.plan(site, answers, ask=lambda prompt: reply)
            reasons = " ".join(w for _, w in held)
            check("答えのとおりの文だけ登録する", [x["no"] for x in f], ["1"])
            check("答えに無い数字（20・35）を入れた文は保留", "答えに無い数字" in reasons and "20" in reasons, True)
            check("既存と食い違う答えは保留（記録の id を添える）", "zz-hiyou" in reasons, True)
            check("掲載不可とお客様の声は登録しない", sorted(n for n, w in held if "記事に出さない" in w), ["4", "7"])
            ok, msg = MS.apply(p, False, ask=lambda prompt: reply)
            check("--write なしは書かない（見るだけ）", (ok, "--write" in msg), (True, True))
    finally:
        MS.requests_of, MS.display_name, MS.existing, MS.fact_gaps = old


def test_auto_measure_mail_and_voice_draft():
    import material_sheet as MS
    print("\n■ 施策の前後を自動で測る・質問シートを週次に送る・お客様の声を下書きへ（運用者 2026-10-21）")
    # 前後の測定: 31日たった自社の書き直しだけ・前が0回（まだ検索に出ていなかった）と小さすぎる数は一次情報にしない
    src = (ROOT / "scripts" / "material_sheet.py").read_text(encoding="utf-8")
    check("前が0回・30回未満は一次情報にしない・お客様の社は測らない・31日",
          ("if imp[0] == 0:" in src, "MIN_IMP = 30" in src, "if P.is_client(site) or meta.get(\"noindex\"):" in src,
           "timedelta(days=31)" in src), (True, True, True, True))
    # メール: 依頼が変わった社だけ・送った後は印を残して同じ依頼では送らない
    old = MS.targets, MS.requests_of, MS.MAIL_STATE, MS.make
    with tempfile.TemporaryDirectory() as d:
        sent = []
        MS.targets = lambda only="": ["ai-lab"]
        MS.requests_of = lambda s: {"zz": {"keyword": "庭木", "questions": ["a"]}}
        MS.MAIL_STATE = Path(d) / "mail.json"
        MS.make = lambda s, out, today=None: (Path(d) / "x.xlsx", 3)
        (Path(d) / "x.xlsx").write_bytes(b"x")
        env = dict(__import__("os").environ)
        try:
            __import__("os").environ.update({"RESEND_API_KEY": "k", "NOTIFY_TO_EMAIL": "a@example.com"})
            first = MS.mail(Path(d), send=sent.append)
            second = MS.mail(Path(d), send=sent.append)
        finally:
            __import__("os").environ.clear()
            __import__("os").environ.update(env)
            MS.targets, MS.requests_of, MS.MAIL_STATE, MS.make = old
        check("依頼が変わった週だけ送り、同じ依頼では送らない", (first, second, len(sent), len(sent[0]["attachments"])), (1, 0, 1, 1))
    check("手元からは送らない", "MATERIAL_MAIL=skip" in src, True)
    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("週次で送る・測る", ("material_sheet.py --mail" in wf, "material_sheet.py --auto-measure 3" in wf), (True, True))
    # お客様の声: 掲載可否は「要確認」で下書きへ（機械は掲載しない）
    with tempfile.TemporaryDirectory() as d:
        ans = [{"no": "7", "key": "tpl:voice", "item": "お客様の声", "question": "", "answer": "助かりました",
                "period": "2025年", "count": "", "ok": "可"}]
        p = MS.voice_draft("ai-lab", ans, out_dir=d, today="2026-10-21")
        from openpyxl import load_workbook
        ws = load_workbook(p)["お客様の声"]
        rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if any(r)]
        check("声は「要確認」で下書きに1行だけ（雛形の「可」の空行は消す）", [(r[0], r[5]) for r in rows], [("要確認", "助かりました")])


def test_research_reaches_other_sites_with_right_owner():
    import industry_ai_sources as I
    import private_store as P
    import sites as S
    print("\n■ 業種調査は自社の社にだけ渡す（社の設定 research）。お客様の記事に当社の名前と調べたことを書かない（お客様の記事が配信の門で止まった 2026-10-22）")
    old_load, old_client = S.load, P.is_client
    S.load = lambda s: {"research": ["inshoku"]} if s in ("client-zz", "corporate") else {}
    P.is_client = lambda s: s == "client-zz"
    try:
        cl, _ = I.for_site("client-zz", "", "")
        own, _ = I.for_site("corporate", "", "")
        none, _ = I.for_site("subsidy", "", "")
    finally:
        S.load, P.is_client = old_load, old_client
    check("お客様の社には渡さない・自社は「当社は」・設定の無い社は渡さない",
          (cl, all(c.startswith("当社は") for c in own) and bool(own), none), ([], True, []))


def test_measure_claim_needs_past_tense_and_ci_skip():
    print("\n■ 施策の前と後: やったことは「〜した」で書かせる・CI では質問シートを作らない")
    src = (ROOT / "scripts" / "material_sheet.py").read_text(encoding="utf-8")
    check("「〜した」の形でなければ止める・CI では作らない",
          ('if not a.what.endswith("た"):' in src, 'MATERIAL_SHEET=skip' in src), (True, True))
    iw = (ROOT / "scripts" / "intake_watch.py").read_text(encoding="utf-8")
    check("intake_watch --apply の後に、依頼の残る社の質問シートを作り直す", iw.count("MS.make(s, MS.Path.home() / \"Desktop\")"), 2)


def test_fact_fill_rows_update_in_place():
    import json
    import tempfile
    from pathlib import Path
    import facts as F
    import material_sheet as M
    print("\n■ よく使う一次情報の補い: 記事でよく使われ、母数・ばらつき・期間・数え方の欠けた文をシートに出し、答えで書き換える（2026-10-22）")
    fs = [{"id": "zz-a", "claim": "2023年4月から2026年9月までに担当した15社で、入力作業が月平均20時間から6時間になりました",
           "period": "2023-04〜2026-09", "source": "運用者の申告"},
          {"id": "zz-b", "claim": "請求書の台帳を担当者が月末に集計したところ、2025年4月から2026年3月までの42社のうち最も早い会社は5日でした"},
          {"id": "own-x", "claim": "検索の表示は12件でした"}]
    bodies = ["月平均20時間から6時間に減り、15社で…", "15社では20時間が6時間に", "42社のうち5日", "関係ない本文"]
    old = F.load_for
    F.load_for = lambda site: ("", fs)
    try:
        gaps = M.fact_gaps("zz", bodies=bodies)
    finally:
        F.load_for = old
    check("欄（期間・出どころ）ではなく文で欠けを見る・機械の一次情報（own-）は聞かない・使う記事が2本未満は聞かない",
          [(g[0], g[2], [m[:2] for m in g[3]]) for g in gaps], [("zz-a", 2, ["内訳", "数え"])])
    ans = {"no": "1", "key": "fact:zz-a", "question": "今の文: 「…15社で、入力作業が月平均20時間から6時間…」",
           "answer": "各社の担当者が作業の記録を月末に集計。最小は月3時間・最大は月10時間", "period": "", "count": ""}
    check("補いの答えは、今の文の数字も使ってよい（作った数字だけを止める）",
          (M.invented("15社で月平均20時間から6時間（最小3時間・最大10時間）", ans), M.invented("15社で7時間", ans)), ([], ["7"]))
    with tempfile.TemporaryDirectory() as d:
        own_f, own_e = M.OWN_FACTS, M.OWN_EXPS
        M.OWN_FACTS, M.OWN_EXPS = Path(d) / "f.json", Path(d) / "e.json"
        M.OWN_FACTS.write_text(json.dumps({"facts": [dict(fs[0])]}, ensure_ascii=False), encoding="utf-8")
        M.OWN_EXPS.write_text(json.dumps({"experiences": []}), encoding="utf-8")
        try:
            M.write("ai-lab", [{"claim": "新しい文（15社・最小3時間・最大10時間）", "replaces": "zz-a"}], [], today="2026-10-22")
            got = json.loads(M.OWN_FACTS.read_text(encoding="utf-8"))["facts"]
        finally:
            M.OWN_FACTS, M.OWN_EXPS = own_f, own_e
    check("補いは同じ id の文を書き換える（同じ事実を2つ並べない）",
          ([g["id"] for g in got], got[0]["claim"].startswith("新しい文"), "補い" in got[0]["source"]), (["zz-a"], True, True))
