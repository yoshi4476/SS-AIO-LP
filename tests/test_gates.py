# -*- coding: utf-8 -*-
"""品質ゲートの回帰テスト

実行: python tests/test_gates.py

ここにあるのは、実際に起きた不具合をそのまま固定したもの。
どれもコードを読むだけでは見つからず、動かして初めて分かった。

  1. 食い合い検査が、審査対象の語そのものの台帳行を重複と数え、
     どのKWも弾かれてパイプラインがKWを1件も取得できなくなった
  2. 狙う語を文字の類似で比べ、「it導入補助金 学習塾」と
     「it導入補助金 飲食店」を同じものとみなした（88件中ほとんどが誤検出）
  3. タグの開始と終了の数だけを見ていたため、閉じ忘れと余分な閉じが
     相殺し、</content> が公開HTMLに残った
  4. 採点とビルドで文字数の数え方が違い、同じ記事が5,251字と4,922字になった
"""
import os
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

FAIL = []


def check(name, got, want):
    ok = got == want
    print(f"  {'OK' if ok else 'NG'}  {name}" + ("" if ok else f"  （得た {got!r} / 期待 {want!r}）"))
    if not ok:
        FAIL.append(name)


# ── 1. 狙う語の判定 ──────────────────────────────
def test_kw_conflicts():
    from cannibal_check import kw_conflicts, norm_kw
    print("\n■ 狙う語の食い合い判定")
    arts = [{"slug": "a", "kw": "aio診断", "title": "AIO診断のやり方", "desc": "", "h2": [], "cat": "aio"},
            {"slug": "b", "kw": "it導入補助金 学習塾", "title": "学習塾のIT導入補助金", "desc": "", "h2": [], "cat": "aio"},
            {"slug": "c", "kw": "整骨院 集客", "title": "整骨院の集客", "desc": "", "h2": [], "cat": "aio"}]
    check("表記ゆれを同じ語とみなす", bool(kw_conflicts("aio 診断", arts)), True)
    check("全角も同じ語とみなす", norm_kw("ＡＩＯ診断") == norm_kw("aio診断"), True)
    # 業種が違うだけのKWを弾かない（実データで88件中ほとんどが誤検出だった）
    check("業種違いは食い合いにしない", bool(kw_conflicts("it導入補助金 飲食店", arts)), False)
    check("狭い語は広い語に含まれる", bool(kw_conflicts("整骨院 集客 口コミ", arts)), True)
    check("無関係な語は通す", bool(kw_conflicts("動物病院 集客", arts)), False)


# ── 2. タグの対応 ────────────────────────────────
def test_tag_balance():
    sys.path.insert(0, str(ROOT / "scripts"))
    from live_check import tag_balance
    print("\n■ タグの対応検査")
    check("正しい入れ子は通る", tag_balance("<div><p>あ</p></div>"), [])
    check("void要素は数えない", tag_balance("<div><img src='x'><br></div>"), [])
    # 数を数えるだけだと、この2つは相殺して見つからない
    both = tag_balance("<div>あ</content>")
    check("閉じ忘れと余分な閉じを両方見つける", len(both), 2)
    check("開始タグの無い閉じを見つける",
          any("対応する開始タグが無い" in x for x in tag_balance("<p>あ</p></content>")), True)
    check("コードブロック内は無視する", tag_balance("<script>if(a</b>){}</script>"), [])


# ── 3. 文字数の数え方 ────────────────────────────
def test_char_count():
    import md2html
    print("\n■ 文字数の数え方")
    body = "## 見出し\n\n本文です。\n\n| 列A | 列B |\n|:--|:--|\n| あ | い |\n"
    plain = re.sub(r"\s|<[^>]+>", "", md2html.convert(body)[0])
    # 見出しの # や表の | を数えない（採点とビルドで食い違う原因だった）
    check("記法の記号を数えない", ("#" in plain) or ("|" in plain), False)
    check("本文は数える", "本文です。" in plain, True)


# ── 4. 管制塔の食い合い判定（GAS） ──────────────────
def test_hub_gas():
    print("\n■ 管制塔（GAS）の食い合い判定")
    src = (ROOT / "automation" / "gas" / "hub.gs").read_text(encoding="utf-8")

    def fn(name):
        m = re.search(rf"function {name}\(.*?\n\}}\n", src, re.S)
        return m.group(0) if m else ""

    harness = """
var READS = 0, ROWS = [];
for (var i = 0; i < 60; i++) ROWS.push(['ai-lab','重複語' + (i%3), '公開済み','A','','','','','','','']);
for (var i = 0; i < 40; i++) ROWS.push(['ai-lab', i < 20 ? '重複語' + (i%3) : '新しい語' + i,
                                        '未着手','B','','','','','','','']);
function kwRows_() { READS++; return ROWS; }
"""
    js = harness + fn("normKw_") + fn("kwConflict_") + fn("nextKw_") + """
var r = nextKw_('ai-lab');
console.log(JSON.stringify({reads: READS, kw: r.keyword,
                            skipped: (r.skipped_conflict||[]).length}));
"""
    tmp = ROOT / "tests" / "_hub_tmp.js"
    tmp.write_text(js, encoding="utf-8")
    try:
        out = subprocess.run(["node", str(tmp)], capture_output=True, text=True,
                             encoding="utf-8", errors="replace", timeout=60)
        import json
        d = json.loads(out.stdout.strip() or "{}")
    finally:
        tmp.unlink(missing_ok=True)
    # 自分の台帳行を重複と数えると、どのKWも弾かれてパイプラインが止まる
    check("自分の行を重複と数えない", bool(d.get("kw", "").startswith("新しい語")), True)
    # 候補ごとに台帳を読み直すと、未着手が数百件でGASの実行時間の上限に当たる
    check("台帳の読み取りは1回だけ", d.get("reads"), 1)
    check("本当の重複だけを飛ばす", d.get("skipped"), 20)


# ── 5. 執筆後ゲートの自己照合 ────────────────────
def test_self_exclusion():
    """書いた記事が自分自身と「完全一致」で止まった事故の再現。
    正当な新記事3本が全部この誤検出で止まり、その回の執筆が無駄になった。"""
    from kw_guard import judge
    print("\n■ 執筆後ゲートの自己照合")
    import glob, re
    # 実在の記事を1本選び、自分のKWで審査する。
    # 導入直後は記事が無いので、その場合は飛ばす（雛形でも落ちないように）
    files = sorted(glob.glob(str(ROOT / "articles" / "*.md")))
    if not files:
        print("  --  記事が無いため飛ばします（導入直後の状態）")
        return
    f = files[0]
    slug = Path(f).stem
    fm = Path(f).read_text(encoding="utf-8-sig").split("---", 2)[1]
    kw = (re.search(r"^keyword:\s*(.+)$", fm, re.M) or [0, ""])[1].strip()
    lv_with, _ = judge(kw, "", use_gsc=False, exclude_slug=slug)
    lv_without, _ = judge(kw, "", use_gsc=False)
    check("自分を除けば自分とは食い合わない", lv_with < 2, True)
    check("除かなければ完全一致で止まる（検査自体は生きている）", lv_without, 2)


# ── 6. 公開済み記事を新規とみなさない ────────────
def test_published_not_rewritten_as_new():
    """公開中の記事10本が隔離された事故の再現。

    内部リンクの追加などで既存記事が書き換わると、gitの差分では
    「変更」として出る。それを新規として審査すると、別ページが同じ語で
    順位を持っているという理由で、公開中の記事が隔離されてしまう。
    """
    import kw_gate
    print("\n■ 公開済み記事の扱い")
    pairs = kw_gate.written_keyword("ai-lab")
    published = [s for s, _ in pairs
                 if list((ROOT / "site").glob(f"*/{s}/index.html"))]
    check("公開済みの記事を新規として拾わない", published, [])


# ── 7. トークンの権限判定 ──────────────────────
def test_token_check_probes_write():
    """「書き込み可」と出るのに配信が403で落ちた事故の再現。

    GitHub APIの permissions.push は「その利用者の権限」であって
    トークンの権限ではない。細粒度PATが Contents: Read only でも
    true が返るため、これを見ていた検査は通り、配信で落ちた。
    """
    src = (ROOT / "scripts" / "token_check.py").read_text(encoding="utf-8")
    print("\n■ トークンの権限判定")
    check("利用者の権限で判定していない",
          'get("permissions") or {}).get("push")' in src, False)
    check("実際に書き込みを試している", "_probe_write" in src, True)
    check("403を権限不足として扱う", "403" in src, True)


# ── 8. 自動修復の待ち受け名 ──────────────────────
def test_selfheal_watches_real_workflows():
    """待ち受ける名前が1文字でも違うと、落ちても何も起きない。

    最初に書いたときは日本語の名前を並べていたが、実際の name: は
    英語だった。名前は目で合わせず、機械で突き合わせる。
    """
    import yaml
    print("\n■ 自動修復の待ち受け名")
    wf = ROOT / ".github" / "workflows"
    heal = yaml.safe_load((wf / "selfheal.yml").read_text(encoding="utf-8"))
    # PyYAML は on: を True と読む
    trig = heal.get("on") or heal.get(True) or {}
    watched = set((trig.get("workflow_run") or {}).get("workflows") or [])
    actual = {yaml.safe_load(f.read_text(encoding="utf-8")).get("name")
              for f in wf.glob("*.yml") if f.name != "selfheal.yml"}
    missing = sorted(watched - actual)
    check("待ち受け名がすべて実在する", missing, [])
    check("待ち受けが空でない", bool(watched), True)


# ── 9. ラボからコーポレートへの導線 ────────────────
def test_lab_links_to_corporate():
    """カテゴリーの記事を読み切った人の行き先。

    ラボは読む場所、コーポレートは頼む場所。つながっていないと、
    記事を読み終えた人が次にどこへ行けばいいか分からない。
    リンク先を手で書いているので、綴りの間違いを機械で見る。
    """
    import glob
    import json as _json
    print("\n■ ラボからコーポレートへの導線")
    own = [f for f in (ROOT / "sites").glob("*.json")
           if _json.loads(f.read_text(encoding="utf-8")).get("type") == "self-static"]
    if not own:
        print("  --  自前ビルドのサイトが無いため飛ばします")
        return
    conf = _json.loads(own[0].read_text(encoding="utf-8"))
    cats = list((conf.get("categories") or {}).keys())
    # コーポレート側に実在するページ
    src = (ROOT / ".publish-work" / "corporate" / "src" / "lib" / "services.ts")
    known = {"/aio-agent", "/rakushift", "/contact", "/company", "/"}
    if src.exists():
        known |= {"/services/" + s for s in
                  re.findall(r'slug: "([^"]+)"', src.read_text(encoding="utf-8"))}
    missing, nolink = [], []
    for c in cats:
        f = ROOT / "site" / c / "index.html"
        if not f.is_file():
            continue
        html = f.read_text(encoding="utf-8")
        urls = set(re.findall(r'href="https://corp\.7senses\.co\.jp([^"]*)"', html))
        deep = {u for u in urls if u not in ("", "/")}
        if not deep:
            nolink.append(c)
        if src.exists():
            missing += [f"{c}{u}" for u in deep if u.split("#")[0] not in known]
    if not any((ROOT / "site" / c / "index.html").is_file() for c in cats):
        print("  --  カテゴリーページがまだ無いため飛ばします（導入直後の状態）")
        return
    check("全カテゴリーに個別リンクがある", nolink, [])
    check("リンク先が配信先に実在する", missing, [])


# ── 10. 日次監査の本数カウント ────────────────────
def test_daily_audit_ignores_unscored_drafts():
    """score未設定（Phase 5未通過）の記事を「本日公開」に数えた事故の再現。

    publish.pyはscore<90を配信拒否するため、そういう記事はまだサイトに無い。
    それを本数に数えると「本日OK」と誤表示した上で、check_liveが
    「公開したのに404」という偽の不具合をTODOに積む。"""
    import daily_audit
    print("\n■ 日次監査の本数カウント")
    today = daily_audit.today_iso()
    fixture = {"ai-lab": [
        {"slug": "a", "date": today, "score": "95", "title": "t", "category": "c"},
        {"slug": "b", "date": today, "score": "", "title": "t2", "category": "c"},
    ]}
    orig = daily_audit.articles_by_site
    daily_audit.articles_by_site = lambda: fixture
    try:
        by_site = daily_audit.check_volume([])
    finally:
        daily_audit.articles_by_site = orig
    check("scoreのない記事は本数に数えない", len(by_site["ai-lab"]), 1)


def test_every_article_has_a_lead_path():
    """全記事にリード導線があること。

    読んで納得した人の行き先が無いと、記事はそこで終わる。
    実際コーポレートは88本中75本に行き先が無く、記事からの反応がゼロだった。
    新しい記事が増えるたびに漏れるため、機械で見張る。
    """
    lead = re.compile(r"/diagnosis/|/site-audit/|#diagnosis|/contact|/lp/")
    missing = []
    for f in sorted((ROOT / "articles").glob("*.md")):
        body = f.read_text(encoding="utf-8", errors="replace").split("---", 2)[-1]
        if not lead.search(body):
            missing.append(f.stem)
    check("リード導線の無い記事", missing[:5], [])

    # 検査そのものが働いているか。build.py が見張っていないと、
    # 上の判定が通っていても次に増えた記事で崩れる
    src = (ROOT / "scripts" / "build.py").read_text(encoding="utf-8")
    check("build.pyがリード導線を見張っている", "リード導線なし" in src, True)

    # 日次パイプラインで新しい記事にも入ること
    wf = (ROOT / ".github" / "workflows" / "pipeline-multi.yml").read_text(encoding="utf-8")
    check("日次で導線を入れている", "tool_links.py --write" in wf, True)


def test_token_never_in_command_line():
    """配信のコマンドラインにトークンを載せないこと。

    URLに埋めると、プロセス一覧を見るだけでPATが読める。特権は要らない。
    実際、稼働中の配信からPATの全体が読み出せた。失敗時のログにも残る。
    """
    src = (ROOT / "scripts" / "publish.py").read_text(encoding="utf-8")
    check("URLにトークンを埋めていない", "x-access-token:{token}" in src, False)
    check("askpass経由で渡している", "GIT_ASKPASS" in src, True)
    check("失敗ログでトークンを伏せている", "def mask(" in src, True)


def test_kw_intent_separates_click_need():
    """検索結果で用が済む語と、開かないと済まない語を見分けられること。

    同じ順位でもCTRは5倍違った。例文・診断系は8〜9位で表示88回・クリック0、
    不採択理由・書き方は5〜7位でクリックが付いていた。
    この判定が壊れると、読まれない語を書き続けることになる。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import kw_intent
    weak = ["口コミ 返信 例文 医療", "aio 診断", "meo対策 とは"]
    strong = ["ai導入補助金 不採択理由", "実績報告 書き方", "持続化補助金 対象 条件"]
    check("弱い語を弱と判定", [kw_intent.verdict(k)[0] for k in weak], ["弱"] * 3)
    check("強い語を強と判定", [kw_intent.verdict(k)[0] for k in strong], ["強"] * 3)

    # 執筆前のゲートが注意を出すこと。判定だけあっても使われなければ意味がない
    guard = (ROOT / "scripts" / "kw_guard.py").read_text(encoding="utf-8")
    check("食い合いゲートが開く理由を見る", "kw_intent" in guard, True)
    disc = (ROOT / "scripts" / "kw_discover.py").read_text(encoding="utf-8")
    check("候補の並びに反映している", "kw_intent" in disc, True)


def test_lead_funnel_is_watched():
    """リード導線の各段階が、週次で測られ続けること。

    導線を入れただけでは伸びない。押されているのか、押した先で
    落ちているのかが分からないと、次に直す場所が決まらない。
    実際、記事のCTAは長らくクリックが記録されず、判断できなかった。
    """
    f = ROOT / "scripts" / "funnel.py"
    check("段階を測る道具がある", f.is_file(), True)
    src = f.read_text(encoding="utf-8") if f.is_file() else ""
    for ev in ("cta_click", "form_start", "form_submit"):
        check(f"{ev} を見ている", ev in src, True)
    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("週次で走る", "funnel.py" in wf, True)


def test_cta_wording_is_measured_not_guessed():
    """CTAの文言を推測で決めないこと。

    記事を見た544件に対しクリックは38件。どの文言なら押されるかは
    推測では決まらない。半々で出し分けて実測する仕組みを保つ。
    """
    b = (ROOT / "scripts" / "build.py").read_text(encoding="utf-8")
    check("記事のCTAをA/Bの対象にしている", 'data-ab="article_cta"' in b, True)
    js = (ROOT / "site" / "js" / "site.js").read_text(encoding="utf-8")
    # パラメータはGA4の管理画面で登録しないと集計できない。名前に入れる
    check("A/Bを出来事の名前に入れている", "'cta_click_' + abv" in js, True)
    check("表示側も名前に入れている", "'ab_impression_' + v" in js, True)
    check("結果を読む道具がある", (ROOT / "scripts" / "ab_result.py").is_file(), True)
    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("週次で結果を見る", "ab_result.py" in wf, True)


def test_each_site_declares_what_it_sells():
    """サイトごとに、何を売る記事なのかが決まっていること。

    主力が決まっていないと、書きやすい領域に寄る。実際3サイトとも
    主力が最多ではなかった（ラボはAIO 19.8%、コーポレートはBPO 40.8%、
    補助金はAI導入補助金 37.6%）。読まれても売上につながらない。
    """
    import json
    for f in sorted((ROOT / "sites").glob("*.json")):
        c = json.loads(f.read_text(encoding="utf-8"))
        check(f"{f.stem} に主力がある", bool(c.get("main_offer")), True)
        mix = c.get("category_mix") or c.get("scheme_mix") or {}
        real = {k: v for k, v in mix.items() if not k.startswith("_")}
        check(f"{f.stem} に狙う配分がある", bool(real), True)
        if real:
            check(f"{f.stem} の配分が100%", sum(real.values()), 100)
    brief = (ROOT / "scripts" / "site_brief.py").read_text(encoding="utf-8")
    check("執筆前に主力を見せている", "このサイトで売るもの" in brief, True)


def test_next_keyword_prefers_main_offer():
    """次に書く語は、そのサイトの主力を先に選ぶこと。

    台帳は優先度順に返すが、いま台帳の語はすべて優先度Bで実質は行の並び順。
    主力から離れた古い語が先に出ていた（ラボで「工務店 sns 集客」など）。
    配信済みの管制塔コードは優先度付きの追加に対応しておらず、
    台帳側では順序を変えられないため、取り出す側で寄せている。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import hub_client
    for site, must in (("ai-lab", "aio"), ("subsidy", "ai導入補助金"),
                       ("corporate", "bpo")):
        pat = hub_client._main_offer_pattern(site)
        check(f"{site} の主力の目印がある", bool(pat), True)
        check(f"{site} の目印が主力を含む", must in pat, True)
    src = (ROOT / "scripts" / "hub_client.py").read_text(encoding="utf-8")
    check("next_kw が主力を優先している", "主力優先" in src, True)


def test_improvements_apply_themselves():
    """効果の出ていない記事が、毎週自動で手当てされること。

    「対策した」と「効果があった」は別物で、測って終わりでは何も変わらない。
    条件で決まる手当て（内部リンク）は機械に任せ、文章の判断が要るものだけ
    人（AI）へ渡す。全部を自動にすると、主張のずれた記事が量産される。
    """
    for f in ("effect.py", "auto_improve.py"):
        check(f"{f} がある", (ROOT / "scripts" / f).is_file(), True)
    src = (ROOT / "scripts" / "auto_improve.py").read_text(encoding="utf-8")
    check("内部リンクは自動で当てる", "add_links" in src, True)
    # タイトルの書き換えを機械に任せると、中身の無い約束が並ぶ
    check("タイトルは自動で書き換えない", "def rewrite_title" in src, False)
    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("週次で自動改善が走る", "auto_improve.py --write" in wf, True)
    pr = (ROOT / "automation" / "weekly_optimize_prompt.txt").read_text(encoding="utf-8")
    check("人が判断する分はAIへ渡している", "人が判断すること" in pr, True)
    check("一度に触る本数を制限している", "最大5本" in pr, True)


def test_auto_fixes_are_reviewed():
    """自動で当てた修正が、必ず見直されること。

    当てる側は1本ずつしか見ないため、積み上がりを検知できない。実測では
    「関連して、[X]もあわせてご確認ください。」という同じ一文が内部リンク文
    970本中403本（41.5%）を占め、1記事に11本並んだものまであった。
    1本単位では基準内でも、全体で見れば量産の指紋になる。
    """
    check("見直し役がある", (ROOT / "scripts" / "auto_review.py").is_file(), True)
    rv = (ROOT / "scripts" / "auto_review.py").read_text(encoding="utf-8")
    im = (ROOT / "scripts" / "auto_improve.py").read_text(encoding="utf-8")

    # 当てたら自動で見直しへ繋ぐ。人の判断に委ねると忙しい週に飛ばされる
    check("自動修正のあと見直しが走る", "auto_review.py" in im, True)
    # 検算なしで書き込むと、パターンの取りこぼしが本番の記事を壊す
    check("書き込む前に検算する", "def guard" in rv, True)
    for keep in ("cta-button", "diagnosis", "contact"):
        check(f"リード導線「{keep}」は削除対象から外す", keep in rv, True)
    check("HTMLタグの開閉数を見る", "</figure>" in rv, True)
    check("台帳に残す", "auto_fix.jsonl" in rv and "auto_fix.jsonl" in im, True)
    check("ビルドの機械ゲートに通す", "build.py" in rv, True)

    # 言い回しを自前で書くスクリプトがあると、その一文がサイト中に並ぶ
    for f in ("auto_improve.py", "link_boost.py"):
        src = (ROOT / "scripts" / f).read_text(encoding="utf-8")
        check(f"{f} は自前の定型文を書かない",
              "もあわせてご確認ください" in src, False)
        check(f"{f} は link_new の型を使う", "ln.sentence" in src, True)

    # サイト全体の偏りを測れること（記事単位の基準では捕まらない）
    sys.path.insert(0, str(ROOT / "scripts"))
    import auto_review as ar
    check("全体の偏りに上限がある", ar.MAX_SHARE <= 0.2, True)
    check("1記事のリンク段落に上限がある", ar.MAX_LINK_PARA <= 6, True)

    doc = (ROOT / "CLAUDE.md").read_text(encoding="utf-8")
    check("パイプライン定義に載っている", "自動修正と、その見直し" in doc, True)



def test_client_onboarding_is_one_sheet():
    """ヒアリングシートを1枚埋めれば運用が立ち上がること。

    設定が sites/ と data/ に散っていると、埋め忘れたまま記事を作り始めて
    後から気づく。特に一次情報（その会社にしか出せない数値）が無いまま
    書くと、どのサイトでも書ける記事になりAI検索に引用されない。
    """
    check("ヒアリングシートを作る口がある",
          (ROOT / "scripts" / "client_intake.py").is_file(), True)
    sys.path.insert(0, str(ROOT / "scripts"))
    import client_add as ca
    import client_intake as ci

    # シートは共通＋執筆材料＋被リンクで1枚。fields_for がその全部を返す
    keys = {k for k, *_ in ci.fields_for() if not k.startswith("#")}
    # 主力商材が無いと、表示は増えても相談につながらない記事が量産される
    for k in ("main_offer", "main_category", "categories", "category_mix"):
        check(f"シートで「{k}」を聞いている", k in keys, True)
    # 一次情報はAI検索に引用されるかを決める材料。必須で聞く
    check("一次情報を必須で聞いている",
          any(k == "facts.1.claim" and req for k, _, _, _, req in ci.FIELDS), True)
    check("一次情報の時点も聞いている", "facts.1.as_of" in keys, True)
    for k in ("company.name", "company.address", "cta.label", "cta.url"):
        check(f"シートで「{k}」を聞いている", k in keys, True)

    # シートの回答が、そのまま設定として通ること
    nl = chr(10)
    got = {k: "x" for k in keys}
    got.update({
        "id": "sample-media", "type": "external-html", "repo": "o/r",
        "categories": "sample-a: A" + nl + "sample-b: B",
        "category_mix": "sample-a: 60" + nl + "sample-b: 40",
        "main_category": "sample-a",
        "kw_seeds.industries": nl.join("業種%d" % i for i in range(25)),
        "kw_seeds.intents": nl.join("意図%d" % i for i in range(10)),
        "facts.1.as_of": "2026-09",
    })
    cfg = ci.to_config(got)
    check("シートから設定が組み立つ", bool(cfg["main_offer"] and cfg["categories"]), True)
    check("配分が数値として読める", cfg["category_mix"]["sample-a"], 60)
    check("一次情報が取り出せる", len(ci.to_facts(got, "sample-media")) >= 1, True)
    ng, _ = ci.review(got, cfg)
    check("埋まったシートは不備なしで通る", ng, [])
    # JSONから直接入れる口（client_add）の必須項目も満たせること
    check("JSON側の必須項目もすべて埋まる",
          [k for k in ca.REQUIRED if not cfg.get(k)], [])

    # 受託運用で、クライアントの記事に運用会社の実績が混ざらないこと
    import facts
    check("一次情報をクライアント別に読む", hasattr(facts, "load_for"), True)

    # 記事を書くための材料。空だと「どこにでもある記事」になる
    for k in ("service.list", "customer.faq", "author.name", "tone.style",
              "target.persona", "kw.main", "kw.sub"):
        check(f"シートで「{k}」を聞いている", k in keys, True)
    # 被リンクは買えない。すでにある関係を掘り起こすしかない
    for k in ("link.orgs", "link.partners", "link.press"):
        check(f"外部との接点「{k}」を聞いている", k in keys, True)
    # 業種が変われば必要な材料も変わる
    check("業種別のシートがある", "restaurant" in ci.INDUSTRY, True)
    rkeys = {k for k, *_ in ci.fields_for("restaurant") if not k.startswith("#")}
    for k in ("shop.seats", "menu.signature", "attract.reserve", "flink.gourmet"):
        check(f"飲食店シートで「{k}」を聞いている", k in rkeys, True)

    # メイン×サブから主題が組み立つこと。同じ語の重ねは出さない
    got2 = dict(got)
    got2.update({"kw.main": "梅田 個室 居酒屋", "kw.sub": "梅田 宴会 個室",
                 "kw_seeds.intents": "個室" + nl + "予約",
                 "kw.exclude": "求人"})
    subs = [x["keyword"] for x in ci.subjects(got2)]
    check("サブキーワードが主題になる", "梅田 宴会 個室" in subs, True)
    check("同じ語を重ねない", [s for s in subs if s.count("個室") > 1], [])
    check("狙わない語を除く", [s for s in subs if "求人" in s], [])

    # 執筆エージェントがこれを読めること
    import site_brief
    check("執筆ブリーフが材料を読む", hasattr(site_brief, "show_brief"), True)


def test_quality_gate_holds_on_wordpress():
    """納品方式が変わっても、基準に届かない記事が公開されないこと。

    静的サイトはファイルごと生成しないため物理的に公開できない。
    WordPressはデータベースに入ってしまうので、公開ステータスへの
    遷移を先方側で止める必要がある。配信スクリプトだけで見ていると、
    管理画面から直接投稿された記事を止められない。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import client_add as ca
    import publish
    check("WordPress納品に対応している", "wordpress" in ca.TYPES, True)
    check("WordPressへ配信できる", hasattr(publish, "write_wordpress"), True)

    php = ROOT / "automation" / "wordpress" / "ss-quality-gate.php"
    check("品質ゲートのプラグインがある", php.is_file(), True)
    src = php.read_text(encoding="utf-8")
    # 管理画面から止められる場所に置くと「忙しいので一旦切った」が起きる
    check("mu-plugins に置く指示がある", "mu-plugins" in src, True)
    # 個別の入口をふさぐ形にすると、入口が増えたときに漏れる
    check("すべての保存が通る関所で止める", "wp_insert_post_data" in src, True)
    check("公開を下書きへ戻す", "'draft'" in src or '"draft"' in src, True)
    check("予約投稿も対象にする", "future" in src, True)
    check("未採点は公開させない", "品質スコアがありません" in src, True)
    check("止めた理由を管理画面に出す", "admin_notices" in src, True)
    # 配信側でも見る。二重に見ないと、どちらか一方の抜け道が残る
    pub = (ROOT / "scripts" / "publish.py").read_text(encoding="utf-8")
    check("配信側も90点未満を止める", "公開基準未達" in pub, True)
    check("公開されなかったことに気づける", "公開されませんでした" in pub, True)


def test_daily_todo_fits_in_one_run():
    """同日救済が1回で終わる量にとどまること。

    定期実行は数時間ずれる。実際、21:30の救済が翌03:07に走り、
    深夜に「本日0本」を見て6本書こうとしてターン上限で落ちた。
    さらにTODOが13件あり、救済プロンプトが説明しているのは4種類だけだった。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import daily_audit as da
    # 公開の時刻を過ぎていないものを「不足」と言わない
    check("公開時刻を持っている", bool(da.PUBLISH_HOURS), True)
    check("1回で扱う上限がある", da.MAX_TODAY <= 10, True)
    now, later = da.split_todo([
        "TODO: ai-lab の記事を本日あと 2 本作成して公開する",
        "TODO: 「aio診断」で自社ページが競合（ai-lab）。…",
        "TODO: リライト候補が223件ある。…",
        "TODO: x は subsidy の既存記事と重複（類似度0.6・url）。…",
    ])
    check("記事の作成が先頭に来る", "の記事を本日あと" in now[0], True)
    check("カニバリは今日やらない", any("競合" in x for x in later), True)
    check("リライト候補は今日やらない", any("リライト候補" in x for x in later), True)
    check("重複301は今日やる", any("既存記事と重複" in x for x in now), True)

    # 救済プロンプトが、今日やるTODOの全種類を説明していること
    pr = (ROOT / "automation" / "retry_prompt.txt").read_text(encoding="utf-8")
    for key in ("本作成して公開する", "品質基準まで直して", "の領域。",
                "既存記事と重複", "のKWを補充する"):
        check(f"救済プロンプトが「{key}」を説明している", key in pr, True)
    check("NOTE行は触らないと書いてある", "NOTE" in pr, True)

    # 似ているだけの記事を毎日TODOに出さない（実害で判定する）
    import cannibal_check as cc
    check("重複を実績で仕分ける", hasattr(cc, "judge_overlap"), True)
    check("食い合っていないものは出さない",
          "clear" in (cc.judge_overlap.__doc__ or ""), True)


def test_near_page1_is_pushed_every_week():
    """1ページ目に近い記事が、放置されずに毎週底上げされること。

    順位そのものは約束できない。決めるのはGoogleで、反映にも数ヶ月かかる。
    できるのは「こちら側でやれることを漏れなく続ける」ことだけなので、
    そこを仕組みに固定する。
    """
    check("底上げの工程がある", (ROOT / "scripts" / "priority_boost.py").is_file(), True)
    sys.path.insert(0, str(ROOT / "scripts"))
    import priority_boost as pb
    # 主力の語に絞る。関係ない語で上位を取っても相談につながらない
    for sid in ("ai-lab", "corporate", "subsidy"):
        check(f"{sid} の主力の語を持っている", bool(pb.MAIN_PATTERN.get(sid)), True)
    # 11〜30位。20.5で切っていたため21〜30位の138語・表示801回が対象外だった
    check("あと少しの帯を狙う", pb.NEAR[0] >= 10 and pb.NEAR[1] >= 30, True)
    check("被リンクの下限がある", pb.INBOUND_FLOOR >= 10, True)
    # タイトルの判断は人に残す。機械が当てると主張のずれた記事になる
    src = (ROOT / "scripts" / "priority_boost.py").read_text(encoding="utf-8")
    check("タイトルは自動で書き換えない", "人が判断して直す" in src, True)
    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("毎週走る", "priority_boost.py --write" in wf, True)


def test_main_category_actually_grows():
    """主力カテゴリが、狙いの配分に向かって増えること。

    site_brief が「次に書くなら◯◯」と出していても、執筆する側が
    それに従わなければ配分は動かない。実測で、AI集客ラボの主力（aio）は
    21%（狙い45%）まで落ち、直近21日の新規記事でも22%のままだった。
    指示が出ているだけでは足りず、従わせる一文が要る。
    """
    pr = (ROOT / "automation" / "multi_site_prompt.txt").read_text(encoding="utf-8")
    check("配分に従う指示がある", "次に書くなら" in pr, True)
    check("主力が不足すると何が起きるか書いてある",
          "相談につながらない" in pr, True)

    # 主力優先の目印が、主力カテゴリの語だけを拾うこと。
    # 広すぎると別カテゴリの語まで主力扱いになり、配分が動かない
    sys.path.insert(0, str(ROOT / "scripts"))
    import re as _re
    import hub_client as hb
    pat = _re.compile(hb._main_offer_pattern("ai-lab"), _re.I)
    check("主力の語を拾う", bool(pat.search("aio 診断")), True)
    check("別カテゴリの語は拾わない", bool(pat.search("chatgpt 集客 方法")), False)
    for sid in ("ai-lab", "corporate", "subsidy"):
        check(f"{sid} に主力の目印がある", bool(hb._main_offer_pattern(sid)), True)


def test_paragraph_split_keeps_text():
    """長い段落を分けても、地の文が変わらないこと。

    読みにくい段落は機械で分けられるが、文の途中で切れば意味が壊れる。
    「。」の直後だけで分け、分けた前後で文字が1字でも増減したら
    その記事は書き換えない。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import split_paragraphs as sp
    check("分ける工程がある", (ROOT / "scripts" / "split_paragraphs.py").is_file(), True)

    # 文の途中では切らない
    nl = chr(10)
    body = nl.join(["---", "title: x", "---", ("あ" * 120 + "です。") * 3])
    after, n = sp.process(body)
    check("長い段落を分ける", n > 0, True)
    check("地の文は変わらない", sp.plain(body) == sp.plain(after), True)
    check("「。」の直後で分かれる",
          all(p.rstrip().endswith(("。", "」", "）")) or "あ" not in p
              for p in after.split(nl + nl)[1:]), True)

    # 表・リスト・HTMLの塊には触らない
    check("表は触らない", sp.splittable("| 列A | 列B |" + "あ" * 300), False)
    check("リストは触らない", sp.splittable("- " + "あ" * 300), False)
    check("HTMLの塊は触らない", sp.splittable("<div>" + "あ" * 300 + "</div>"), False)

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("毎週走る", "split_paragraphs.py --all --write" in wf, True)


def test_anchor_text_stays_readable():
    """内部リンクのアンカーが、読める長さに収まること。

    自動で入れるリンク文は記事タイトルをそのままアンカーにする。
    タイトルは45字まで許されるため、前後の文と合わせると1文が100字を超える。
    実測で公開317本中49本がこれで警告に引っかかっていた。
    リンク先は絶対に変えず、表示される文字だけを意味の切れ目で詰める。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import link_new as ln
    import shorten_anchors as sa
    check("詰める工程がある", (ROOT / "scripts" / "shorten_anchors.py").is_file(), True)

    long_title = "飲食店のAI導入補助金 対象要件｜資本金・従業員数など5つの基準【2026年】"
    s = sa.shorten(long_title)
    check("長いアンカーを詰める", len(s) < len(long_title), True)
    check("詰めた結果が短すぎない", len(s) >= sa.MIN, True)
    check("元の文字列から作る（作文しない）",
          all(c in long_title for c in s), True)

    # 短いものには触らない
    check("短いアンカーは触らない", sa.shorten("AIO対策の始め方"), "AIO対策の始め方")

    # 今後生成されるリンクも詰まっていること（そうしないと毎週また溜まる）
    for fit in (0.1, 0.9):
        for seed in range(8):
            line = ln.sentence(long_title, "/aio/x/", seed, fit)
            anchor = line[line.index("[") + 1:line.index("](")]
            check(f"生成されるアンカーが{sa.MAX}字以内（seed={seed} fit={fit}）",
                  len(anchor) <= sa.MAX, True)

    # リンク先は変えない
    import re as _re
    body = f"[{long_title}](/subsidy/inshokuten-hojokin/)"
    out = _re.sub(r"\[([^\]]+)\]", lambda m: "[" + sa.shorten(m.group(1)) + "]", body)
    check("リンク先は変わらない", "(/subsidy/inshokuten-hojokin/)" in out, True)

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("毎週走る", "shorten_anchors.py --write" in wf, True)


def test_block_breaks_render_correctly():
    """段落の切れ目が壊れていないこと。

    リンク挿入と段落分割が、それぞれ別の壊し方をしていた。どちらも記事の点数には
    出ず、画面で初めて分かる。表がパイプ記号のまま本文に出て、** がそのまま表示される。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import fix_block_breaks as fb
    import split_paragraphs as sp
    nl = chr(10)

    # 1. 空行なしで続く表・リストの前に空行を入れる
    for block in ("| 列A | 列B |", "1. 最初の手順", "- 箇条書き"):
        out, n = fb.add_blank_lines("本文です。" + nl + block)
        check(f"「{block[:6]}」の前に空行を入れる", n == 1 and nl + nl in out, True)
    out, n = fb.add_blank_lines("1. 一つ目" + nl + "2. 二つ目")
    check("リストの行同士には入れない", n, 0)

    # 2. 装飾の途中で分かれた段落を戻す
    broken = "**失敗1: 土台を飛ばす。" + nl + nl + "**理由はこうです。"
    out, n = fb.rejoin_marks(broken)
    check("** が分かれた段落を戻す", n == 1 and out.count(nl + nl) == 0, True)
    ok = "普通の段落です。" + nl + nl + "次の段落です。"
    check("閉じている段落は触らない", fb.rejoin_marks(ok)[1], 0)

    # 3. 生HTMLの中の ** を <strong> にする（Markdownが処理しないため）
    out, n = fb.strong_in_html('<div class="caution-box">罰金は**100万円**です。</div>')
    check("生HTMLの ** を <strong> にする", n == 1 and "<strong>100万円</strong>" in out, True)
    check("地の文の ** は触らない", fb.strong_in_html("本文の**強調**です。")[1], 0)

    # 4. 分割側でも、装飾の内側では切らない（同じ壊れ方を作らない）
    body = nl.join(["---", "title: x", "---",
                    "あ" * 100 + "。**" + "い" * 100 + "。" + "う" * 100 + "**。"])
    after, _ = sp.process(body)
    for para in after.split(nl + nl):
        check("分けた各段落で ** が閉じている", para.count("**") % 2, 0)

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("毎週走る", "fix_block_breaks.py --write" in wf, True)


def test_rendered_html_is_checked():
    """出来上がったHTMLを見る検査が、ビルドの中で動くこと。

    これまでの検査は全部、原稿（Markdown）だけを見ていた。そのため
    「原稿は正しいが変換すると壊れる」崩れを長く見逃した。表がパイプ記号の
    まま出る・** がそのまま表示される・リンクが押せない、の3種。
    どれも記事の点数には出ず、読者の画面にだけ出る。
    """
    import tempfile
    sys.path.insert(0, str(ROOT / "scripts"))
    import md2html
    import render_check as rc

    def md2html_convert(src):
        return md2html.convert(src)[0]

    # ビルドが必ず通ること（単体スクリプトのままだと呼ばれずに終わる）
    b = (ROOT / "scripts" / "build.py").read_text(encoding="utf-8")
    check("ビルドが描画検査を呼ぶ", "render_check.scan(" in b, True)

    # 実際に捕まえること。素通りする検査は無いのと同じ
    cases = {
        "強調記号がそのまま出ている": "<article><p>罰金は**100万円**です。</p></article>",
        "Markdownリンクが押せない形で出ている":
            "<article><p>詳細は[解説](/seo/x/)です。</p></article>",
        "表の区切り記号が本文に出ている":
            "<article><p>| 列A |" + chr(10) + "|:--|</p></article>",
        "altの無い画像がある": '<article><p><img src="a.png"></p></article>',
        "押せないリンクがある": '<article><p><a href="">文字</a></p></article>',
        "句点が続いている": "<article><p>そうです。。</p></article>",
    }
    d = Path(tempfile.mkdtemp())
    for i, (name, src) in enumerate(cases.items()):
        f = d / str(i) / "index.html"
        f.parent.mkdir(parents=True)
        f.write_text(src, encoding="utf-8")
        check(f"崩れを見つける: {name}", name in rc.scan([f]), True)

    # 正しいページを間違いと言わないこと
    f = d / "ok" / "index.html"
    f.parent.mkdir(parents=True)
    f.write_text('<article><p>罰金は<strong>100万円</strong>、詳細は'
                 '<a href="/seo/x/">解説</a>。</p>'
                 '<img src="a.png" alt="図"></article>', encoding="utf-8")
    check("正しいページを誤検出しない", dict(rc.scan([f])), {})

    # 配信の全方式が同じ検査を通ること。writerごとに書くと方式が増えたとき漏れる。
    # 実際 external-html と wordpress には検査が無く、表がパイプ記号のまま
    # 5本配信されていた
    src = (ROOT / "scripts" / "publish.py").read_text(encoding="utf-8")
    body = re.search(r"(?s)def main\(\):.*?if cfg\[.type.\] == .self-static.", src)
    check("配信前の検査が main にある（方式ごとではなく）",
          bool(body and "render_check.problems(" in body.group(0)), True)
    check("writerの中に個別の検査を残していない",
          "raw_markdown_left(" in src, False)

    # 表の列数が合わないと、Markdownは表として扱わず本文にそのまま出す
    n = chr(10)
    broken = md2html_convert("| A | B | C |" + n + "|:--|:--|" + n + "| 1 | 2 | 3 |")
    check("列数の合わない表を崩れとして見つける",
          any("表の区切り記号" in x[0] for x in rc.problems(broken)), True)
    good = md2html_convert("| A | B | C |" + n + "|:--|:--|:--|" + n + "| 1 | 2 | 3 |")
    check("列数の合う表は崩れと数えない", rc.problems(good), [])

    # コード例の中の記法は崩れではない（誤検出すると検査が信用されなくなる）
    f2 = d / "code" / "index.html"
    f2.parent.mkdir(parents=True)
    f2.write_text("<article><pre><code>**太字**の書き方</code></pre></article>",
                  encoding="utf-8")
    check("コード例は崩れと数えない", dict(rc.scan([f2])), {})


def test_import_has_no_side_effects():
    """読み込むだけで外へ出ていくスクリプトが無いこと。

    notify_indexnow.py は処理をモジュール直下に書いていたため、
    `import notify_indexnow` した瞬間にIndexNowへ送信していた。
    実際、検査のために読み込んだだけで136件のURLを送ってしまった。
    副作用は必ず main() の中に置き、__main__ ガードで囲う。
    """
    import ast
    print("\n■ 読み込みの副作用")
    OUT = ("urlopen", "urlretrieve", "post", "put", "get", "request", "run",
           "check_output", "call", "Popen", "write_text", "write_bytes", "unlink")
    ng = []
    for p in sorted((ROOT / "scripts").glob("*.py")):
        src = p.read_text(encoding="utf-8", errors="replace")
        tree = ast.parse(src)
        has_main = any(isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
                       and n.name == "main" for n in tree.body)
        if has_main and '__name__ == "__main__"' not in src \
                and "__name__ == '__main__'" not in src:
            ng.append(f"{p.name}: main() があるのに __main__ ガードが無い")
        # モジュール直下で外に出ていく呼び出しをしていないか
        for n in tree.body:
            if not isinstance(n, ast.Expr):
                continue
            for c in ast.walk(n):
                if isinstance(c, ast.Call):
                    f = c.func
                    name = f.attr if isinstance(f, ast.Attribute) else \
                        (f.id if isinstance(f, ast.Name) else "")
                    if name in OUT:
                        ng.append(f"{p.name}: 読み込み時に {name}() を呼んでいる")
    check("読み込むだけで動くスクリプトが無い", ng, [])

    # 実際に読み込んでも何も起きないこと（構文だけ見ても分からない）
    import importlib
    sys.path.insert(0, str(ROOT / "scripts"))
    broke = []
    for p in sorted((ROOT / "scripts").glob("*.py")):
        try:
            importlib.import_module(p.stem)
        except SystemExit:
            broke.append(f"{p.name}: 読み込みで終了した")
        except Exception as e:
            broke.append(f"{p.name}: {type(e).__name__}")
    check("全スクリプトが安全に読み込める", broke, [])

    # requirements.txtに無い外部ライブラリをimportしていないか（関数内の遅延importも含む）
    # （実行環境にたまたま入っていただけの依存は、入っていない環境で上のチェックが再現しない。
    #  openpyxlが関数内でしかimportされておらず、tree.body直下しか見ていなかったために
    #  requirements.txtへの記載漏れを検出できなかった実例がある。ast.walkで全体を見る）
    import importlib.metadata
    dist_map = importlib.metadata.packages_distributions()
    req_names = set()
    for line in (ROOT / "requirements.txt").read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#"):
            req_names.add(re.split(r"[<>=\[]", line)[0].strip().lower().replace("_", "-"))
    local_modules = {p.stem for p in (ROOT / "scripts").glob("*.py")}
    missing_req = []
    for p in sorted((ROOT / "scripts").glob("*.py")):
        tree = ast.parse(p.read_text(encoding="utf-8", errors="replace"))
        for n in ast.walk(tree):
            mods = []
            if isinstance(n, ast.Import):
                mods = [a.name.split(".")[0] for a in n.names]
            elif isinstance(n, ast.ImportFrom) and n.level == 0 and n.module:
                mods = [n.module.split(".")[0]]
            for m in mods:
                if m in sys.stdlib_module_names or m in local_modules:
                    continue
                dists = {d.lower().replace("_", "-") for d in dist_map.get(m, [])}
                if dists and not (dists & req_names):
                    missing_req.append(f"{p.name}: {m} ({'/'.join(dists)}) がrequirements.txtに無い")
    check("外部importがrequirements.txtに宣言されている（関数内含む）", missing_req, [])


def test_long_sentences_are_split():
    """長すぎる1文が、意味を変えずに分かれること。

    切ってよいのは「左がそれだけで文として成り立つ形」だけ。
    「〜によると、」「〜し、」で切ると主語と述語がねじれる。
    「が」は逆接とは限らず、前置きの「が」に「ただし」を足すと意味がずれる。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import split_sentences as ss
    nl = chr(10)

    long_ok = "あ" * 40 + "を運用しており、" + "い" * 40 + "という結果になりました。"
    out = ss.split_sentence(long_ok)
    check("連用中止法で分ける", out and "しています。" in out, True)

    reason = "あ" * 40 + "が公開されているため、" + "い" * 40 + "を先に確認します。"
    out = ss.split_sentence(reason)
    check("理由の「ため」を「そのため」で受け直す",
          out and "そのため、" in out, True)

    # 左が文にならない形では切らない
    for bad in ("によると、", "を使い、", "が必要なため、"):
        s = "あ" * 40 + bad + "い" * 40 + "です。"
        out = ss.split_sentence(s)
        if bad == "が必要なため、":
            continue          # 「ため」は受け直せるので対象
        check(f"「{bad}」では切らない", out, None)

    # 前置きの「が」には「ただし」を足さない
    preface = "あ" * 40 + "について解説していますが、" + "い" * 20 + "は対象外です。"
    check("前置きの「が」は切らない", ss.split_sentence(preface), None)

    # 装飾の内側では切らない
    deco = "あ" * 30 + "**" + "い" * 20 + "しており、" + "う" * 20 + "**" + "え" * 30 + "です。"
    out = ss.split_sentence(deco)
    check("装飾の内側では切らない", out is None or out.count("**") % 2 == 0, True)

    # 検算: 記録した書き換え以外は動いていない
    body = nl.join(["---", "title: x", "---", long_ok])
    after, n, edits = ss.process(body)
    check("分けた数を記録する", n, len(edits))
    check("検算が通る", ss.verify(body, after, edits), "")
    check("食い違いを見つける",
          bool(ss.verify(body, after + "余計な文字", edits)), True)

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("毎週走る", "split_sentences.py --write" in wf, True)


def test_rewrite_is_verified_and_reverted():
    """文章の判断を任せた直しが、必ず検算され、外れたら戻ること。

    auto_improve は、タイトルと検索意図の見直しを人へ回していた。
    判断そのものは機械に任せてよいが、**通してよいかを決めるのは機械側**でなければ
    ならない。検算が無い書き換えは、主張のずれた記事を静かに量産する。
    """
    src = (ROOT / "scripts" / "auto_rewrite.py").read_text(encoding="utf-8")
    check("直しの工程がある", (ROOT / "scripts" / "auto_rewrite.py").is_file(), True)

    # 外れたら必ず元に戻す
    check("検算に外れたら元に戻す", 'git", "checkout", "--"' in src, True)
    check("戻したあとビルドし直す",
          src.count("build.py") >= 2, True)

    # 検算の中身。1つでも欠けると、その穴から壊れた記事が通る
    for name, key in (("数字を作らせない", "numbers("),
                      ("出典を消させない", "sources("),
                      ("狙う語を変えさせない", "狙う語が変わりました"),
                      ("タイトルの字数", "TITLE_MIN"),
                      ("警告を増やさせない", "警告が増えました"),
                      ("食い合いを作らせない", "kw_guard"),
                      ("ビルドを通す", "ビルドが通りません"),
                      ("他の記事に触らせない", "別の記事まで変わっています")):
        check(f"検算: {name}", key in src, True)

    # 台帳に残す（何をいつ直したか追えないと、戻す判断ができない）
    check("台帳に残す", "auto_fix.jsonl" in src, True)

    # 1回に触る本数を絞る（まとめて当てると原因が分からなくなる）
    m = re.search(r'"--limit", type=int, default=(\d+)', src)
    check("1回に触る本数を絞る", bool(m) and int(m.group(1)) <= 5, True)

    # 全権限を飛ばさない。使う道具を読み書きだけに絞る
    check("権限を全部は飛ばさない", "dangerously-skip-permissions" in src, False)
    check("使う道具を絞る", "--allowedTools" in src, True)

    # 対象を取りこぼさない（画面表示は8件で打ち切られる）
    sys.path.insert(0, str(ROOT / "scripts"))
    import auto_improve
    check("全件を返す口がある", hasattr(auto_improve, "human_items"), True)

    # 少ない表示回数の増減はノイズ。ここを根拠に記事を書き換えてはいけない
    import effect
    def act(before_imp, after_imp, pos=30.0):
        return effect.actions([{"slug": "x", "site": "ai-lab",
                                "before": (before_imp, 0, pos),
                                "after": (after_imp, 0, pos)}])
    check("表示1→0では書き換えない", act(1, 0), [])
    check("表示4→0では書き換えない", act(4, 0), [])
    check("表示15→2では書き換えない", act(15, 2), [])
    check("表示94→9なら見直す",
          [x["do"] for x in act(94, 9)], ["review"])
    check("下限は他の分岐と同じ値", effect.MIN_IMP, 20)

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("毎週走る", "auto_rewrite.py --write" in wf, True)
    # 検算が壊れたままエージェントを動かすと、悪い書き換えがそのまま公開される。
    # 先に自己診断を通し、落ちたら書き換えを飛ばす
    check("先に検算の自己診断を通す",
          wf.index("--selftest") < wf.index("auto_rewrite.py --write"), True)
    check("診断に落ちたら書き換えない",
          "steps.guard.outputs.ok == 'yes'" in wf, True)
    check("自己診断の口がある", "--selftest" in src, True)
    check("自己診断は必ず元に戻す", "finally:" in src, True)
    # 積み上がりの見直しより前に置く。後ろだと、当てた週は見直されないまま公開される
    check("見直しより前に置く",
          wf.index("auto_rewrite.py") < wf.index("auto_review.py --fix"), True)


def test_output_matches_source():
    """原稿に書いたものが、出力に同じ数だけ出ていること。

    症状を並べる検査は「こちらが知っている壊れ方」しか見つけられない。
    実際、表が生で出る・** がそのまま出る・リンクが押せない・段落が繋がる、の
    4種類が長期間だれにも気づかれずに公開されていた。どれも点数には出なかった。
    ここでは中身を見ず、塊の数だけを突き合わせる。知らない壊れ方でも数は合わなくなる。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import md2html
    import render_check as rc
    nl = chr(10)
    print(nl + "■ 原稿と出力の突き合わせ")

    def gap(src):
        return rc.structure_gap(src, "<article>" + md2html.convert(src)[0] + "</article>")

    # 実際に起きた4種類を、症状名を使わずに見つけられること
    check("列数の合わない表を見つける",
          bool(gap("| A | B | C |" + nl + "|:--|:--|" + nl + "| 1 | 2 | 3 |")), True)
    check("空行なしのリストを見つける",
          bool(gap("本文です。" + nl + "- 一つ目" + nl + "- 二つ目")), True)
    check("空行なしの表を見つける",
          bool(gap("本文です。" + nl + "| A | B |" + nl + "|:--|:--|" + nl + "| 1 | 2 |")), True)
    check("生HTML内のリンクを見つける",
          bool(gap('<div class="box">詳細は[解説](/seo/x/)です。</div>')), True)

    # 正しいものを間違いと言わない（誤検出すると検査が信用されなくなる）
    check("正しい表は素通り", gap("| A | B |" + nl + "|:--|:--|" + nl + "| 1 | 2 |"), [])
    check("正しいリストは素通り",
          gap("本文です。" + nl + nl + "- 一つ目" + nl + "- 二つ目"), [])
    check("コード例の中は数えない",
          gap("```" + nl + "| これは表ではない |" + nl + "```"), [])

    b = (ROOT / "scripts" / "build.py").read_text(encoding="utf-8")
    check("ビルドが突き合わせる", "structure_gap(" in b, True)


def test_inquiry_body_is_not_truncated():
    """届いた相談の本文が、途中で切れないこと。

    clean_() は件名や会社名のための関数で、改行を空白に潰して80字で切る。
    これを相談内容にも使っていたため、実際に届いた問い合わせが
    「インド人新卒採用支援サービスを…検索され」でちょうど80字で途切れた。
    メールだけでなく台帳にも切れたまま保存されており、続きは復元できない。

    本文には body_() を使う。改行を残し、送信側の上限（2000字）を超える値にする。
    """
    import re as _re
    print(chr(10) + "■ 問い合わせ本文の扱い")
    gs = (ROOT / "automation" / "gas" / "contact.hub.gs").read_text(encoding="utf-8")

    check("本文用の関数がある", "function body_(" in gs, True)
    # 本文に clean_ を使っていないこと（ここが事故の原因だった）
    for field in ("d.message", "d.body", "d.message || d.body"):
        check(f"clean_({field}) を使っていない",
              f"clean_({field})" in gs, False)
    check("台帳へは body_ で書く", "body_(d.message || d.body)" in gs, True)
    check("メールへは body_ で出す", "body_(d.message)" in gs, True)

    # 上限が、送信側（lead.js）の上限より大きいこと。小さいとそこで切れる
    m = _re.search(r"function body_\(s\)[\s\S]*?slice\(0, *(\d+)\)", gs)
    js = (ROOT / "functions" / "api" / "lead.js").read_text(encoding="utf-8")
    m2 = _re.search(r"String\(v\)\.slice\(0, *(\d+)\)", js)
    gas_max = int(m.group(1)) if m else 0
    web_max = int(m2.group(1)) if m2 else 0
    check(f"上限が送信側（{web_max}字）を上回る", gas_max > web_max, True)

    # 改行を潰していないこと（段落が読めなくなる）
    body_fn = m.group(0) if m else ""
    check("改行を空白に潰していない",
          _re.search(r"replace\(/\[\r\n\]\+/g, *['\"] ['\"]\)", body_fn) is None, True)

    # 件名や会社名は従来どおり clean_（1行に収める必要がある）
    check("件名は clean_ のまま", "clean_(d.company)" in gs, True)


def test_articles_are_not_uniform():
    """記事の長さが1つの帯に固まっていないこと。

    CLAUDE.md は「全記事が同じ長さに揃うと、それ自体が量産の指紋になる」と
    書いているのに、実測では312本（92%）が5,000字台に固まっていた。
    depth も273本（81%）が standard のまま。仕組みはあるのに使われていなかった。
    原因は、執筆プロンプトに長さの指示が1行も無かったこと。

    透かし（SynthID）は関係ない。この仕組みは画像をPillowで描いており、
    AI画像生成を使っていないので透かし自体が存在しない。
    見られているのは「書き方が揃っていること」のほうである。
    """
    print(chr(10) + "■ 記事の均一さ")
    pr = (ROOT / "automation" / "multi_site_prompt.txt").read_text(encoding="utf-8")
    check("執筆プロンプトが長さを指示する", "depth:" in pr, True)
    for d in ("quick", "standard", "deep"):
        check(f"{d} の使いどころが書いてある", d in pr, True)
    check("水増しを禁じている", "水増し" in pr, True)
    check("冒頭の型を固定しないよう書いてある", "とは" in pr and "型を固定" in pr, True)

    ad = (ROOT / "scripts" / "daily_audit.py").read_text(encoding="utf-8")
    check("日次監査が偏りの中身まで出す", "最も多い帯" in ad, True)
    # ばらつきの数字だけでは、どれだけ固まっているか読み取れなかった
    check("1帯への集中を見ている", "share >= 0.5" in ad, True)


def test_facts_are_per_article():
    """一次情報が、記事ごとに違うものになること。

    会社の実績（3,200店舗）はどの記事に書いても同じ一文になる。
    実測では126本（37%）が同じ数字を載せていた。肩書きとしては正しいが、
    「その記事にしかない情報」にはならず、AI検索の引用先にも選ばれにくい。

    自サイトのGSC実測は記事ごとに違う。「この語が何位で、何回表示され、
    何回クリックされたか」は、ほかのどのサイトも書けない。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import facts
    print(chr(10) + "■ 記事ごとの一次情報")

    check("記事ごとの実測を出す口がある", hasattr(facts, "own_query_facts"), True)

    rows = [{"kw": "口コミ 返信 例文 医療", "pos": 8.7, "imp": 96, "clicks": 0, "ctr": 0.0},
            {"kw": "aio対策 自分で", "pos": 26.2, "imp": 28, "clicks": 1, "ctr": 3.6},
            {"kw": "原口優", "pos": 1.3, "imp": 36, "clicks": 4, "ctr": 11.1},
            {"kw": "少ない語", "pos": 5.0, "imp": 3, "clicks": 0, "ctr": 0.0}]
    import json as _j, tempfile
    d = Path(tempfile.mkdtemp())
    (d / "t.json").write_text(_j.dumps({"2026-09-17": rows}), encoding="utf-8")
    old = facts.RANKS
    facts.RANKS = d
    try:
        a = facts.own_query_facts("t", "口コミ 返信 例文 医療")
        b = facts.own_query_facts("t", "aio対策 自分で")
        check("語ごとに違う数字が出る",
              bool(a) and bool(b) and a[0]["claim"] != b[0]["claim"], True)
        # 上位なのにクリックが無い語は、自社の無力さではなく検索の性質として書く
        check("0クリックを自社の弱みとして書かない",
              "クリックは0回" not in a[0]["claim"], True)
        check("0クリックを検索の性質として説明する",
              "答えが済む" in a[0]["claim"], True)
        check("数字はそのまま出す", "96回" in a[0]["claim"] and "8.7位" in a[0]["claim"], True)
        check("出典と時点が付く",
              bool(a[0].get("source")) and bool(a[0].get("as_of")), True)
        # 表示が少ない語は偶然と区別できないので出さない
        check("表示が少ない語は出さない", facts.own_query_facts("t", "少ない語"), [])
        check("語が無ければ何も出さない", facts.own_query_facts("t", ""), [])
    finally:
        facts.RANKS = old

    pr = (ROOT / "automation" / "multi_site_prompt.txt").read_text(encoding="utf-8")
    check("プロンプトが記事ごとの数字を優先させる", "その記事にしかない数字" in pr, True)
    check("丸めを禁じている", "丸めたり" in pr, True)


def test_unindexed_pages_are_chased():
    """公開したのに検索に出ていないページを、放置しないこと。

    書いて配信したところで満足すると、載っていないページが溜まる。
    実測で、公開30日以上たっても28日間まったく表示されないページが19件あった。
    コーポレートは sitemap の40%（53件）が表示ゼロだった。

    拾う仕組み（index_status.py / reindex.py）は作られていたが、
    **どこからも呼ばれていなかった**。作っただけで動いていない道具は無いのと同じ。
    """
    print(chr(10) + "■ 未掲載ページの追跡")
    for f in ("index_status.py", "reindex.py"):
        check(f"{f} がある", (ROOT / "scripts" / f).is_file(), True)

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("毎週、登録状況を見て再通知する", "reindex.py" in wf, True)

    ad = (ROOT / "scripts" / "daily_audit.py").read_text(encoding="utf-8")
    check("日次監査が気づく", "def check_unseen" in ad, True)
    check("監査から呼ばれている", "check_unseen(todo)" in ad, True)
    # 30日未満は評価が定まっていないだけ。手を当てても意味がない
    check("30日を境にしている", "days < 30" in ad, True)
    # 記事は所属サイトでだけ数える（全サイトで照合すると1本を3回数える）
    check("所属サイトだけで数える", "find_category_owner" in ad, True)
    check("直し方を示す", "reindex.py" in ad, True)


def test_rate_claims_need_evidence():
    """割合を書くなら、母数と集計期間を必ず添えること。

    「採択率90%」だけでは、何社中何社か、いつからいつまでかが分からない。
    根拠を示さない割合は優良誤認になる（景品表示法）。
    人の注意に任せると、忙しいときに抜ける。登録の時点で機械が止める。
    """
    import json as _j, subprocess as _s, tempfile
    print(chr(10) + "■ 一次情報の登録条件")
    d = Path(tempfile.mkdtemp())

    def judge(obj):
        f = d / "f.json"
        f.write_text(_j.dumps(obj, ensure_ascii=False), encoding="utf-8")
        r = _s.run([sys.executable, str(ROOT / "scripts" / "add_fact.py"),
                    "--check", str(f)], cwd=ROOT, capture_output=True,
                   text=True, encoding="utf-8", errors="ignore")
        return r.returncode == 0

    base = {"sites": ["subsidy"], "topic": ["補助金"],
            "source": "自社実績", "as_of": "2026-09"}
    check("母数も期間も無い割合は止める",
          judge({**base, "id": "x1", "claim": "当社の採択率は90%です"}), False)
    check("母数を本文に書いていなければ止める",
          judge({**base, "id": "x2", "claim": "当社の採択率は90.5%です",
                 "denominator": 42, "period": "2025-04〜2026-03"}), False)
    check("母数が小さすぎれば止める",
          judge({**base, "id": "x3", "denominator": 3, "period": "2026-01〜2026-09",
                 "claim": "2026年に支援した3社のうち3社が採択されました（採択率100%）"}), False)
    check("雛形のままは止める",
          judge({**base, "id": "x4", "claim": "＜例＞◯◯社を支援しました"}), False)
    check("数値が無ければ止める",
          judge({**base, "id": "x5", "claim": "多くの企業を支援してきました"}), False)
    check("母数と期間を本文に書いてあれば通す",
          judge({**base, "id": "x6", "denominator": 42, "period": "2025-04〜2026-03",
                 "claim": "2025年4月〜2026年3月に支援した42社のうち38社が採択されました"
                          "（採択率90.5%）"}), True)
    # 実数（割合でない）は母数を求めない
    check("割合でなければ母数は求めない",
          judge({**base, "id": "x7", "claim": "2026年9月までに42社の申請を支援しました"}), True)


def test_submissions_are_not_double_counted():
    """送信の数を、二重に数えないこと。

    site.js は1回の送信で form_submit と lead_capture の両方を発火させる。
    両方を足していたため「送信10件」と出ていたが、実際は4件だった
    （form_submit 4 + lead_capture 6）。転換率を2倍以上に見せてしまう。

    問い合わせと購読も分ける。同じ箱に入れると、商談につながる数が分からない。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import funnel
    print(chr(10) + "■ 送信数の数え方")
    names = dict(funnel.STEPS)["送信した"]
    check("lead_capture を送信に数えない", "lead_capture" in names, False)
    check("form_submit は数える", "form_submit" in names, True)
    # 同じ送信で飛ぶイベントを2つ以上足していないか
    fired_together = {"form_submit", "lead_capture", "lead_newsletter"}
    check("同時に飛ぶイベントを重ねて数えない",
          len(fired_together & set(names)) <= 1, True)

    routes = dict(funnel.LEAD_ROUTES)
    check("問い合わせと購読を分ける口がある",
          "問い合わせ・相談" in routes and "ニュースレター購読" in routes, True)
    check("購読を問い合わせに含めない",
          "newsletter" in routes["問い合わせ・相談"], False)

    src = (ROOT / "scripts" / "funnel.py").read_text(encoding="utf-8")
    # 内訳が出せないときは、黙らずに何をすればよいかを出す
    check("設定が要るときに案内する", "カスタムディメンションを作成" in src, True)
    check("取れなくても落ちない", "except Exception as e:" in src, True)


def test_boost_finds_link_sources():
    """押し上げるページに、リンク元を見つけられること。

    検索語を空白で割るだけだと、空白の無い日本語の語が1つの塊のまま残る。
    実測で「aiかんたん集客」は分割されず、その文字列を3回以上含む記事が
    0本だったため、11.3位・表示88回のページに1本もリンクを送れなかった。
    「0本足しました」と表示されるだけで、原因は何も示されない。

    主力語の判定にも漏れがあった。サイト名が「AI集客ラボ」で主力商材が
    AIを使った集客そのものなのに、判定に ai集客 が入っていなかった。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import priority_boost as P
    import re as _re
    print(chr(10) + "■ 押し上げのリンク元探し")

    # サイトの主力語が判定に入っていること
    pat = _re.compile(P.MAIN_PATTERN["ai-lab"], _re.I)
    for kw in ("aiかんたん集客", "ai集客", "aio対策", "meo 病院", "seo 対策"):
        check(f"主力語として拾う: {kw}", bool(pat.search(kw)), True)
    check("無関係な語は拾わない", bool(pat.search("確定申告 やり方")), False)

    # 空白の無い語からも、リンク元を探せる語を作れること
    nl2 = chr(10)
    texts = {"x": nl2.join(["---", "title: AI集客とは？かんたんに始める方法",
                            "category: ai-marketing", "---", "本文"])}
    w = P.topic_words("x", [("aiかんたん集客", 11.3, 88)], texts)
    check("空白の無い語でも語を作れる", len(w) >= 2, True)
    check("タイトルからも語を取る", any("集客" in x for x in w), True)
    check("1文字の語は使わない", all(len(x) >= 2 for x in w), True)
    check("語を取りすぎない（無関係な記事に当たる）", len(w) <= 8, True)


def test_rank_data_is_verified():
    """順位改善が使う数字が、別の取り方と一致すること。

    GSCは見る次元で数字が変わる。query次元は検索数の少ない語を返さないため、
    実測で page次元50クリックが query次元では0と出た。この食い違いに
    気づかず「クリックが出ていない」と判断し、誤った結論を重ねた。

    同じページが「末尾スラッシュあり/なし」で2行に分かれることもあり、
    上書きすると表示が消える（実測で2ページ・37表示が消えた）。

    取り方は rank_up.fetch() の1か所に固定し、page次元を正とする。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    import rank_up as R
    print(chr(10) + "■ 順位データの取り方")

    src = (ROOT / "scripts" / "rank_up.py").read_text(encoding="utf-8")
    check("取得の入口が1つ", src.count("def fetch(") == 1, True)
    check("page次元を使う", '["page"], 25000' in src, True)
    check("URLを鍵にする（slugで潰さない）", 'pages[url] =' in src, True)
    check("末尾スラッシュの重複を合算する", 'd["urls"] += 1' in src, True)
    check("表示の少ないページを対象にしない", R.MIN_IMP >= 20, True)
    # 1つの帯だけ見ても全体は動かない。実測で11〜20位は全323ページの22%だった
    names = [b[2] for b in R.BANDS]
    check("全順位帯を見る", len(R.BANDS), 5)
    for n in ("1〜3位", "4〜10位", "11〜20位", "21〜50位", "51位以下"):
        check(f"帯がある: {n}", n in names, True)
    # 帯ごとに違う手を当てる（1ページ目にいるページへリンクを足しても効かない）
    check("帯ごとにやることが違う", len({b[3] for b in R.BANDS}), 5)
    check("順位相応のクリック率を持っている", R.expected_ctr(5.0) > R.expected_ctr(15.0), True)
    check("効く順に渡す口がある", hasattr(R, "human_items"), True)

    # 直した記録を残し、効果を測れること
    check("記録を残す", "def save_log" in src, True)
    check("前後を比べる口がある", "--effect" in src, True)

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("毎週走る", "rank_up.py --write" in wf, True)
    check("効果を先に記録する", wf.index("--effect") < wf.index("rank_up.py --write"), True)


def test_measurement_pitfalls_are_documented():
    """間違えやすい測り方が、手順として残っていること。

    1つのセッションで9回、誤った数字を報告した。どれも30秒の再確認で防げた。
    共通していたのは「自分の測り方が正しい」と仮定したこと。
    対象の性質ではなく、道具の癖を見ていた。

    人の注意に任せると忘れるので、手順を文書に残し、検算の道具を用意する。
    """
    print(chr(10) + "■ 測り方の落とし穴")
    md = (ROOT / "CLAUDE.md").read_text(encoding="utf-8")
    check("数字を出す前の手順がある", "数字を出す前に必ず通す手順" in md, True)
    for k in ("別の方法で同じ数字を出す", "内訳に分ける",
              "逆を探す", "前提を口に出して確かめる"):
        check(f"手順に含まれる: {k}", k in md, True)

    # 実際に踏んだ落とし穴が、名指しで残っていること
    for k in ("query次元", "末尾スラッシュ", "指名検索", "lead_capture"):
        check(f"落とし穴が明記されている: {k}", k in md, True)

    check("検算の道具がある", (ROOT / "scripts" / "data_sanity.py").is_file(), True)
    ds = (ROOT / "scripts" / "data_sanity.py").read_text(encoding="utf-8")
    for name, key in (("二重計上", "check_double_count"), ("無効な流入", "check_invalid"),
                      ("GA4とGSCの食い違い", "check_ga_vs_gsc"),
                      ("指名検索の割合", "check_brand_share")):
        check(f"検算する: {name}", key in ds, True)


def test_built_tools_actually_run():
    """作った道具が、実際に動いていること。

    scripts に96本あるうち22本が、ワークフローからも他スクリプトからも
    呼ばれていなかった。index_status.py と reindex.py はその典型で、
    未登録ページを拾う仕組みがあるのに一度も走っていなかった。
    作っただけで動いていない道具は、無いのと同じ。

    手で使う道具（営業資料・クライアント導入など）は定期実行しなくてよい。
    ここでは「毎週走るべきもの」だけを固定する。
    """
    print(chr(10) + "■ 作った道具が動いているか")
    wf = " ".join(p.read_text(encoding="utf-8", errors="replace")
                  for p in (ROOT / ".github" / "workflows").glob("*.yml"))
    # 読むだけの検査は findings.py がまとめて呼ぶ。YAMLに名前は出ないため、
    # その一覧も「毎週走るもの」として数える（間接呼び出しを追わないと、
    # まとめ役に移しただけで「動いていない」と誤検出する）
    check("まとめ役が毎週走る", "findings.py" in wf, True)
    import findings
    wf += " " + " ".join(sc for _, sc, _ in findings.CHECKS)
    for name, why in (("reindex.py", "登録されていないページを拾う"),
                      ("rank_up.py", "順位を上げる"),
                      ("anchor_audit.py", "リンクの文言に狙う語を入れる"),
                      ("aio_check.py", "AIO基盤のずれを見る"),
                      ("nap_check.py", "会社表記のずれを見る"),
                      ("data_sanity.py", "計測が壊れていないか")):
        check(f"毎週走る: {why}", name in wf, True)

    # 順序が逆だと、入れた狙う語が短縮で落ちる
    check("狙う語を入れてから詰める",
          wf.index("anchor_audit.py") < wf.index("shorten_anchors.py"), True)


def test_quality_fixes_run_before_publishing():
    """機械で直せる崩れは、公開する前に直すこと。

    記事は1日6本公開されるのに、崩れを直す工程は週1回しか走っていなかった。
    最大6日間そのまま公開される計算で、実測でも直近7日の記事に1本残っていた。
    読者が崩れたページを見てから直しても遅い。

    どの工程も地の文を変えず、検算に外れたら書き込まない作りになっている。
    公開前に置いても壊れない。
    """
    print(chr(10) + "■ 公開前の品質修正")
    for f in ("pipeline-multi.yml", "pipeline.yml"):
        wf = (ROOT / ".github" / "workflows" / f).read_text(encoding="utf-8")
        for name, label in (("fix_block_breaks.py", "段落の崩れ"),
                            ("split_paragraphs.py", "長すぎる段落"),
                            ("split_sentences.py", "長すぎる1文"),
                            ("shorten_anchors.py", "長すぎるアンカー"),
                            ("anchor_audit.py", "アンカーに狙う語")):
            check(f"{f} で公開前に直す: {label}", name in wf, True)
        # ビルド（＝公開物の生成）より前でなければ意味がない
        check(f"{f}: ビルドより前に直す",
              wf.index("fix_block_breaks.py") < wf.rindex("python scripts/build.py"), True)
        # 狙う語を入れてから詰める（逆だと語が落ちる）
        check(f"{f}: 狙う語を入れてから詰める",
              wf.index("anchor_audit.py") < wf.index("shorten_anchors.py"), True)


def test_findings_judges_by_marker_not_exit_code():
    """「見つかった」と「壊れた」を、終了コードで混同しないこと。

    auto_review は検出だけで終了コード1を返していた。CIのログでは
    本当に落ちたときと区別がつかず、正常な検出を不具合と読み違えた。
    判定は印（*_OK=）で行い、終了コードは検査が動いたかだけに使う。
    """
    import findings
    print("\n■ 検出と故障の見分け")
    cases = [
        ("見つかった（終了コードは0）", "NAP_OK=no", 0, "要対応"),
        ("見つからなかった", "NAP_OK=yes", 0, "問題なし"),
        ("印が故障より優先される", "AIO_OK=no", 1, "要対応"),
        ("印が無ければ終了コードを見る", "なにも出ていません", 1, "要対応"),
        ("印が無く終了コードも0", "なにも出ていません", 0, "問題なし"),
        ("検査が動かない", "（起動できません）", 127, "動かせず"),
    ]
    for name, text, rc, want in cases:
        check(name, findings.judge(text, rc), want)

    # 検査対象のスクリプトが実在すること（名前を変えたら気づける）
    for label, script, _ in findings.CHECKS:
        # 引数つきの指定（"measure.py --log"）も許す。実在はファイル名で見る
        check("検査が実在: " + label,
              (ROOT / "scripts" / script.split()[0]).exists(), True)


def test_detection_scripts_do_not_fail_the_run():
    """検出だけのスクリプトが、終了コード1で止めないこと。

    例外は nap_check と deploy_check（人が手を動かすまで消えない問題）。
    """
    print("\n■ 検出＝0 の規約")
    src = (ROOT / "scripts" / "auto_review.py").read_text(encoding="utf-8")
    check("auto_review が印を出す", "REVIEW_OK=" in src, True)
    check("auto_review の検出が return 1 でない",
          'print("REVIEW_OK=no")\n        return 0' in src, True)
    src = (ROOT / "scripts" / "data_sanity.py").read_text(encoding="utf-8")
    check("data_sanity が印を出す", "SANITY_OK=" in src, True)


def test_notify_does_not_send_junk():
    """中身の無い通知を、本物の宛先に送らないこと。

    .env に RESEND_API_KEY と LEAD_TO_EMAIL があるため、手元での動作確認が
    info.ai へ本物のメールを出していた。実際に届いたのは2通:
      - 「（メッセージなし）」… 引数なしで呼んだもの
      - 「$icon …」        … 呼び出し側の引用符が壊れ、変数が展開されなかったもの
    通知は人の注意を使う。中身の無いものを送ると、次から読まれなくなる。
    """
    import notify_slack as N
    print("\n■ 送ってはいけない通知")
    check("空は送らない", bool(N.unsendable("")), True)
    check("空白だけも送らない", bool(N.unsendable("   \n  ")), True)
    check("未展開の変数は送らない", bool(N.unsendable("$icon 週次最適化: success")), True)
    check("Actionsの式も送らない", bool(N.unsendable("${{ job.status }}")), True)
    check("正常な本文は送れる", N.unsendable("✅ 記事パイプライン: success"), "")
    # 金額の $ を誤って弾かないこと（先頭でなければ通す）
    check("本文中のドルは弾かない", N.unsendable("月額 $99 のプランに変更"), "")

    src = (ROOT / "scripts" / "notify_slack.py").read_text(encoding="utf-8")
    check("手元では送らない歯止めがある", "GITHUB_ACTIONS" in src, True)


def test_routine_success_is_not_emailed():
    """うまくいっただけの回を、メールで送らないこと。

    届けるのは4つだけ: 問い合わせ / 異常のアラート / 月次レポート /
    変えたほうが良い点。1日2〜3通の成功報告が届くと、本当に見てほしい回が
    埋もれる。定時の工程は --routine を付けて呼び、知らせることがある回だけ送る。
    """
    import notify_slack as N
    print(chr(10) + "■ 定時の報告を送らない")
    ok = "✅ 記事パイプライン [ai-lab]: success" + chr(10) + "publish(ai-lab)"
    check("成功だけの回は送らない", N.worth_sending(ok, True), False)
    check("--routine が無ければ送る", N.worth_sending(ok, False), True)
    check("失敗は送る",
          N.worth_sending("🚨 記事パイプライン: failure", True), True)
    check("要対応があれば送る",
          N.worth_sending("✅ 週次最適化: success" + chr(10) + "要対応: 会社表記のゆれ", True), True)
    check("検査がすべてOKなら送らない",
          N.worth_sending("🔧 週次最適化: success" + chr(10) + "検査4件すべて問題なし", True), False)
    check("公開が止まっていれば送る",
          N.worth_sending("✅ パイプライン: success" + chr(10) + "BLOCKED記事あり", True), True)

    # 定時の工程が --routine を付けて呼んでいること（付け忘れると元に戻る）
    wfs = {p.name: p.read_text(encoding="utf-8", errors="replace")
           for p in (ROOT / ".github" / "workflows").glob("*.yml")}
    for name in ("pipeline-multi.yml", "pipeline.yml", "weekly-optimize.yml",
                 "digest.yml", "monthly-report.yml"):
        src = wfs.get(name, "")
        if "notify_slack.py" not in src:
            continue
        check(f"{name} が定時の報告を絞っている",
              "notify_slack.py --routine" in src, True)

    # 問い合わせは Apps Script が直接送るため、この絞り込みを通らない
    gas = " ".join(p.read_text(encoding="utf-8", errors="replace")
                   for p in (ROOT / "automation" / "gas").glob("*.gs"))
    check("問い合わせは別経路で届く", "MailApp.sendEmail" in gas, True)


def test_every_site_gets_articles():
    """サイトを増やしたとき、どのサイトも0本にならないこと。

    日次は6つの時刻枠に 0/1/2 を直接書いていたため、4社目以降は一度も
    選ばれず記事が1本も書かれなかった。サイトを足しても静かに無視される。
    枠の番号をサイト数で割った余りにすると、全社へ行き渡る。
    ただし余りだけだと1社のとき同じサイトに6回回るため、1社2本の上限もかける。
    """
    print(chr(10) + "■ サイト数ごとの記事の割り当て")
    SLOTS = 20                      # cron の数（2本 × 10社）

    def assign(n):
        out = {}
        for slot in range(SLOTS):
            if slot >= n * 2:       # 1サイト2本まで
                continue
            i = slot % n
            out[i] = out.get(i, 0) + 1
        return out

    # 10社までは、どの社も1日2本ずつ受け取れること
    for n in (1, 2, 3, 5, 8, 10):
        got = assign(n)
        zero = [i for i in range(n) if got.get(i, 0) == 0]
        print("   %2d社 → 合計%2d本 / 1社あたり %d〜%d本"
              % (n, sum(got.values()), min(got.values()), max(got.values())))
        check("%d社で全サイトに記事が回る" % n, zero, [])
        check("%d社は全社2本ずつ" % n, (min(got.values()), max(got.values())), (2, 2))
        check("%d社の合計本数" % n, sum(got.values()), n * 2)

    # 上限を超えても0本の社は出さない（本数が減るだけ）。そのうえで知らせる
    over = assign(11)
    print("   11社 → 合計%d本 / 1社あたり %d〜%d本（上限超え）"
          % (sum(over.values()), min(over.values()), max(over.values())))
    check("11社でも0本の社は出ない", [i for i in range(11) if over.get(i, 0) == 0], [])

    wf = (ROOT / ".github" / "workflows" / "pipeline-multi.yml").read_text(
        encoding="utf-8", errors="replace")
    # 直書きに戻ると4社目が静かに消えるため、余りで選んでいることを固定する
    check("余りでサイトを選んでいる", "slot % N" in wf, True)
    check("1サイト2本の上限がある", "N * 2" in wf, True)
    check("枠が20ある（10社ぶん）", wf.count("* * *\"") >= 20, True)
    # 10社を超えたら、実行のたびに知らせる（黙って本数が減ると気づけない）
    check("10社超をエラーで知らせる", '[ "$N" -gt 10 ]' in wf, True)
    check("別リポジトリへ分けるよう示す", "別リポジトリに分けて" in wf, True)


def test_intake_sheets_are_not_published():
    """ヒアリングシートがリポジトリに入らないこと。

    シートには会社名・住所・電話・担当者のメールが入る。このリポジトリは
    public なので、置いたまま commit すると誰でも読める。取り込みは手元で
    実行し、生成された sites/*.json だけを commit する。
    """
    import subprocess
    print(chr(10) + "■ ヒアリングシートの扱い")
    for name in ("intake/A社.xlsx", "intake/done/A社.xlsx", "intake/todo/A社.xlsx"):
        r = subprocess.run(["git", "check-ignore", "-q", name],
                           cwd=str(ROOT), capture_output=True)
        check(f"{name} は公開されない", r.returncode == 0, True)

    # 置き場の説明だけは残す（フォルダが消えると使い方が分からなくなる）
    check("置き場の説明がある", (ROOT / "intake" / "README.md").is_file(), True)

    src = (ROOT / "scripts" / "intake_watch.py").read_text(encoding="utf-8")
    # 不備のあるシートを通すと、どのサイトでも書ける記事が量産される
    check("不備があれば登録しない", "intake/todo/ へ移しました" in src, True)
    # 上限を超えて受け入れると、記事の枠が足りず全社の本数が減る
    check("10社の上限で止める", "MAX_SITES = 10" in src, True)


def test_kw_plan_keeps_only_buyers():
    """計画の一新で、関連しない語・見込み客でない語が入らないこと。

    ラッコのLSIは「経理代行 求人」「記帳代行 儲かる」を重要度highで返す。
    検索している人は顧客ではなく求職者・副業者で、記事にしても問い合わせに
    ならない。担当領域の条件（kw_discover と同じ）に加えて、ここで落とす。
    一新は非破壊で、台帳の「未着手」だけを対象外にし、公開済み・執筆中は触らない。
    """
    import kw_plan
    import kw_intent
    print(chr(10) + "■ 計画の一新で入れてよい語")
    S = {"own_terms": ("経理", "記帳", "請求書"), "domain_terms": ("経理", "記帳", "費用", "相場"),
         "ng_terms": ("aio",), "industries": ["経理"], "intents": ["経理代行 費用"]}
    corpus, arts, owned, picked = [], [], [], []

    def ok(kw, vol=100):
        return kw_plan.relevant({"kw": kw, "vol": vol, "imp": 0}, S, corpus, arts, owned, picked)

    check("見込み客の語は通す", ok("経理代行 費用 相場"), "")
    check("求人は落とす", ok("経理代行 求人"), "見込み客でない")
    check("副業は落とす", ok("記帳代行 副業 在宅"), "見込み客でない")
    check("儲かるは落とす", ok("記帳代行 儲かる"), "見込み客でない")
    # 他社製品の指名検索。「請求書freee」を書いても経理BPOの相談には来ない
    check("他社製品名は落とす（freee）", ok("請求書freee"), "他社名")
    check("他社製品名は落とす（マネーフォワード）", ok("請求書マネーフォワード 使い方"), "他社名")
    check("他社サービス名は落とす（ドットコム）", ok("記帳代行ドットコム"), "他社名")
    check("他サイトの領域語は落とす", ok("経理 aio対策"), "除外語")
    check("領域語の無い語は落とす", ok("ホームページ 作成 費用"), "領域語なし")
    # 部分一致の罠。「日記帳」は「記帳」ではなく、「給付金請求書」は「請求書」ではない
    check("別の語の一部は落とす（日記帳）", ok("おすすめ 日記帳"), "領域語なし")
    check("別の語の一部は落とす（給付金請求書）", ok("年金生活者支援給付金請求書"), "領域語なし")
    check("番号を調べるだけの語は落とす", ok("適格請求書発行事業者登録番号"), "領域語なし")
    check("先頭一致は通す（請求書の封筒の書き方）", ok("請求書の封筒の書き方"), "")
    check("設定にある複合語は通す（経理代行）", ok("経理代行 相場"), "")
    subs = {"own_terms": ("補助金", "申請"), "domain_terms": ("補助金",), "ng_terms": (),
            "industries": ["運送業"], "intents": ["小規模事業者持続化補助金", "申請 代行"]}
    check("設定にある複合語は通す（〜補助金）",
          kw_plan.relevant({"kw": "小規模事業者持続化補助金 運送業", "vol": 50, "imp": 0}, subs, corpus, arts, owned, picked), "")
    # 業種はサブジェクトであって領域ではない。業種名だけでは通さない
    lab = {"own_terms": ("aio", "meo", "集客"), "domain_terms": ("aio",), "ng_terms": (),
           "industries": ["リフォーム", "飲食店"], "intents": ["AIO対策 失敗"]}
    def lab_ok(kw):
        return kw_plan.relevant({"kw": kw, "vol": 5000, "imp": 0}, lab, corpus, arts, owned, picked)
    check("業種名だけの語は落とす（リフォーム キッチン 費用）", lab_ok("リフォーム キッチン 費用"), "領域語なし")
    check("業種名だけの語は落とす（飲食店 近くの）", lab_ok("飲食店 近くの"), "領域語なし")
    check("業種×領域語は通す（リフォーム meo 対策）", lab_ok("リフォーム meo 対策"), "")
    lab2 = dict(lab, own_terms=("aio", "meo", "口コミ"))
    def lab2_ok(kw):
        return kw_plan.relevant({"kw": kw, "vol": 5000, "imp": 0}, lab2, corpus, arts, owned, picked)
    check("患者の評判検索は落とす（クリニック 口コミ）", lab2_ok("レジーナ クリニック 口コミ"), "領域語なし")
    check("事業者側の口コミは通す（口コミ 返信）", lab2_ok("クリニック 口コミ 返信 例文"), "")
    check("領域語を含む複合意図は通す（aio対策 失敗）", lab_ok("工務店 aio対策 失敗"), "")
    # 業種に掛ける領域語は設定で明示できる。無ければ owns の先頭
    check("core があればそれを使う",
          kw_plan.core_terms({"cfg": {"kw_seeds": {"core": ["集客", "MEO"]}}, "own_terms": ("aio", "llmo")}),
          ["集客", "meo"])
    check("core が無ければ owns の先頭",
          kw_plan.core_terms({"cfg": {}, "own_terms": ("aio", "llmo", "seo", "meo", "x")}),
          ["aio", "llmo", "seo", "meo"])
    check("買い手の意図（申請 代行）が上限内に残る", "申請 代行" in kw_plan.intents_for(
        {"intents": ["申請 代行"] + ["意図%d" % i for i in range(30)]}), True)
    check("検索数も表示も無い語は落とす", ok("経理 記帳 手順", vol=None), "検索数が少ない")
    check("検索数が無くても表示があれば通す",
          kw_plan.relevant({"kw": "経理 記帳 手順", "vol": None, "imp": 30}, S, corpus, arts, owned, picked), "")

    # 語順・助詞が違うだけの語は同じ検索。別枠で採ると採用枠を食い合う
    check("語順違いは同じ語", kw_plan.same("請求書 書き方 封筒", "請求書 封筒 書き方"), True)
    check("助詞の有無も同じ語", kw_plan.same("請求書の封筒の書き方", "請求書 封筒 書き方"), True)
    check("空白なし・助詞ありも同じ語", kw_plan.same("請求書封筒書き方", "請求書の書き方 封筒"), True)
    check("違う語は別", kw_plan.same("経理代行 費用", "記帳代行 副業"), False)
    # 優先業種（月1件の成約で費用が回収できる業種）は点が上がり、枠も厚い
    base = kw_plan.score({"kw": "美容室 meo 対策", "vol": 100, "kd": 30, "subject": "美容室"}, ["クリニック"])
    top = kw_plan.score({"kw": "クリニック meo 対策", "vol": 100, "kd": 30, "subject": "クリニック"}, ["クリニック"])
    check("優先業種の語に加点される", round(top - base, 2), kw_plan.PRIORITY_BONUS)
    check("優先業種の枠は他業種より厚い", kw_plan.PER_PRIORITY > kw_plan.PER_SUBJECT, True)
    check("ai-lab の優先業種が設定にある",
          set(kw_plan.priority_subjects({"cfg": __import__("json").load(open(ROOT / "sites" / "ai-lab.json", encoding="utf-8"))}))
          >= {"クリニック", "不動産", "リフォーム"}, True)
    check("補助金の優先が個人事業主と中小企業",
          set(kw_plan.priority_subjects({"cfg": __import__("json").load(open(ROOT / "sites" / "subsidy.json", encoding="utf-8"))})),
          {"個人事業主", "中小企業"})

    # 買い手の語は、検索数が少なくても上に来ること
    buyer = kw_plan.score({"kw": "経理代行 費用 相場", "vol": 210, "kd": 31})
    diy = kw_plan.score({"kw": "請求書 封筒 書き方", "vol": 5400, "kd": 33})
    check("外注を考える語が、自分でやる語より上", buyer > diy, True)

    # ルール: 課金の前に必ず見積もり、予算内でだけ取得する（手で走らせても同じ）
    src_plan = (ROOT / "scripts" / "kw_plan.py").read_text(encoding="utf-8")
    check("課金の前に見積もる", "if not budget_ok(S) and not DRY" in src_plan, True)
    import rakko as _rk
    S_est = {"cfg": {"kw_seeds": {"core": ["集客"]}}, "own_terms": ("aio",),
             "industries": ["クリニック", "歯科医院"], "intents": []}
    check("見積もりは問い合わせ数×1.5＋一括15", kw_plan.estimate(S_est), 2 * 1.5 + 15)
    # 一括調査でSEO難易度を取ると1語0.75で、500語なら375になる（今日の主因）。取らない
    check("一括調査で難易度を取らない", '"seoDifficulty": False' in src_plan and '"seoDifficulty": True' not in src_plan, True)
    saved_ms, saved_mb = _rk.month_spent, _rk.MONTHLY_BUDGET
    try:
        _rk.month_spent, _rk.MONTHLY_BUDGET = (lambda month=None: 995.0), 1000
        check("月の目安を超えるなら取得しない", kw_plan.budget_ok(S_est), False)
        _rk.month_spent = lambda month=None: 0.0
        check("予算内なら取得する", kw_plan.budget_ok(S_est), True)
    finally:
        _rk.month_spent, _rk.MONTHLY_BUDGET = saved_ms, saved_mb
    wf_m = (ROOT / ".github" / "workflows" / "monthly-report.yml").read_text(encoding="utf-8", errors="replace")
    check("月次は dry-run を先に記録してから本番", wf_m.index("--dry-run") < wf_m.index("--if-needed --replace"), True)
    src_disc = (ROOT / "scripts" / "kw_discover.py").read_text(encoding="utf-8")
    check("週次補充も月の目安を見る", "rakko.month_spent() > rakko.MONTHLY_BUDGET" in src_disc, True)

    # 一括調査に送るのは価値の高い語だけ、1回ぶんまで（1万件を20回送って300クレジット払った）
    many = [{"kw": "補助金 その%d" % i} for i in range(700)] \
         + [{"kw": "補助金 申請 代行 費用"}, {"kw": "補助金 不採択 理由"}, {"kw": "補助金 とは", "imp": 30}]
    sel = kw_plan.worth_lookup(many)
    check("価値の無い語は一括調査に送らない", len(sel) <= kw_plan.LOOKUP_MAX and all(
        c.get("imp") or kw_plan.BUYER.search(c["kw"]) or kw_intent.score(c["kw"])[0] >= 2 for c in sel), True)
    check("買い手の語が先頭に来る", sel[0]["kw"], "補助金 申請 代行 費用")

    # 検索数を取る前に、安い条件で落とす（1万件を一括登録して500エラーになった）
    check("事前選別: 見込み客でない語", kw_plan.cheap_reject({"kw": "経理代行 求人"}, S), "見込み客でない")
    check("事前選別: 通る語", kw_plan.cheap_reject({"kw": "経理代行 費用 相場"}, S), "")
    check("一括登録は500件ずつ", [len(c) for c in kw_plan.chunks(list(range(1201)), kw_plan.BULK)], [500, 500, 201])

    # ラッコの 502/503/504 は数分続く。1回で諦めると計画が半分で止まる
    import io as _io, urllib.error, rakko
    calls = {"n": 0}
    def flaky(req, timeout=0):
        calls["n"] += 1
        if calls["n"] < 3:
            raise urllib.error.HTTPError("u", 502, "Bad Gateway", {}, _io.BytesIO(b"<html>502</html>"))
        return _io.BytesIO(b'{"result":true,"data":{"ok":1}}')
    saved = (rakko.urllib.request.urlopen, rakko.api_key, rakko.RETRY_WAIT, rakko._OUT_OF_CREDIT)
    rakko.urllib.request.urlopen, rakko.api_key, rakko.RETRY_WAIT = flaky, (lambda: "k"), (0, 0, 0)
    # 本物のキャッシュに触らない（前回のテストの応答が残っていると呼び出し回数が0になる）
    import tempfile, pathlib
    saved_cache = rakko.CACHE_DIR
    rakko.CACHE_DIR = pathlib.Path(tempfile.mkdtemp())
    try:
        r = rakko.call("/v1/x", {"a": 1})
        check("5xx は待ってやり直す", (r or {}).get("data"), {"ok": 1})
        check("やり直しの回数", calls["n"], 3)
        # 自動課金だと尽きずに請求が伸びる。上限に達したら以降は呼ばない
        def paid(req, timeout=0):
            return _io.BytesIO(b'{"result":true,"meta":{"consumedCredit":1.5},"data":{}}')
        rakko.urllib.request.urlopen = paid
        rakko.CONSUMED, rakko.BUDGET = 0.0, 3.0
        # 問い合わせを変える（同じ内容はキャッシュから返って課金されないため）
        got = [rakko.call("/v1/x", {"i": i}) is not None for i in range(4)]
        check("上限まで呼べる（1.5×2=3.0）", got[:2], [True, True])
        check("上限に達したら止まる", got[2:], [False, False])
        check("消費を数えている", rakko.spent(), 3.0)
        # 同じ問い合わせは期限内なら課金なしで返す。条件の調整で取り直して約1,000クレジット無駄にした
        rakko.CACHE_DIR = pathlib.Path(tempfile.mkdtemp())
        rakko.CONSUMED, rakko.BUDGET = 0.0, None
        hits = {"n": 0}
        def counted(req, timeout=0):
            hits["n"] += 1
            return _io.BytesIO(b'{"result":true,"meta":{"consumedCredit":1.5},"data":{"items":[1]}}')
        rakko.urllib.request.urlopen = counted
        a = rakko.call("/v1/suggest-keywords", {"keyword": "x"})
        b = rakko.call("/v1/suggest-keywords", {"keyword": "x"})
        check("同じ問い合わせは1回しか課金しない", (hits["n"], rakko.spent()), (1, 1.5))
        check("2回目も同じ中身が返る", a == b, True)
        rakko.call("/v1/search-volume", {"keywords": ["x"]}); rakko.call("/v1/search-volume", {"keywords": ["x"]})
        check("一括調査の登録は残さない（毎回別物）", hits["n"], 3)
    finally:
        rakko.urllib.request.urlopen, rakko.api_key, rakko.RETRY_WAIT, rakko._OUT_OF_CREDIT = saved
        rakko.CONSUMED, rakko.BUDGET = 0.0, None
        rakko.CACHE_DIR = saved_cache

    # 管制塔の一時的な404も待ってやり直す。取り下げ直後の追加で当たると未着手が0件で残る
    import hub_client
    hits = {"n": 0}
    def flaky_hub(req, timeout=0):
        hits["n"] += 1
        if hits["n"] < 3:
            raise urllib.error.HTTPError("u", 404, "Not Found", {}, _io.BytesIO(b""))
        return _io.BytesIO(b'{"ok":true}')
    saved2 = (hub_client.urllib.request.urlopen, hub_client.RETRY_WAIT)
    hub_client.urllib.request.urlopen, hub_client.RETRY_WAIT = flaky_hub, (0, 0, 0)
    try:
        with hub_client._open(hub_client.urllib.request.Request("https://example.invalid/")) as r:
            check("管制塔の404は待ってやり直す", r.read(), b'{"ok":true}')
        check("やり直しの回数（管制塔）", hits["n"], 3)
    finally:
        hub_client.urllib.request.urlopen, hub_client.RETRY_WAIT = saved2

    # 計画ファイルの表を読み戻せること（レポートが検索数・難易度を引く経路）
    tmp = ROOT / "docs" / "kw-plan-_gate_.md"
    tmp.write_text("| 優先 | キーワード | 月間 | 難易度 | 開く理由 | 12か月 | 表示 | 出どころ |" + chr(10)
                   + "|:--|:--|--:|--:|--:|--:|--:|:--|" + chr(10)
                   + "| A | 記帳代行 相場 | 480 | 29 | +2 | -33% | — | rakko |" + chr(10)
                   + "| C | 経理 やり方 | — | — | +3 | — | 12 | suggest |" + chr(10),
                   encoding="utf-8")
    try:
        m = kw_plan.plan_metrics("_gate_")
        check("表から月間・難易度・優先を読める", m.get("記帳代行相場"), (480, 29, "A"))
        check("欠けている値は None", m.get("経理やり方"), (None, None, "C"))
    finally:
        tmp.unlink(missing_ok=True)

    # 意図の掛け合わせは上限まで（補助金サイトは26×30で時間切れになった）
    many = {"intents": ["意図%d" % i for i in range(30)]}
    check("意図は上限まで", len(kw_plan.intents_for(many)) <= kw_plan.MAX_INTENTS, True)

    # 新計画にもある語は取り下げない。配備中の管制塔は取り下げ済みの語の再追加を弾き、
    # その語が計画から消える（実際に17本消えた）
    r, k = kw_plan.split_retire(["経理 アウトソース 費用", "古い語 A", "請求書 封筒 書き方"],
                                ["経理 アウトソース 費用", "請求書の書き方 封筒"])
    check("新計画にある語は残す", sorted(k), ["経理 アウトソース 費用", "請求書 封筒 書き方"])
    check("新計画に無い語だけ取り下げる", r, ["古い語 A"])
    # kw_plan が以前に積んだ語は、今回の計画に無くても取り下げない（再実行で在庫が減る）
    r2, k2 = kw_plan.split_retire(["前回の計画の語", "古い語 A"], ["別の語"], own_norms={kw_plan.norm("前回の計画の語")})
    check("以前に積んだ語は残す", (r2, k2), (["古い語 A"], ["前回の計画の語"]))
    check("積んだ語の記録がある", (ROOT / "data" / "kw_plan_added.json").is_file(), True)

    # 一新は「未着手」だけを対象外にする。公開済み・執筆中を落とすと生きている記事が消える
    src = (ROOT / "scripts" / "kw_plan.py").read_text(encoding="utf-8")
    check("未着手だけを取り下げる", 'get("status") == "未着手"' in src, True)
    check("force で公開済みを落とさない", "force=True" not in src, True)


def test_reports_carry_diagnosis_and_next_actions():
    """レポートに「現在のサイト診断と、次にやること」が入っていること。

    数字の推移だけでは、いまどこが弱く次に何をするかが読み手に委ねられていた。
    順位帯・台帳・指名検索・導線・検査の要対応から機械的に組み立てる。
    週次は検査（findings）の後に作らないと、要対応が載らない。
    """
    import site_diagnosis
    print(chr(10) + "■ レポートの診断と次にやること")
    d = {"site": "x", "name": "テスト", "domain": "example.invalid", "days": 28,
         "bands": {"4〜10位": {"pages": 3, "imp": 300, "clicks": 3, "gap": 9.0},
                   "11〜20位": {"pages": 2, "imp": 100, "clicks": 1, "gap": 1.0}},
         "funnel": [("記事を見た", 200), ("CTAを押した", 2), ("フォームを開いた", 1), ("送信した", 1)],
         "brand": {"prev": 10, "cur": 30, "growth": "+200%"},
         "stock": {"todo": 12, "doing": 1, "done": 50, "A": 4},
         "fixes": [{"slug": "a-b", "why": "9位・表示101・クリック1｜1ページ目にいるのにクリックが取れていない"}],
         "findings": ["要対応: 会社表記のゆれ（NAP）"]}
    acts = site_diagnosis.next_actions(d)
    check("1ページ目の取りこぼしを自動の手当てに出す", any("タイトル・説明文" in a for _, a in acts), True)
    check("在庫が薄ければ自動の手当てに出す", any("在庫が 12 本" in a for _, a in acts), True)
    check("CTA押下率が低ければ人の手当てに出す", any("CTAの押下率" in a for w, a in acts if w == "人"), True)
    check("検査の要対応を人の手当てに出す", any("会社表記のゆれ" in a for w, a in acts if w == "人"), True)
    h = site_diagnosis.html_block(d)
    check("HTMLに順位帯の表がある", "順位相応なら増える" in h, True)
    check("HTMLに次にやることの表がある", "次にやること（効く順）" in h, True)
    # 材料が取れなくてもHTMLは出る
    h2 = site_diagnosis.html_block({"site": "x", "name": "テスト", "domain": "e", "days": 28,
                                     "bands": None, "funnel": None, "brand": None, "stock": None,
                                     "fixes": [], "findings": []})
    check("材料が無くても落ちない", "取得できません" in h2, True)
    wk = (ROOT / "scripts" / "weekly_report.py").read_text(encoding="utf-8")
    mo = (ROOT / "scripts" / "monthly_report.py").read_text(encoding="utf-8")
    check("週次レポートが診断を載せる", "site_diagnosis.html_block(" in wk or "site_diagnosis.html(" in wk, True)
    check("月次レポートが診断を載せる", "site_diagnosis.html(" in mo, True)
    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8", errors="replace")
    check("週次は検査の後にレポートを作る",
          wf.index("name: 見つかったものを集めて知らせる") < wf.index("name: 週次レポートの作成"), True)


def test_rewrites_reach_the_sheet():
    """週次で直した内容が、管制塔の「リライトログ」に届くこと。

    hub_client.rewrite_log は定義だけあって呼ぶ側が無く、リライトログは4行の
    ままだった。手元の台帳（auto_fix.jsonl / rank_up.json）にしか残らず、
    シートを見る人には直した記録が見えなかった。
    """
    print(chr(10) + "■ 直した記録がシートに届く")
    for name in ("auto_rewrite.py", "rank_up.py"):
        src = (ROOT / "scripts" / name).read_text(encoding="utf-8")
        check(f"{name} がリライトログに書く", "hub_client.rewrite_log(" in src, True)


def test_hub_has_one_kpi_writer():
    """管制塔のKPIは、書き手が1つであること。

    GASの毎朝6時の updateKpi と、CIの daily_kpi.py が両方書いていた。GAS側は
    サイト一覧のGA4/GSCが空のまま「全部0」の行を毎日3本足し、最後に書いた側として
    ダッシュボードを0で上書きしていた（合計が全て0で表示されていた）。
    """
    print(chr(10) + "■ 管制塔の書き手")
    kpi = (ROOT / "automation" / "gas" / "kpi.gs").read_text(encoding="utf-8")
    hub = (ROOT / "automation" / "gas" / "hub.gs").read_text(encoding="utf-8")
    check("GASは updateKpi のトリガーを作らない", "newTrigger('updateKpi')" not in kpi, True)
    check("計測設定の無いサイトは書かない", "if (!cfg.ga4 && !cfg.gsc) return;" in kpi, True)
    check("全部0でダッシュボードを上書きしない", "allZero && hadValue" in hub, True)
    check("エラーログは同じ未対応を積まない", "same_open" in hub, True)
    check("0行を掃除する保守タスクがある", "case 'clean_kpi'" in hub, True)
    dk = (ROOT / "scripts" / "daily_kpi.py").read_text(encoding="utf-8")
    check("AIO計測に推定を添える", "aio_est" in dk and "aio_est" in hub, True)
    # ダッシュボードは項目名で上書きする。全行を消すと別の書き手の項目が消える
    dash = (ROOT / "automation" / "gas" / "dashboard.gs").read_text(encoding="utf-8")
    body = hub.split("function writeDashboard_")[1].split(chr(10) + "function ")[0]
    body2 = dash.split("function refreshDashboard")[1].split(chr(10) + "function ")[0]
    check("ダッシュボードの書き手は行を消さない", "deleteRows" not in body2 and "deleteRows" not in body, True)
    check("両方の書き手が upsert を使う", dash.count("upsertDashboard_(") >= 1 and hub.count("upsertDashboard_(") >= 2, True)
    # CTRは数値で書く（'2.4%' は 0.024 に解釈され「0.02%」と出ていた）
    check("CTRを文字列で書かない", "(r.ctr || 0) + '%'" not in hub, True)


def test_five_hub_features_are_wired():
    """管制塔まわりの5機能が、書く側と受ける側の両方につながっていること。

    片側だけだと黙って何も起きない（rewrite_log がそうだった）。
    """
    print(chr(10) + "■ 管制塔の5機能の配線")
    c = (ROOT / "automation" / "gas" / "contact.hub.gs").read_text(encoding="utf-8")
    h = (ROOT / "automation" / "gas" / "hub.gs").read_text(encoding="utf-8")
    d = (ROOT / "automation" / "gas" / "dashboard.gs").read_text(encoding="utf-8")
    hc = (ROOT / "scripts" / "hub_client.py").read_text(encoding="utf-8")
    ru = (ROOT / "scripts" / "rank_up.py").read_text(encoding="utf-8")
    dk = (ROOT / "scripts" / "daily_kpi.py").read_text(encoding="utf-8")
    wf = (ROOT / ".github" / "workflows" / "pipeline.yml").read_text(encoding="utf-8", errors="replace")
    check("温度に流入経路を使う", "LEAD_HOT_PATHS" in c and "leadTemp_(type, d.message, d," in c, True)
    check("HOTはSlackにも送る", "SLACK_WEBHOOK_URL" in c and "temp === 'HOT'" in c, True)
    check("問い合わせを生んだページを出す", "問い合わせを生んだページ" in d, True)
    check("エラーログの同期: 受ける側", "case 'error_sync'" in h and "function errorSync_" in h, True)
    check("エラーログの同期: 送る側", "error_sync rescue" in wf and "def error_sync" in hc, True)
    check("後順位: 受ける側", "case 'rewrite_effect'" in h and "function rewriteEffect_" in h, True)
    check("後順位: 送る側", "hub_client.rewrite_effect(" in ru and "def rewrite_effect" in hc, True)
    check("AI参照の着地ページ: 受ける側", "'AI参照'" in h and "ai_pages" in h, True)
    check("AI参照の着地ページ: 送る側", 'out["ai_pages"]' in dk and "landingPage" in dk, True)


def test_merge_is_wired():
    """統合の自動化が、候補の判定・検算・配信先の取り下げまでつながっていること"""
    print(chr(10) + "■ 食い合う記事の統合")
    import auto_merge as M
    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8", errors="replace")
    check("週次が検算を先に通す", "auto_merge.py --selftest" in wf and "auto_merge.py --write" in wf, True)
    check("配信先からも外す", "retract.py --pending --push" in wf, True)
    st = {"ai-lab": {"a": {"imp": 100, "clicks": 5, "pos": 8.0}, "b": {"imp": 40, "clicks": 1, "pos": 15.0},
                     "c": {"imp": 200, "clicks": 0, "pos": 12.0}}}
    arts = {"a": {"title": "整骨院のMEO対策", "kw": "整骨院 meo", "h2": ["口コミ"]},
            "b": {"title": "整骨院のMEO対策と口コミ", "kw": "整骨院 meo 口コミ", "h2": ["口コミ"]},
            "c": {"title": "農業用倉庫の補助金", "kw": "農業用倉庫 補助金", "h2": ["倉庫"]}}
    two = [{"kw": "x"}, {"kw": "y"}]
    j = lambda s, l, rev=0: M.judge({"site": "ai-lab", "survivor": s, "loser": l, "imp": 50, "kws": list(two)}, st, rev, arts)["skip"]
    one = lambda s, l: M.judge({"site": "ai-lab", "survivor": s, "loser": l, "imp": 50, "kws": [{"kw": "x"}]}, st, 0, arts)["skip"]
    check("同じ語が1つだけなら見送る（偶然の重なり）", "同じ語が1つだけ" in one("a", "b"), True)
    check("同じ語が2つでも主題が違えば見送る", "主題が違う" in j("c", "b"), True)
    check("同じ語が2つで題名が似ていれば通す", "主題が違う" in j("a", "b"), False)
    import tempfile as _tf, pathlib as _pl
    import retract as R
    d = _pl.Path(_tf.mkdtemp())
    (d / "content").mkdir(); (d / "content" / "old.json").write_text("{}", encoding="utf-8")
    (d / "blog" / "old").mkdir(parents=True); (d / "blog" / "old" / "index.html").write_text("x", encoding="utf-8")
    (d / "article-manifest.json").write_text('{"old": "abc", "new": "def"}', encoding="utf-8")
    (d / "sitemap.xml").write_text("<urlset>\n  <url>\n    <loc>https://ex.jp/blog/old/</loc>\n  </url>\n  <url>\n    <loc>https://ex.jp/blog/new/</loc>\n  </url>\n</urlset>\n", encoding="utf-8")
    (d / "llms.txt").write_text("- [old](/blog/old/): x\n- [new](/blog/new/): y\n", encoding="utf-8")
    cfg = {"domain": "ex.jp", "content_dir": "content"}
    R.apply({"slug": "old", "from": "/blog/old/", "to": "/blog/new/", "at": "2026-09-21", "reason": "統合"}, cfg, d)
    check("配信先の記事ファイルを消す", (d / "content" / "old.json").exists() or (d / "blog" / "old").exists(), False)
    check("manifest から外す", "old" in (d / "article-manifest.json").read_text(encoding="utf-8"), False)
    check("sitemap から外し、残す方は残る", ("/blog/old/" in (d / "sitemap.xml").read_text(encoding="utf-8"), "/blog/new/" in (d / "sitemap.xml").read_text(encoding="utf-8")), (False, True))
    check("llms.txt から外す", "/blog/old/" in (d / "llms.txt").read_text(encoding="utf-8"), False)
    check("301 を書く", "/blog/old/ /blog/new/ 301" in (d / "_redirects").read_text(encoding="utf-8"), True)
    check("クリックのある記事は消さない", "クリック" in j("b", "a"), True)
    check("表示が多い側は消さない", "表示が多い" in j("a", "c"), True)
    check("逆向きの組が強ければ見送る", "逆向き" in j("a", "b", rev=99), True)
    check("GSCが引けなければ判断しない", "GSC" in M.judge({"site": "x", "survivor": "a", "loser": "b", "imp": 1, "kws": []}, {}, 0)["skip"], True)
    lo = "---\ntitle: t\n---\n## 開業届の書き方\n本文\n## まとめ\n"
    check("吸収の証拠になる語を取り出す", M.distinct_terms(lo, "これは本文です"), ["開業届"])
    # 「触ったのは1本だけか」は git ではなく直前の指紋と比べる（週次では先の工程の未コミットの直しがある）
    import auto_rewrite as ARw
    snap = ARw.snapshot()
    check("指紋が全記事ぶんある", len(snap) > 100, True)
    check("変わっていなければ空", ARw.changed_since(snap), [])
    check("既にある語は数えない", M.distinct_terms(lo, "開業届の書き方を説明します"), [])
    check("骨組みの語（まとめ）は数えない", "まとめ" in M.distinct_terms(lo, ""), False)
    check("本文の長さは装飾を除いて数える", M.plain_len("---\na: b\n---\n<p>あ い</p>\n"), 2)
    import retract as R
    check("Next.js は public/_redirects に書く", str(R.redirects_file({"images_dir": "public/images"}, ROOT / "nope")).replace("\\", "/").endswith("public/_redirects"), True)
    check("静的サイトは直下の _redirects に書く", str(R.redirects_file({}, ROOT / "nope")).replace("\\", "/").endswith("nope/_redirects"), True)


def test_jisseki_intake_never_invents_numbers():
    """実績・お客様の声は、シートに入れた数字だけを母数・期間の条件で通すこと"""
    print(chr(10) + "■ 実績・お客様の声の記入シート")
    import tempfile as _tf, pathlib as _pl
    import jisseki_intake as J
    from openpyxl import load_workbook
    d = _pl.Path(_tf.mkdtemp())
    p = J.make_sheet(d / "実績記入シート.xlsx")
    check("記入シートが4タブで作られる", load_workbook(p).sheetnames, ["書き方", "経理BPOの効果", "継続率", "お客様の声"])
    f, v, ng, warn = J.review(J.read(p))
    check("空のシートは登録しない", (f, v, bool(ng)), ([], [], True))
    wb = load_workbook(p)
    w = wb["経理BPOの効果"]; w["B2"], w["B3"], w["B4"], w["B5"] = 12, "2025-01", "2026-08", 3
    w = wb["継続率"]; w["B2"], w["B3"], w["B4"], w["B5"] = 8, 7, "2025-04", "2026-03"
    w = wb["お客様の声"]
    w.append(["可", "", "大阪市の歯科医院", "歯科", "院長", "新患が増えました", "3", "11", "月の問い合わせ件数", "2026-01〜2026-06", "ai-lab"])
    w.append(["不可", "A社", "", "不動産", "", "載せないで", "", "", "", "", ""])
    w.append(["可", "", "", "リフォーム", "", "表示名が無い", "", "", "", "", ""])
    wb.save(p)
    f, v, ng, warn = J.review(J.read(p))
    claims = {x["id"].split("-2")[0]: x["claim"] for x in f}
    check("BPOの効果は社数と期間つきの実数", "12社" in claims.get("keiri-bpo-effect", "") and "2025-01〜2026-08" in claims.get("keiri-bpo-effect", ""), True)
    check("母数10件未満の継続率は割合にしない", "%" in claims.get("keizoku", "x"), False)
    check("掲載可否が可の声だけ載る", [x["who"] for x in v], ["大阪市の歯科医院"])
    check("表示名も会社名も無い声は不備", any("表示名" in x for x in ng), True)
    check("不可の声は警告で知らせる", any("可」でない" in x for x in warn), True)
    html = J.render(v, "lp")
    check("掲載HTMLに数字の期間と許可の注記がある", "2026-01〜2026-06" in html and "許可" in html and J.START in html, True)
    placed = J.place("<x>\n<!-- 代表メッセージ -->\n<y>", html, "<!-- 代表メッセージ -->")
    check("印が無ければ目印の直前に入る", placed.index(J.START) < placed.index("<!-- 代表メッセージ -->"), True)
    again = J.place(placed, J.render(v, "lp"), "<!-- 代表メッセージ -->")
    check("2回目は置き換えで増えない", again.count(J.START), 1)
    wf = (ROOT / "scripts" / "intake_watch.py").read_text(encoding="utf-8")
    check("intake_watch が実績シートを振り分ける", "jisseki_intake" in wf, True)
    # 声は「何がどう変わったか」を先に出す。長い引用が先頭だと、並べたときに読めない
    v = [{"who": "製造業（従業員約50名）", "industry": "製造業", "person": "", "quote": "増えました",
          "number": {"metric": "問い合わせ", "before": "月3件", "after": "月8件", "period": "導入から2年"}, "sites": ["ai-lab"]}]
    h = J.render(v, "lp")
    check("数字が引用より前に出る", h.index("voice-metric") < h.index("<blockquote>"), True)
    check("業種を二重に書かない", h.count("製造業"), 1)
    check("期間の札が出る", "voice-term" in h and "導入から2年" in h, True)
    css = (ROOT / "site" / "css" / "style.css").read_text(encoding="utf-8")
    check("声の見た目はCSSで持つ", ".voice-grid" in css and ".voice-metric" in css, True)
    # 外部プロフィールは、機械（sameAs）だけでなく読者からもたどれること
    import social_footer as SF
    check("フッターの外部リンクが公開ページに入っている",
          all(SF.MARK in (ROOT / "site" / p).read_text(encoding="utf-8")
              for p in ("index.html", "about/index.html", "lp/index.html")), True)
    check("記事の雛形にも入っている", SF.MARK in (ROOT / "templates" / "article.html").read_text(encoding="utf-8"), True)
    check("2回当てても増えない", SF.insert(SF.insert("<p class=\"addr\">x</p>", SF.AI_ANCHOR), SF.AI_ANCHOR).count(SF.MARK), 1)


def test_speed_fix_keeps_pages_light():
    """日本語Webフォントと先読みの計測タグが戻っていないこと。1ページ1MB超の主因だった"""
    print(chr(10) + "■ 表示速度")
    import speed_fix as S
    sample = ('<link rel="preconnect" href="https://fonts.googleapis.com">\n'
              '<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>\n'
              '<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+JP:wght@400;700&family=Outfit:wght@500&display=swap" rel="stylesheet">\n'
              '<script async src="https://www.googletagmanager.com/gtag/js?id=G-TEST1"></script>\n'
              "<script>window.dataLayer=window.dataLayer||[];function gtag(){dataLayer.push(arguments);}gtag('js',new Date());gtag('config','G-TEST1');</script>\n"
              '<img src="/a.png"><img src="/b.png">')
    once = S.fix_html(sample, "ai-lab", lazy_images=True)
    twice = S.fix_html(once, "ai-lab", lazy_images=True)
    check("日本語Webフォントの読み込みを外す", "Noto+Sans+JP" in once, False)
    check("Google Fonts を一切読まない", "fonts.googleapis" in once, False)
    check("計測タグは描画後に読む", "<script async src" not in once and "addEventListener('load'" in once and "G-TEST1" in once, True)
    check("最初の画像以外を遅延読み込み", (once.count('loading="lazy"'), '<img src="/a.png">' in once), (1, True))
    check("2回当てても変わらない", once == twice, True)
    check("公開HTMLに重い読み方が残っていない", S.leftovers(ROOT / "site"), [])
    tpl = (ROOT / "templates" / "article.html").read_text(encoding="utf-8")
    check("記事の雛形も軽い", "Noto+Sans+JP" not in tpl and "<script async src" not in tpl, True)
    css = (ROOT / "site" / "css" / "style.css").read_text(encoding="utf-8")
    check("本文フォントは端末のフォント", "--sans: \"Hiragino Kaku Gothic ProN\"" in css, True)


def test_entities_link_articles_to_official_sources():
    """記事の実体（about/mentions）が実在の公式URLに結ばれ、3サイトで同じ表を使うこと"""
    print(chr(10) + "■ 実体の構造化データ")
    import entities as E, json as _j
    check("実体のURLは全て https", all(e["sameAs"].startswith("https://") for e in E.ENTITIES), True)
    check("実体名は重複しない", len({e["name"] for e in E.ENTITIES}), len(E.ENTITIES))
    a, m = E.about_and_mentions("IT導入補助金のパソコン購入", "gBizIDを取得。SEOULは無関係。ChatGPTで下書き")
    check("題名の実体は about に", [x["name"] for x in a], ["IT導入補助金"])
    check("本文の実体は mentions に（英字は語の途中に当てない）", [x["name"] for x in m], ["gBizID", "ChatGPT"])
    check("about は最大3件", len(E.about_and_mentions("IT導入補助金 ものづくり補助金 持続化補助金 事業再構築補助金", "")[0]), 3)
    corp = ROOT / ".publish-work" / "corporate" / "src" / "lib" / "entities.json"
    if corp.is_file():
        check("コーポレートも同じ表", _j.loads(corp.read_text(encoding="utf-8"))["entities"] == E.ENTITIES, True)
    b = (ROOT / "scripts" / "build.py").read_text(encoding="utf-8")
    check("build.py が about/mentions/knowsAbout を書く", "about_and_mentions" in b and "knowsAbout" in b, True)
    p = (ROOT / "scripts" / "publish.py").read_text(encoding="utf-8")
    check("補助金の配信も同じ実体を書く", "ABOUT_JSONLD" in p and "MENTIONS_JSONLD" in p, True)


def test_aio_rewrites_and_citation_measurement():
    """AI Overviewに取られている記事を最優先で直し、引用は実測できる形になっていること"""
    print(chr(10) + "■ AI引用の対策と実測")
    import auto_rewrite as A, ai_cite_check as C
    check("aio 種別の直し方がある", "aio" in A.WHAT and "{facts}" in A.WHAT["aio"], True)
    check("一次情報の数字は検算で許される", A.numbers("継続率86.5%") - A.numbers("") - A.numbers("74件中10件（86.5%）"), A.numbers(""))
    check("一次情報に無い数字は今までどおり止める", bool(A.numbers("継続率86.5%") - A.numbers("") - A.numbers("")), True)
    src = (ROOT / "scripts" / "ai_cite_check.py").read_text(encoding="utf-8")
    check("キーが無ければ課金せずに飛ばす", "AI_CITE=skipped" in src and "if not engines:" in src, True)
    check("自社URLはドメインで判定する", C.domain_of("https://www.ai.7senses.co.jp/aio/x/"), "ai.7senses.co.jp")
    check("1サイトの上限は20語", C.MAX_QUERIES, 20)
    wf = (ROOT / ".github" / "workflows" / "monthly-report.yml").read_text(encoding="utf-8", errors="replace")
    check("月次で実測を回す", "ai_cite_check.py" in wf, True)


def test_data_intake_publishes_only_grounded_numbers():
    """一次データは母数・期間つきで、人が入れた数字だけを公開すること。ページは構造化データつき"""
    print(chr(10) + "■ 一次データページ")
    import tempfile as _tf, pathlib as _pl, json as _j
    import data_intake as D
    from openpyxl import load_workbook
    d = _pl.Path(_tf.mkdtemp())
    p = D.make_sheet(d / "データ記入シート.xlsx")
    ds, ng, warn = D.review(D.read(p))
    check("空のシートは公開しない", (ds, bool(ng)), (None, True))
    wb = load_workbook(p); w = wb["概要"]
    vals = {"slug": "meo-reply-rate", "題名": "業種別・口コミ返信率と新患数（3,200店舗）", "説明（1〜2文）": "G-ranで運用する店舗の口コミ返信率を区分し、新患数の相対値を集計しました。",
            "母数（件数）": 3200, "母数の単位": "店舗", "対象期間（開始）": "2024-01", "対象期間（終了）": "2026-08",
            "集計方法・出典": "G-ranの運用データから月次で集計", "値の単位": "倍", "関連カテゴリ": "meo"}
    for row in w.iter_rows(min_row=2):
        k = str(row[0].value)
        for kk, v in vals.items():
            if k.startswith(kk): row[1].value = v
    wd = wb["データ"]; wd.append(["返信率80%以上", 1.6, ""]); wd.append(["返信率20%未満", 1.0, "基準"]); wb.save(p)
    ds, ng, warn = D.review(D.read(p))
    check("条件を満たせば通る", ng, [])
    check("引用用の一文に母数と期間が入る", "3,200店舗" in ds["sentence"] and "2024-01〜2026-08" in ds["sentence"], True)
    h = D.page_html(ds)
    ld = [x for x in re.findall(r'<script type="application/ld\+json">(.*?)</script>', h, re.S)]
    ok = any(_j.loads(x).get("@type") == "Dataset" for x in ld)
    check("Dataset の構造化データが入る", ok, True)
    check("表と棒グラフと引用用の一文がある", "<table>" in h and "<svg" in h and ds["sentence"] in h, True)
    w = wb["概要"]
    for row in w.iter_rows(min_row=2):
        if str(row[0].value).startswith("母数（件数）"): row[1].value = 8
    wb.save(p)
    ds2, ng2, _ = D.review(D.read(p))
    check("母数10未満は公開しない", ds2 is None and any("10未満" in x for x in ng2), True)
    # 分子・分母があれば、実数・信頼区間・振れ幅まで出す（割合だけでは根拠を確かめられない）
    ds3 = {"rows": [{"label": "A", "value": 86.5, "num": 64, "den": 74, "window": "1年以内"},
                    {"label": "B", "value": 90.0, "num": 9, "den": 10}],
           "n_unit": "件", "unit": "%", "den_label": "契約件数", "num_label": "継続", "neg_label": "解約"}
    check("実数があるデータと判定できる", D.has_counts(ds3), True)
    check("合計は実数の足し算", D.totals(ds3), (73, 84))
    lo, hi = D.wilson(9, 10)
    check("母数10の95%区間は幅が広い", (round(lo), round(hi)), (60, 98))
    lo2, hi2 = D.wilson(64, 74)
    check("母数74なら幅は狭い", (round(lo2), round(hi2)), (77, 92))
    tbl = D.counts_table(ds3)
    check("表に実数と区間と振れ幅が出る", ("64" in tbl and "10.0pt" in tbl and "合計" in tbl), True)
    check("積み上げ図に凡例が出る", ("継続" in D.stack_svg(ds3) and "解約" in D.stack_svg(ds3)), True)
    check("判定の条件が違えば節を出す", "判定の条件" in D.window_table(ds3), True)
    check("読み方の節は実際の値で説明する", "10.0ポイント" in D.reading_section(ds3), True)
    # 95%の範囲は表の数字だけでは伝わらない。図と「言えること/言えないこと」まで出す
    rs = D.range_svg(ds3)
    check("95%の範囲を図にする", ("<svg" in rs and "95%の範囲" in rs), True)
    ds4 = {**ds3, "rows": [{**ds3["rows"][0], "months": 12, "row_start": "2023-05", "row_end": "2026-09"},
                           {**ds3["rows"][1], "months": 3, "row_start": "2026-05", "row_end": "2026-09"}]}
    check("期間の差を月で数える", D.span_months("2023-05", "2026-09"), 40)
    tl = D.timeline_svg(ds4)
    check("観測期間と判定期間を図にする", ("40か月" in tl and "判定 12か月" in tl), True)
    check("期間が無ければ図を出さない", D.timeline_svg(ds3), "")
    lt = D.limits_table(ds4)
    check("言えること・言えないことを出す", ("言えること" in lt and "実力値としては読めない" in lt), True)
    bad = dict(ds3); bad = {**ds3, "rows": [{"label": "A", "value": "70", "num": "64", "den": "74", "window": "", "note": ""}]}
    _, ng2, _ = D.review({"overview": {"slug": "x-y-z", "題名": "十分に長い題名です", "説明（1〜2文）": "何をどう数えたかを説明する十分な長さの文です。",
                                        "母数（件数）": 74, "母数の単位": "件", "対象期間（開始）": "2024-01", "対象期間（終了）": "2026-08",
                                        "集計方法・出典": "契約一覧から集計しました", "値の単位": "%", "関連カテゴリ": "seo"},
                          "rows": bad["rows"]})
    check("値が分子分母と合わなければ止める", any("合いません" in x for x in ng2), True)
    b = (ROOT / "scripts" / "build.py").read_text(encoding="utf-8")
    check("sitemap と記事の枠に配線されている", 'glob("*/index.html")' in b and "datasets_box" in b, True)
    check("intake_watch がデータシートを振り分ける", "data_intake" in (ROOT / "scripts" / "intake_watch.py").read_text(encoding="utf-8"), True)


def test_site_has_two_axes_and_no_orphans():
    """手法だけでなく業種でも記事に行けること。入口ページが孤立していないこと。

    実測で、一次データの入口（/data/）に内部リンクが1本も無く、誰もたどり着けなかった。
    記事は手法（AIO/SEO/MEO）でしか分類されておらず、「クリニックの集客」を探す読者は
    4カテゴリに散った21本を自力で集めるしかなかった。
    """
    print(chr(10) + "■ サイト構成（2軸と入口）")
    import collections
    import industry_hub as IH
    inds, mn = IH.load()
    check("業種の定義がある", len(inds) >= 5 and mn >= 3, True)
    check("具体的な業種を先に判定する（歯科は医院を含む）",
          IH.detect("歯科医院のMEO対策", "歯科 meo"), "shika")
    check("当てはまらない記事は業種なし", IH.detect("AIO対策とは", "aio 対策"), None)
    # 内部リンクの実測（生成済みの公開ページで見る）
    site = ROOT / "site"
    pages, links, inbound = {}, collections.defaultdict(set), collections.Counter()
    for p in site.rglob("index.html"):
        u = "/" + str(p.parent.relative_to(site)).replace("\\", "/") + "/"
        pages[u if u != "/./" else "/"] = p
    for u, p in pages.items():
        body = re.sub(r"<script.*?</script>|<style.*?</style>", "",
                      p.read_text(encoding="utf-8", errors="surrogateescape"), flags=re.S)
        for h in re.findall(r'href="(/[^"#?]*)"', body):
            h = h if h.endswith("/") else h + "/"
            if h in pages and h != u:
                links[u].add(h)
                inbound[h] += 1
    depth, q = {"/": 0}, ["/"]
    while q:
        u = q.pop(0)
        for v in links.get(u, ()):
            if v not in depth:
                depth[v] = depth[u] + 1
                q.append(v)
    lost = [u for u in sorted(pages) if u not in depth and u != "/thanks/"]
    check("トップからたどり着けないページが無い", lost, [])
    for u in ("/data/", "/industry/", "/download/"):
        check(f"{u} に内部リンクがある", inbound[u] > 0, True)
    hubs = sorted(p.parent.name for p in (site / "industry").glob("*/index.html"))
    check("業種ハブが作られている", len(hubs) >= 5, True)
    check("業種ハブは一覧からたどれる", all(inbound[f"/industry/{h}/"] > 0 for h in hubs), True)
    llms = (site / "llms.txt").read_text(encoding="utf-8")
    check("llms.txt に業種の目次がある", "## 業種から探す" in llms, True)
    check("llms.txt の業種は1ブロックだけ", llms.count("## 業種から探す"), 1)
    sm = (site / "sitemap.xml").read_text(encoding="utf-8")
    check("sitemap に業種ハブが載る", all(f"/industry/{h}/" in sm for h in hubs), True)
    # ナビはサイト全体で1つ。固定ページは手書きで、生成ページと中身が食い違っていた
    import sync_nav
    navs = set()
    for p in site.rglob("*.html"):
        m = sync_nav.G_RX.search(p.read_text(encoding="utf-8", errors="surrogateescape"))
        if m:
            navs.add(tuple(re.findall(r'href="([^"]+)"', m.group(2))))
    check("ナビは全ページで同じ", len(navs), 1)
    check("ずれているページが無い", sync_nav.run(False), [])


def test_search_engines_are_told_about_all_sites():
    """公開した記事を、3サイトぶん検索エンジンへ知らせていること。

    以前は AI集客ラボ の sitemap しか見ておらず、コーポレートと補助金は記事を出しても
    通知していなかった。実測で、公開から13日以内に一度でも検索結果に出た割合が
    AI集客ラボ72%に対しコーポレート41%と差が出ていた（2026-09-22）。
    """
    print(chr(10) + "■ 検索エンジンへの通知")
    now = (ROOT / "scripts" / "notify_indexnow.py").read_text(encoding="utf-8")
    idx = (ROOT / "scripts" / "notify_indexing.py").read_text(encoding="utf-8")
    check("IndexNow が全サイトを回る", "S.load_all()" in now, True)
    check("鍵ファイルの有無を先に確かめる", "def key_ok" in now, True)
    check("Indexing API が全サイトを回る", "S.load_all()" in idx, True)
    check("1サイト固定のURLが残っていない",
          ('SITE_URL = "https://ai.7senses.co.jp"' in idx), False)
    import notify_indexnow as N
    check("鍵ファイルが無いドメインは通知しない", N.key_ok("example.com", "dummykey0000000000000000"), False)


def test_lessons_are_learned_and_pruned():
    """やって分かったことが台帳に残り、記事を書く前に読まれ、増えすぎない仕組みであること。

    以前は kpi_feedback.md に文章で積むだけで、36件11.7KBを毎回まるごと読ませていた。
    機械で守れる学びまで人とAIが覚え続けており、クライアントが増えれば破綻する形だった。
    """
    print(chr(10) + "■ 学びの台帳")
    import lessons as L
    rows = L.load()
    check("台帳に学びがある", len(rows) >= 10, True)
    check("すべて必須の項目を持つ",
          all(set(("id", "at", "kind", "phase", "rule", "gate", "status")) <= set(r) for r in rows), True)
    # 機械が守るものは読ませない（覚える対象を増やさないため）
    gated = [r for r in rows if r.get("gate")]
    check("機械が守る学びがある", len(gated) >= 1, True)
    br = L.brief("ai-lab")
    check("機械が守る学びは読ませない", any(g["rule"] in br for g in gated), False)
    check("読ませる量に上限がある", len(br) <= L.BRIEF_CHARS + 200, True)
    check("読ませる件数に上限がある", br.count(chr(10) + "- ") <= L.BRIEF_MAX, True)
    # 同じ学びを二重に積まない
    before = len(L.load())
    r1, made1 = L.add("failure", "kw", "（検査用）同じ学びは二重に積まない")
    r2, made2 = L.add("failure", "kw", "（検査用）同じ学びは二重に積まない")
    check("同じ学びは1件だけ", (made1, made2), (True, False))
    rows = [r for r in L.load() if r["id"] != r1["id"]]
    L.save(rows)
    check("後片付けできた", len(L.load()), before)
    del r2
    # 次からこうする、の一文を本文から取れる
    detail = "2026-09-01、図解が5個に切られた。原因は描画の上限。次回はPhase 3の構成審査時点でflow型は5個以内に収めること。"
    check("対処の一文を拾う", "5個以内に収めること" in L.best_rule(detail), True)
    # 配線
    prompt = (ROOT / "automation" / "multi_site_prompt.txt").read_text(encoding="utf-8")
    check("記事を書く前に読ませる", "lessons.py --brief" in prompt, True)
    check("終わったら積ませる", "lessons.py --add" in prompt, True)
    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8", errors="replace")
    check("週次で棚卸しする", "lessons.py --review" in wf, True)


def test_totals_never_come_from_a_dimensioned_query():
    """表示・クリックの合計を GSC の次元つき問い合わせから数えていないか。

    query 次元は検索数の少ない語を返さない。合計をこれで数えると大きく減る。
    実測（28日・2026-08-23〜09-19）では、補助金サイトのクリックが
    実際51回のところ2回、コーポレートが43回のところ17回に見えていた。
    その結果 growth_guard は伸び率を +38% と報告していたが、実際は +107% で、
    data_sanity は毎週「GA4とGSCが食い違う」と誤報し、指名検索の割合を
    35%以下のところ88%と出していた。数字は施策の優先順位まで動かすため、
    書き方そのものを禁じる（CLAUDE.md 0.1）。
    """
    bad = []
    for name in ("growth_guard.py", "data_sanity.py", "gsc_check.py", "daily_kpi.py",
                 "monthly_report.py", "brand_search.py", "local_kpi.py"):
        p = ROOT / "scripts" / name
        if not p.is_file():
            continue
        src = p.read_text(encoding="utf-8")
        for i, line in enumerate(src.splitlines(), 1):
            s = line.strip()
            # 「合計を出す」形（sum(... for ... in G.q(..., ["query"...]))）だけを止める。
            # 内訳のために語ごとに取るのは正しい使い方なので触らない
            if "sum(" in s and "G.q(" in s and '["query"' in s:
                bad.append(f"{name}:{i} {s[:78]}")
    if bad:
        print("  NG  合計を次元つきで数えています（次元なしで取ること）:")
        for b in bad:
            print(f"      {b}")
        FAIL.append("totals_from_dimensioned_query")
        return
    # 直したあとの形が残っているかも確かめる（消されたら気づけない）
    gg = (ROOT / "scripts" / "growth_guard.py").read_text(encoding="utf-8")
    ds = (ROOT / "scripts" / "data_sanity.py").read_text(encoding="utf-8")
    # 次元なしで直接取っているか、measure（2通りの照合）に任せているか。
    # measure 経由のほうが強い（一致しなければ値そのものが返らない）
    ms = (ROOT / "scripts" / "measure.py")
    for name, src in (("growth_guard.py", gg), ("data_sanity.py", ds)):
        direct = "None, 1)" in src
        delegated = "measure" in src and ms.is_file() and "None, 1)" in ms.read_text(
            encoding="utf-8")
        if not (direct or delegated):
            print(f"  NG  {name} が次元なしの合計を取っていません")
            FAIL.append("totals_from_dimensioned_query")
            return
    print("  OK  表示・クリックの合計は次元なしで取っている（measure で2通り照合）")


def test_stuck_articles_are_pushed_every_week():
    """11〜30位で止まった記事に、毎週手が入る形になっているか。

    伸び率の主因は順位で、詰まりは11〜30位にある（実測で244語・表示1,496回・
    クリック0）。ここに手が入らなければ、新規記事をいくら足しても伸びは戻らない。

    見るのは4つ。
      1. 押し上げの対象帯が21〜30位まで届いているか（以前は20.5位で切れていた）
      2. 診断が週次で動くようになっているか
      3. 内部リンクが「順位で止まっている記事」へ寄るようになっているか
      4. 書き換えの検算で kw_guard に自分自身を除外させているか
         （渡さないと、その記事自身が食い合い相手になり100%差し戻される）
    """
    pb = (ROOT / "scripts" / "priority_boost.py").read_text(encoding="utf-8")
    m = re.search(r"^NEAR = \((\d+\.?\d*),\s*(\d+\.?\d*)\)", pb, re.M)
    if not m or float(m.group(2)) < 30.0:
        print(f"  NG  priority_boost の対象帯が狭すぎます（{m.group(0) if m else '不明'}）"
              "。21〜30位が対象外になります")
        FAIL.append("stuck_articles_are_pushed")
        return

    rr = ROOT / "scripts" / "rank_rescue.py"
    if not rr.is_file():
        print("  NG  scripts/rank_rescue.py がありません（止まった記事の診断）")
        FAIL.append("stuck_articles_are_pushed")
        return

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    for need, why in (("rank_rescue.py", "止まった記事の診断が週次で動きません"),
                      ("--rescue", "内部リンクが順位で止まった記事へ寄りません")):
        if need not in wf:
            print(f"  NG  週次に {need} がありません: {why}")
            FAIL.append("stuck_articles_are_pushed")
            return

    ar = (ROOT / "scripts" / "auto_rewrite.py").read_text(encoding="utf-8")
    if "--exclude-slug" not in ar:
        print("  NG  auto_rewrite が kw_guard に --exclude-slug を渡していません。"
              "その記事自身が食い合い相手になり、リライトは必ず差し戻されます")
        FAIL.append("stuck_articles_are_pushed")
        return
    if '"stuck"' not in ar:
        print("  NG  auto_rewrite に stuck 種別がありません（語の欠落を直せません）")
        FAIL.append("stuck_articles_are_pushed")
        return

    import shutil
    exe = "claude" if shutil.which("claude") else ""
    note = "" if exe else "（この端末には claude がないため、書き換えは動きません）"
    print(f"  OK  11〜30位の記事に毎週手が入る（診断・内部リンク・書き換え）{note}")


def test_publish_gate_actually_blocks():
    """品質検査に落ちた新規記事が、本当に公開されないか。

    CLAUDE.md は「警告ゼロが公開条件」と書いているが、build.py は長らく
    WARN を印字するだけで、記事はそのまま site/ に書き出され配信されていた。
    実測（2026-09-23）で、文字数不足・図解の項目超過・狙う語がタイトルに無い
    記事が公開済みだった。印字は検査ではない。ここで実際に止まるかを見る。

    止めるのは未公開の記事だけにする。公開済みを後から取り下げると、
    取れている順位まで失うため。
    """
    import shutil
    import subprocess

    src = (ROOT / "scripts" / "build.py").read_text(encoding="utf-8")
    for need, why in (("QUALITY_ISSUES", "品質不合格を集める入れ物がありません"),
                      ("published.json", "公開済みの台帳がありません"),
                      ("title_has_keyword", "狙う語がタイトルにあるかの検査がありません")):
        if need not in src:
            print(f"  NG  {why}")
            FAIL.append("publish_gate_blocks")
            return

    slug = "zz-gate-selftest-kiji"
    NL = chr(10)
    md = ROOT / "articles" / f"{slug}.md"
    body = (
        "---" + NL
        + "title: 門番の自己診断に使う記事｜公開されてはいけません" + NL
        + "description: このファイルはゲートが実際に公開を止めるかを確かめるために"
        + "その場で作られ、確認が終わると必ず消されます。公開されたら不具合です。" + NL
        + f"slug: {slug}" + NL
        + "keyword: 門番 自己診断" + NL
        + "category: meo" + NL
        + "date: 2026-01-01" + NL
        + "depth: standard" + NL
        + "score: 95" + NL
        + "---" + NL + NL
        + "## 短すぎる見出し" + NL + NL
        + "基準に届かない短い本文です。" + NL)
    try:
        md.write_text(body, encoding="utf-8", newline="")
        env = dict(os.environ, PYTHONIOENCODING="utf-8")
        r = subprocess.run([sys.executable, "scripts/build.py"], cwd=ROOT, env=env,
                           capture_output=True, text=True, encoding="utf-8",
                           errors="replace", timeout=1800)
        out = r.stdout or ""
        blocked = f"BLOCKED(公開不可): {slug}" in out
        html = list((ROOT / "site").glob(f"*/{slug}/index.html"))
        smap = (ROOT / "site" / "sitemap.xml").read_text(encoding="utf-8")
        broke = [l for l in out.splitlines() if "リンク切れ" in l and slug in l]
        if not blocked:
            print("  NG  品質検査に落ちた新規記事が BLOCKED になりません")
            FAIL.append("publish_gate_blocks")
        elif html:
            print("  NG  止めたはずの記事のHTMLが作られています")
            FAIL.append("publish_gate_blocks")
        elif slug in smap:
            print("  NG  止めたはずの記事が sitemap に載っています")
            FAIL.append("publish_gate_blocks")
        elif broke:
            print(f"  NG  止めた記事への内部リンクが残っています（{len(broke)}件）")
            FAIL.append("publish_gate_blocks")
        else:
            print("  OK  品質検査に落ちた新規記事は公開されない（HTML・sitemap・リンクとも）")
    finally:
        md.unlink(missing_ok=True)
        for d in (ROOT / "site").glob(f"*/{slug}"):
            shutil.rmtree(d, ignore_errors=True)
        subprocess.run([sys.executable, "scripts/build.py"], cwd=ROOT,
                       env=dict(os.environ, PYTHONIOENCODING="utf-8"),
                       capture_output=True, text=True, encoding="utf-8",
                       errors="replace", timeout=1800)


def test_coverage_matrix_shows_gaps():
    """どの業種×手法が空いているかを、毎週見ているか。

    記事は1本ずつ選ばれており、全体としてどこが埋まっているかを誰も見ていなかった。
    実測（2026-09-23）で、優先業種の工務店とリフォームに AIO の記事が1本も無い
    一方、コーポレートは経理×経理に67本が集中していた。
    検索エンジンとAIが「この分野を扱うサイト」と見るのは、1本の出来ではなく
    面の広さと深さによる。空いたマスは、そのまま次に書くべき記事になる。

    業種の軸はサイト固有のものを使う。共通の一覧を当てると、経理BPOのサイトに
    歯科医院のマスが並び、埋まるはずのない空きが37マス出る。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    p = ROOT / "scripts" / "coverage.py"
    if not p.is_file():
        print("  NG  scripts/coverage.py がありません（盤面の検査）")
        FAIL.append("coverage_matrix")
        return

    import importlib
    import coverage as CV
    importlib.reload(CV)

    ok = True
    for sid in ("ai-lab", "corporate", "subsidy"):
        inds = CV.industries(sid)
        cs = CV.cores(sid)
        if not inds or not cs:
            print(f"  NG  {sid}: 盤面の軸が作れません（業種{len(inds)} / 手法{len(cs)}）")
            ok = False
            continue
        cells, _, _, arts, _ = CV.matrix(sid)
        if not arts:
            print(f"  NG  {sid}: 記事を1本も拾えていません")
            ok = False
    if not ok:
        FAIL.append("coverage_matrix")
        return

    # サイト固有の軸を使っているか（共通一覧を全サイトに当てていないか）
    a = {i["name"] for i in CV.industries("ai-lab")}
    c = {i["name"] for i in CV.industries("corporate")}
    if a == c:
        print("  NG  業種の軸が全サイトで同じです。サイト固有の起点を使ってください")
        FAIL.append("coverage_matrix")
        return

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    fd = (ROOT / "scripts" / "findings.py").read_text(encoding="utf-8")
    if "coverage.py" not in wf:
        print("  NG  週次で盤面を出していません")
        FAIL.append("coverage_matrix")
        return
    if "coverage.py" not in fd:
        print("  NG  盤面の結果が通知に載りません（findings に入っていません）")
        FAIL.append("coverage_matrix")
        return
    print("  OK  業種×手法の盤面を毎週出し、空きマスを通知に載せている（3サイト）")


def test_howto_is_emitted_for_step_articles():
    """手順記事に HowTo 構造化データが出ているか。

    CLAUDE.md は「手順系記事のみ HowTo 追加」と書いているが実装が無く、
    手順が3つ以上ある記事95本すべてで0件だった（2026-09-23 実測）。
    HowTo は2023年にGoogleのリッチリザルトから外れたため検索結果の見た目は
    変わらない。効くのは機械可読性で、AI検索は本文より先に JSON-LD を読む。
    手順が「何番目に何をするか」で取り出せるかは、引用の正確さに効く。

    あわせて、出している JSON-LD が全ページで壊れていないことも見る。
    壊れていても画面には出ないため、検査しないと気づけない。
    """
    import glob
    import json as _json
    n_howto = n_block = broken = 0
    for p in glob.glob(str(ROOT / "site" / "**" / "index.html"), recursive=True):
        h = Path(p).read_text(encoding="utf-8", errors="replace")
        for m in re.finditer(r'<script type="application/ld\+json">(.*?)</script>', h, re.S):
            n_block += 1
            try:
                d = _json.loads(m.group(1))
            except Exception:
                broken += 1
                continue
            for g in (d.get("@graph") or [d]):
                if g.get("@type") == "HowTo":
                    n_howto += 1
                    if len(g.get("step") or []) < 3:
                        print(f"  NG  HowTo の手順が3つ未満です: {p}")
                        FAIL.append("howto_emitted")
                        return
    if broken:
        print(f"  NG  壊れた JSON-LD が {broken}件あります（画面には出ないので検査でしか見つかりません）")
        FAIL.append("howto_emitted")
        return
    if n_howto < 20:
        print(f"  NG  HowTo が {n_howto}件しか出ていません（手順記事に出ていない疑い）")
        FAIL.append("howto_emitted")
        return
    print(f"  OK  手順記事に HowTo を出している（{n_howto}件 / JSON-LD {n_block}件すべて正常）")


def test_guarantee_rate_is_measured():
    """保証できる部分が数字になっているか。

    順位そのものは保証できない。決めるのはGoogleで、アルゴリズムの更新も
    競合の動きもこちらの外側にある。保証できるのは「順位を決める要因のうち
    こちら側にあるものを毎回100%満たすこと」で、それは数字にできる。

    あわせて、達成率に混ぜる基準に根拠があることも見る。
    根拠の無い基準（実測で否定された下限12本など）を混ぜると、
    達成率そのものが意味を失う。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    p = ROOT / "scripts" / "guarantee.py"
    if not p.is_file():
        print("  NG  scripts/guarantee.py がありません")
        FAIL.append("guarantee_rate")
        return
    src = p.read_text(encoding="utf-8")

    # 保証しないものを「保証する」と書いていないか
    for ng in ("順位を保証", "上位表示を保証", "確実に上位"):
        if ng in src:
            print(f"  NG  保証できないことを保証すると書いています: {ng}")
            FAIL.append("guarantee_rate")
            return
    if "順位そのものは保証できない" not in src:
        print("  NG  何を保証しないのかが書かれていません")
        FAIL.append("guarantee_rate")
        return

    # 実測で否定された下限を、達成率に数えていないか
    if '"止まっている記事の内部リンクが12本以上", n_stuck_ok' in src:
        print("  NG  裏づけの無い基準（内部リンク12本）を達成率に数えています")
        FAIL.append("guarantee_rate")
        return

    import importlib
    import guarantee as G
    importlib.reload(G)
    for fn in ("published", "article_checks", "schema_coverage", "site_checks"):
        if not hasattr(G, fn):
            print(f"  NG  guarantee.{fn} がありません")
            FAIL.append("guarantee_rate")
            return

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    fd = (ROOT / "scripts" / "findings.py").read_text(encoding="utf-8")
    if "guarantee.py" not in wf or "guarantee.py" not in fd:
        print("  NG  達成率が週次か通知に載っていません")
        FAIL.append("guarantee_rate")
        return
    print("  OK  こちら側で決まる要因の達成率を毎週出し、通知に載せている")


def test_interventions_are_measured_against_control():
    """打った手を、触っていない記事と比べているか。

    サイト全体が28日で+107%伸びている時期には、何をしても「効いた」に見える。
    実際 effect.py は「表示が増えた記事 55/126本」と出すが、触っていない記事でも
    同じ割合で増えているなら、その手は効いていない。
    対照群と比べて初めて、効かない施策を落とせる。

    あわせて、記録が残っているかも見る。残っていない施策は測りようがない。
    実測（2026-09-23）で link_boost は1件も記録しておらず、内部リンクが
    効いたかを判定できなかった。
    """
    p = ROOT / "scripts" / "effect_ab.py"
    if not p.is_file():
        print("  NG  scripts/effect_ab.py がありません（対照群との比較）")
        FAIL.append("effect_ab")
        return
    src = p.read_text(encoding="utf-8")
    for need, why in (("touched_days", "触った記事を除いた対照群を作っていません"),
                      ("ctrl_imp", "対照群の動きを測っていません")):
        if need not in src:
            print(f"  NG  {why}")
            FAIL.append("effect_ab")
            return

    # 自動修正は、どれも台帳に記録していること
    for name in ("auto_review.py", "auto_rewrite.py", "link_boost.py"):
        s = (ROOT / "scripts" / name).read_text(encoding="utf-8")
        if "auto_fix.jsonl" not in s:
            print(f"  NG  {name} が台帳に記録していません（あとで効果を測れません）")
            FAIL.append("effect_ab")
            return

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    fd = (ROOT / "scripts" / "findings.py").read_text(encoding="utf-8")
    if "effect_ab.py" not in wf or "effect_ab.py" not in fd:
        print("  NG  対照群との比較が週次か通知に載っていません")
        FAIL.append("effect_ab")
        return
    print("  OK  打った手を対照群と比べ、記録は3つの自動修正すべてが残している")


def test_numbers_come_from_two_methods():
    """数字が2通りで一致したものだけになっているか。

    CLAUDE.md 0.1 は「1つの測り方の結果をそのまま事実として報告しない」と
    決めているが、守るかどうかが人とAIの側にあったため同じ誤りが繰り返された。
    実測（2026-09-23）だけで、合計をquery次元で数える誤りが2箇所、
    検出器の書き方が違う誤りが3箇所あった。
    measure.verified を通せば、一致しないかぎり値が取れない。

    あわせて、自動で書き込む工程すべてが検算を持っていることも見る。
    link_boost は検算を持たない唯一の工程で、実際に表を壊した。
    """
    sys.path.insert(0, str(ROOT / "scripts"))
    p = ROOT / "scripts" / "measure.py"
    if not p.is_file():
        print("  NG  scripts/measure.py がありません（2通りの照合）")
        FAIL.append("two_method_numbers")
        return

    import importlib
    import measure as M
    importlib.reload(M)
    # 食い違ったときに本当に止まるか、その場で試す
    try:
        M.verified("ゲートの確認", lambda: 100, lambda: 51, 0.0)
        print("  NG  食い違っても値を返しています（検算が効いていません）")
        FAIL.append("two_method_numbers")
        return
    except M.Disagree:
        pass
    try:
        if M.verified("ゲートの確認", lambda: 100, lambda: 100, 0.0) != 100:
            raise AssertionError
    except Exception:
        print("  NG  一致しているのに値を返しません")
        FAIL.append("two_method_numbers")
        return

    # 合計を measure 経由で取っているか
    for name in ("growth_guard.py", "data_sanity.py"):
        s = (ROOT / "scripts" / name).read_text(encoding="utf-8")
        if "measure" not in s:
            print(f"  NG  {name} が2通りの照合を通していません")
            FAIL.append("two_method_numbers")
            return

    # 自動で書き込む工程は、すべて検算を持つこと
    checks = {"auto_rewrite.py": "def check(", "auto_review.py": "def guard(",
              "auto_merge.py": "def ", "link_boost.py": "def insert_ok("}
    for name, need in checks.items():
        s = (ROOT / "scripts" / name).read_text(encoding="utf-8")
        if need not in s:
            print(f"  NG  {name} に検算がありません（書きっぱなしは壊れます）")
            FAIL.append("two_method_numbers")
            return
    print("  OK  数字は2通りで一致したものだけ／書き込む工程はすべて検算を持つ")


def test_structure_is_proposed_monthly():
    """サイト構成の案が、月次レポートに自動で載るか。

    「次に何を作るか」を人が考えていた。1本ずつ選ぶと同じマスに重なり、
    空いたマスが残る（実測でクリニック21本に対し工務店のAIOは0本）。
    盤面・止まっている記事・食い合いから機械が案を作り、月次に載せる。
    """
    p = ROOT / "scripts" / "structure_plan.py"
    if not p.is_file():
        print("  NG  scripts/structure_plan.py がありません")
        FAIL.append("structure_proposed")
        return
    mr = (ROOT / "scripts" / "monthly_report.py").read_text(encoding="utf-8")
    if "structure_plan" not in mr or "{plan_html}" not in mr:
        print("  NG  月次レポートに構成の提案が入っていません")
        FAIL.append("structure_proposed")
        return
    sys.path.insert(0, str(ROOT / "scripts"))
    import importlib
    import structure_plan as SP
    importlib.reload(SP)
    rows = SP.plan("ai-lab")
    if not rows:
        print("  OK  盤面に空きが無いため提案なし（空きが出れば載ります）")
        return
    for r in rows[:3]:
        if not r.get("why"):
            print("  NG  提案に根拠がありません（実測から作ってください）")
            FAIL.append("structure_proposed")
            return
    print(f"  OK  月次レポートに構成の提案が載る（いま{len(rows)}件・すべて根拠つき）")


def test_rules_are_validated_against_outcomes():
    """使っている判断が、実際の成果を言い当てているかを確かめているか。

    この仕組みは点数・語の性質・内部リンクの下限で記事を選び、直している。
    どれも「そう決めた」だけで、当たっているかを誰も確かめていなかった。
    確かめると、当たっていないものが2件あった（2026-09-23 実測）。

      内部リンクの下限   上位ほど多いはずが、1〜10位は中央値5本で逆
      品質スコア        96点以上の順位が91〜93点より4.1位悪い（逆相関）

    **向きを見ないと、逆相関を「差が出ている」と読んでしまう。**
    最初の版が実際にそう報告した。逆に出る判断は、無い判断より悪い。
    """
    p = ROOT / "scripts" / "validate_rules.py"
    if not p.is_file():
        print("  NG  scripts/validate_rules.py がありません")
        FAIL.append("rules_validated")
        return
    src = p.read_text(encoding="utf-8")
    if "expect" not in src or "逆になって" not in src:
        print("  NG  関係の向きを見ていません（逆相関を見逃します）")
        FAIL.append("rules_validated")
        return

    sys.path.insert(0, str(ROOT / "scripts"))
    import importlib
    import validate_rules as VR
    importlib.reload(VR)
    # 逆相関を、その場で作って検出できるか試す
    p_fake = {f"a{i}": [100, 1, 100 * (10.0 + i)] for i in range(8)}
    p_fake.update({f"b{i}": [100, 1, 100 * (30.0 + i)] for i in range(8)})
    got = VR.report("診断", [("低いはず", list(p_fake)[:8]),
                             ("高いはず", list(p_fake)[8:])],
                    p_fake, expect="高いはず")
    if got is not False:
        print("  NG  逆相関を検出できません（期待と逆でも合格にしています）")
        FAIL.append("rules_validated")
        return

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    fd = (ROOT / "scripts" / "findings.py").read_text(encoding="utf-8")
    if "validate_rules.py" not in wf or "validate_rules.py" not in fd:
        print("  NG  判断の検証が週次か通知に載っていません")
        FAIL.append("rules_validated")
        return
    print("  OK  使っている判断を毎週実測で検証し、逆相関も検出する")


def test_score_is_audited_and_cta_is_enforced():
    """点数を別工程が付け直しているか。CTAの本数を機械が見ているか。

    公開条件は score >= 90 だが、その点数は記事を書いたエージェント自身が
    付けていた。372本すべてが90点以上で90点未満は0本、95点に177本が集中。
    **一度も門として働いていない。**
    別工程（公開HTMLを読む・Readのみ）で11本を採点し直すと、
    11本すべてが90点を割り、11本すべてが1観点16点未満の足切りだった。
    主因はデザイン観点で、CTAが2箇所に満たない記事が59本あったこと。
    CLAUDE.md は「CTA 最低2箇所」と決めているのに検査が無かった。
    """
    for name, why in (("score_audit.py", "別工程の採点がありません"),
                      ("cta_fill.py", "CTAを埋める工程がありません")):
        if not (ROOT / "scripts" / name).is_file():
            print(f"  NG  {why}")
            FAIL.append("score_audited")
            return

    # 採点の基準が1か所にまとまり、版が付いていること
    rb = ROOT / "scripts" / "rubric.py"
    if not rb.is_file():
        print("  NG  scripts/rubric.py がありません（採点の基準）")
        FAIL.append("score_audited")
        return
    sys.path.insert(0, str(ROOT / "scripts"))
    import importlib
    import rubric as R
    importlib.reload(R)
    if not getattr(R, "VERSION", ""):
        print("  NG  採点の基準に版がありません（変えたときに比べられません）")
        FAIL.append("score_audited")
        return
    # 各軸に、点の目安（アンカー）が書かれていること。無いと採点がぶれる
    for ax in R.AXES:
        if len(ax.get("anchors") or []) < 3:
            print(f"  NG  {ax['name']} に点の目安がありません（採点がぶれます）")
            FAIL.append("score_audited")
            return
    # 足切りが効くこと。合計だけで壊滅軸を隠さない
    # 満点は R.MAX（資料 v2 で各軸100点になった。10で固定すると基準を変えた瞬間に落ちる）
    if R.judge({a["key"]: R.MAX for a in R.AXES})["ok"] is not True:
        print("  NG  満点が合格になりません")
        FAIL.append("score_audited")
        return
    low = {a["key"]: R.MAX for a in R.AXES}
    low[R.AXES[0]["key"]] = R.PASS_EACH - 1
    if R.judge(low)["ok"] is not False:
        print("  NG  1軸が下限を割っても合格にしています（足切りが効いていません）")
        FAIL.append("score_audited")
        return
    # 書く側にも同じ基準が渡っていること
    wp = (ROOT / "automation" / "multi_site_prompt.txt").read_text(encoding="utf-8")
    for ax in R.AXES:
        if ax["name"] not in wp:
            print(f"  NG  記事を書く指示に「{ax['name']}」がありません"
                  "（採点だけ変えても書く側が古い基準を見ます）")
            FAIL.append("score_audited")
            return

    sa = (ROOT / "scripts" / "score_audit.py").read_text(encoding="utf-8")
    # 採点者が元の点数を見ないこと（公開HTMLを読ませる）
    if "site" not in sa or "Read" not in sa:
        print("  NG  採点が公開HTMLを読む形になっていません（点数に引きずられます）")
        FAIL.append("score_audited")
        return
    if '"--allowedTools", "Read"' not in sa:
        print("  NG  採点工程が書き換えられる状態です（Readだけに絞ってください）")
        FAIL.append("score_audited")
        return

    # CTAの本数を build が見ていること
    bd = (ROOT / "scripts" / "build.py").read_text(encoding="utf-8")
    if "CTA不足" not in bd:
        print("  NG  build.py がCTAの本数を見ていません")
        FAIL.append("score_audited")
        return

    # 実際に2箇所未満の記事が残っていないか
    import re as _re
    short = 0
    for p in (ROOT / "articles").glob("*.md"):
        if p.name.startswith("_"):
            continue
        t2 = p.read_text(encoding="utf-8-sig", errors="ignore")
        if not _re.search(r"^score:\s*(9[0-9]|100)\s*$", t2, _re.M):
            continue
        if t2.count("cta-button") < 2:
            short += 1
    if short:
        print(f"  NG  CTAが2箇所未満の記事が{short}本あります"
              "（python scripts/cta_fill.py --write）")
        FAIL.append("score_audited")
        return

    wf = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    if "score_audit.py" not in wf or "cta_fill.py" not in wf:
        print("  NG  採点し直しかCTAの補充が週次に入っていません")
        FAIL.append("score_audited")
        return
    print("  OK  点数は別工程が付け直し、CTAは全記事2箇所以上（機械で確認）")


def test_no_control_characters_anywhere():
    """置換が壊れて制御文字が紛れていないか。

    実測（2026-09-23）で、376個の画像が134ページで表示されていなかった。
    build.py の置換で後方参照が制御文字として埋め込まれ、
    src が chr(1) になっていた。原稿は正しく、出力だけが壊れていた。
    さらに否定先読みの \b も chr(8) になっており、条件が常に成立していた。

    **src が空ではなかったため、「空srcを探す」検査では見つからなかった。**
    見えない壊れは、範囲で探さないと見つからない。
    """
    bad = []
    for p in sorted((ROOT / "scripts").glob("*.py")) + sorted((ROOT / "tests").glob("*.py")):
        s = p.read_text(encoding="utf-8", errors="replace")
        hits = [c for c in s if ord(c) < 9 or (13 < ord(c) < 32)]
        if hits:
            bad.append(f"{p.name}（{len(hits)}個 {[hex(ord(c)) for c in hits[:3]]}）")
    if bad:
        print("  NG  スクリプトに制御文字が混ざっています: " + " / ".join(bad[:3]))
        print("      置換でエスケープが壊れています。生成物も壊れます")
        FAIL.append("no_control_chars")
        return

    out = []
    for p in (ROOT / "site").rglob("index.html"):
        b = p.read_bytes()
        if any(bytes([c]) in b for c in list(range(0, 9)) + [11, 12] + list(range(14, 32))):
            out.append(p.parent.name)
    if out:
        print(f"  NG  生成HTMLに制御文字が残っています（{len(out)}ページ / 例: {out[0]}）")
        FAIL.append("no_control_chars")
        return

    # 検査そのものが効くか、壊れた例で試す
    sys.path.insert(0, str(ROOT / "scripts"))
    import importlib
    import render_check as RC
    importlib.reload(RC)
    broken = '<img src="' + chr(1) + '" alt="x">'
    names = [n for n, fn, _ in RC.SYMPTOMS if fn(broken, broken)]
    if not names:
        print("  NG  壊れた画像を検出できません（検査が効いていません）")
        FAIL.append("no_control_chars")
        return
    print(f"  OK  制御文字なし（スクリプト・生成HTMLとも）／壊れた画像を検出できる")


def test_report_actions_close_the_loop():
    """レポートの改善点は、機械が当てて・見直して・残りだけ人へ届く形になっているか。

    レポートは「改善プラン」を毎月書いていたが、実行するのは人だった。
    月初のCIに人はいないので、レポートの直後に report_actions が動く。
    当てる道具が消えていたり、通知文から人の項目が落ちていたりすれば、
    改善は「書いただけ」に戻る。
    """
    import report_actions as RA
    for kind, cmds in RA.RUNNERS.items():
        for c in cmds:
            check(f"report_actions: {kind} の道具 scripts/{c[0]} がある",
                  (ROOT / "scripts" / c[0]).is_file(), True)
    for k in ("measure", "lp", "hub_external"):
        check(f"report_actions: 人の種別 {k} に説明がある", bool(RA.HUMAN.get(k)), True)
    txt = RA.body([{"site": "s", "finding": "f", "action": "a"}], [], "", 1)
    check("report_actions: 通知文に人の項目が出る", "人の手が要る" in txt and "f → a" in txt, True)
    txt = RA.body([], [], "build.py", 1)
    check("report_actions: 検算に落ちた回は通知文で目立つ", "検算" in txt and "捨てました" in txt, True)
    yml = (ROOT / ".github" / "workflows" / "monthly-report.yml").read_text(encoding="utf-8")
    check("月次CIがレポートの直後に改善を当てる", "report_actions.py --apply" in yml, True)

    import ai_kw_research as AK
    check("ai_kw_research: 質問形を見分ける", bool(AK.Q_PAT.search("aio対策 とは")), True)
    check("ai_kw_research: 質問形でない語は点が低い",
          AK.score("aio 導入 とは 費用", 100, 15, None) > AK.score("aio 導入", 100, 15, None), True)
    check("ai_kw_research: AIが答えを出して自社が無い語が最上位",
          AK.score("aio とは", 10, 15, {"answered": True, "ours": False})
          > AK.score("aio とは", 10, 15, {"answered": True, "ours": True}), True)
    wk = (ROOT / ".github" / "workflows" / "weekly-optimize.yml").read_text(encoding="utf-8")
    check("週次CIがAI検索の語を調べる", "ai_kw_research.py" in wk, True)

    # 3倍計画: 道筋は単調に増え、6か月目で起点の3倍。遅れは「要対応」で知らせる
    import growth_plan as GP
    base = {"sessions": 100, "clicks": 50, "cv": 4, "ai": 10}
    p6 = GP.path_for(base, GP.MONTHS)
    check("growth_plan: 6か月目で3倍", p6["sessions"], 300)
    check("growth_plan: 道筋が単調に増える",
          all(GP.path_for(base, k)["clicks"] <= GP.path_for(base, k + 1)["clicks"] for k in range(GP.MONTHS)), True)
    check("growth_plan: 打ち手にYouTube言及が最上位にある", GP.LEVERS[0][0].startswith("YouTube"), True)
    check("月次CIが3倍計画の進みを見る", "growth_plan.py --check" in yml, True)
    import article_videos as AV
    import inspect
    check("article_videos: 鍵が無いときの要対応文がある",
          "youtube-token.json" in inspect.getsource(AV.note_token_missing), True)
    check("週次CIが記事動画を作る", "article_videos.py --write" in wk, True)
    import hub_client as HC
    check("next_kw: AI回答ありの語を読む関数がある", callable(getattr(HC, "_ai_targets", None)), True)

    # 当てた手は戻せる・表示ゼロは整理する・鮮度は更新する・FAQと動画は記事から機械が作る
    import rewrite_rollback as RB
    import retire_stale as RS
    import video_embed as VE
    import industry_hub as IH
    import auto_rewrite as AR2
    check("rollback: 対照の8割を割ったら戻す", RB.WORSE, 0.8)
    check("retire_stale: 近い記事が無ければ統合しない（Dice下限）", RS.MIN_SIM >= 0.3, True)
    check("auto_rewrite: 鮮度の種別がある", "fresh" in AR2.WHAT, True)
    check("auto_rewrite: 直す前のタイトルを台帳に残す", "before_title" in inspect.getsource(AR2.run_one), True)
    fb = IH.faq_body({"slug": "x", "name": "テスト"},
                     [{"faq": [{"q": f"q{i}", "a": f"a{i}"}], "title": "t", "category": "c", "slug": f"s{i}"} for i in range(5)],
                     lambda m: "/c/" + m["slug"] + "/")
    check("industry_hub: FAQが5問以上ならページを作る", bool(fb) and fb[1]["@type"] == "FAQPage", True)
    check("industry_hub: 4問以下は作らない", IH.faq_body({"slug": "x", "name": "t"}, [], lambda m: "/"), None)
    check("video_embed: 台帳に無い記事は何もしない", VE.prepend("<p>x</p>", {"slug": "__none__"}), "<p>x</p>")
    check("週次CIが戻す・整理・鮮度を回す",
          all(s in wk for s in ("rewrite_rollback.py --write", "retire_stale.py --write", "--kind fresh")), True)
    check("月次CIが15日に途中経過を出す", '"0 0 15 * *"' in yml and "steps.day.outputs.through" in yml, True)
    gs = (ROOT / "automation" / "gas" / "contact.hub.gs").read_text(encoding="utf-8")
    check("自動返信に資料の案内が入る", "videos/aio-pr.mp4" in gs and "body + materials + foot" in gs, True)

    # 実測系: 検出器は「見つかるはずの例」で試す（0.1節）
    import brand_spelling as BS
    got = BS.scan("セブンセンシズ(株) と AI 集客ラボ。https://7senses.co.jp/")
    check("brand_spelling: ゆれを見つける", sum(sum(v.values()) for v in got.values()), 2)
    check("brand_spelling: 正規表記を誤検出しない", BS.scan("セブンセンシズ株式会社のAI集客ラボ"), {})
    check("brand_spelling: URLは触らない", "https://7senses.co.jp/" in BS.fix("セブンセンシズ(株) https://7senses.co.jp/"), True)
    import cwv_check as CW
    check("cwv_check: 基準は 8.4節と同じ", (CW.LIMITS["LCP"], CW.LIMITS["INP"], CW.LIMITS["CLS"]), (2500, 200, 0.1))
    import cert_check as CC2
    check("cert_check: 30日前に知らせる", CC2.WARN_DAYS, 30)
    check("auto_rewrite: 質問形の種別がある", "question" in AR2.WHAT, True)
    check("auto_rewrite: 質問形はH2の本数を守る", "H2の本数が変わりました" in inspect.getsource(AR2.run_one), True)
    import social_post as SP2
    check("social_post: note 用の長文を作る", '"note"' in inspect.getsource(SP2.compose), True)
    check("週次CIが実測（速度・表記・インデックス・AI引用）を回す",
          all(s in wk for s in ("cwv_check.py", "brand_spelling.py --fix", "index_status.py", "ai_cite_check.py --limit 5", "--kind question")), True)
    sh_ = (ROOT / ".github" / "workflows" / "selfheal.yml").read_text(encoding="utf-8")
    check("日次CIが証明書・404・トークンを見る", "cert_check.py" in sh_ and "token_check.py" in sh_, True)


def main():
    for t in (test_kw_conflicts, test_tag_balance, test_char_count, test_hub_gas,
              test_self_exclusion, test_published_not_rewritten_as_new,
              test_token_check_probes_write, test_selfheal_watches_real_workflows,
              test_lab_links_to_corporate, test_daily_audit_ignores_unscored_drafts,
              test_every_article_has_a_lead_path,
              test_token_never_in_command_line,
              test_kw_intent_separates_click_need,
              test_lead_funnel_is_watched,
              test_cta_wording_is_measured_not_guessed,
              test_each_site_declares_what_it_sells,
              test_next_keyword_prefers_main_offer,
              test_improvements_apply_themselves,
              test_auto_fixes_are_reviewed,
              test_client_onboarding_is_one_sheet,
              test_quality_gate_holds_on_wordpress,
              test_daily_todo_fits_in_one_run,
              test_near_page1_is_pushed_every_week,
              test_main_category_actually_grows,
              test_paragraph_split_keeps_text,
              test_anchor_text_stays_readable,
              test_block_breaks_render_correctly,
              test_rendered_html_is_checked,
              test_import_has_no_side_effects,
              test_long_sentences_are_split,
              test_rewrite_is_verified_and_reverted,
              test_output_matches_source,
              test_inquiry_body_is_not_truncated,
              test_articles_are_not_uniform,
              test_facts_are_per_article,
              test_unindexed_pages_are_chased,
              test_rate_claims_need_evidence,
              test_submissions_are_not_double_counted,
              test_boost_finds_link_sources,
              test_rank_data_is_verified,
              test_measurement_pitfalls_are_documented,
              test_built_tools_actually_run,
              test_quality_fixes_run_before_publishing,
              test_findings_judges_by_marker_not_exit_code,
              test_detection_scripts_do_not_fail_the_run,
              test_notify_does_not_send_junk,
              test_routine_success_is_not_emailed,
              test_every_site_gets_articles,
              test_intake_sheets_are_not_published,
              test_kw_plan_keeps_only_buyers,
              test_reports_carry_diagnosis_and_next_actions,
              test_rewrites_reach_the_sheet,
              test_hub_has_one_kpi_writer,
              test_five_hub_features_are_wired,
              test_merge_is_wired,
              test_jisseki_intake_never_invents_numbers,
              test_speed_fix_keeps_pages_light,
              test_entities_link_articles_to_official_sources,
              test_aio_rewrites_and_citation_measurement,
              test_data_intake_publishes_only_grounded_numbers,
              test_site_has_two_axes_and_no_orphans,
              test_search_engines_are_told_about_all_sites,
              test_lessons_are_learned_and_pruned,
              test_totals_never_come_from_a_dimensioned_query,
              test_stuck_articles_are_pushed_every_week,
              test_publish_gate_actually_blocks,
              test_coverage_matrix_shows_gaps,
              test_howto_is_emitted_for_step_articles,
              test_guarantee_rate_is_measured,
              test_interventions_are_measured_against_control,
              test_numbers_come_from_two_methods,
              test_structure_is_proposed_monthly,
              test_rules_are_validated_against_outcomes,
              test_score_is_audited_and_cta_is_enforced,
              test_no_control_characters_anywhere,
              test_report_actions_close_the_loop):
        try:
            t()
        except Exception as e:
            print(f"  NG  {t.__name__} が例外で止まりました: {type(e).__name__} {e}")
            FAIL.append(t.__name__)
    print(f"\n{'失敗 ' + str(len(FAIL)) + '件: ' + ', '.join(FAIL) if FAIL else 'すべて通りました'}")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
