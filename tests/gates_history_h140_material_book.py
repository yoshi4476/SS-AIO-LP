# -*- coding: utf-8 -*-
"""2026-10-31 運用者「今まで出した一次情報をまとめたスプシを社ごとに・過去に図表にした数字も・止まる要素がないほど追加のヒアリングを」。
material_sheet.py --book が社ごとに1冊を作る: 質問のタブ（--make と同じ列と隠しタブ。返せば --apply がそのまま登録する）・
採点の指摘の分類・一次情報の台帳・図表の数字。新しい質問の種類（更新 upd:・語のまとまり kw:・指摘の分類 gap:）を --apply が登録できること。
AI は差し替えて呼ばない。社は架空の client-zz。
"""
import json
import sys
import tempfile
from pathlib import Path

from test_gates import check, ROOT

sys.path.insert(0, str(ROOT / "scripts"))

BODY = """## 費用の目安
| 規模 | 費用 | 期間 |
|---|---|---|
| 10名以下 | 30万円 | 2か月 |
| 50名以下 | 80万円 | 4か月 |

担当した15社で入力が月20時間から6時間になりました。
<img src="/images/zz-a/hiyou.png"><img src="/images/zz-a/nazo.png">
"""
META = {"title": "庭木の剪定の費用", "keyword": "庭木 剪定 費用", "category": "zz",
        "diagrams": [{"name": "hiyou", "type": "bars", "title": "規模ごとの費用", "items": ["小|30", "大|80"], "unit": "万円"},
                     {"name": "nagare", "type": "flow", "title": "流れ", "items": ["聞く", "見積もる"]}]}
FACTS = [{"id": "zz-a", "claim": "2023年4月から2026年9月までに担当した15社で、入力が月20時間から6時間になりました",
          "period": "2023-04〜2026-09", "as_of": "2026-09", "source": "運用者の申告"},
         {"id": "own-x", "claim": "検索の表示は12件でした", "as_of": "2026-10"}]


def _stub(MS):
    import facts as F
    import private_store as P
    old = (MS._articles, MS._todo, MS.weak_of, MS.requests_of, MS.display_name, F.load_for, P.is_client,
           MS._stuck, MS._volumes)
    MS._articles = lambda site: {"zz-a": (META["title"], BODY, META), "zz-b": ("別の記事", "15社で20時間が6時間に", {})}
    MS._todo = lambda site: [{"keyword": "庭木 剪定 費用 相場", "priority": "A"}, {"keyword": "庭木 剪定 時期", "priority": "B"},
                             {"keyword": "生垣 手入れ", "priority": "C"}]
    MS.weak_of = lambda site: [("zz-a", "自社の数字は15社の1つだけで、母数と期間が無い。"), ("zz-a", "費用の実例が無い。")]
    MS.requests_of = lambda site: {"zz-a": {"keyword": "庭木 剪定 費用", "since": "2026-10-20",
                                            "questions": ["「庭木 剪定 費用」そのものについて、当社で数えた数字"], "scorer": ["母数が無い"]}}
    MS._stuck = lambda site, arts, weak, reqs: []
    MS._volumes = lambda: {}
    MS.display_name = lambda site: "テストの社"
    F.load_for = lambda site: ("", FACTS)
    P.is_client = lambda site: site == "client-zz"
    return old


def _unstub(MS, old):
    import facts as F
    import private_store as P
    (MS._articles, MS._todo, MS.weak_of, MS.requests_of, MS.display_name, F.load_for, P.is_client,
     MS._stuck, MS._volumes) = old


def test_book_has_questions_ledger_and_figures():
    import material_sheet as MS
    from openpyxl import load_workbook
    print("\n■ 一冊: 質問は --make と同じ列・隠しタブで優先の高い順。一次情報の台帳と図表の数字（取れない図は「取れなかった」と書く）")
    old = _stub(MS)
    try:
        with tempfile.TemporaryDirectory() as d:
            p, counts = MS.book("client-zz", d, today="2026-10-31")
            wb = load_workbook(p)
            ws = wb[MS.TAB]
            head = [c.value for c in ws[4]]
            check("質問の列は COLS が先頭（--apply が読む A〜H 列を変えない）", head[:len(MS.COLS)], MS.COLS)
            meta = {str(r[0]): str(r[1]) for r in wb[MS.META].iter_rows(values_only=True)}
            rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
            keys = [meta.get(str(r[0]), "") for r in rows]
            tiers = [int(r[len(MS.COLS)]) for r in rows]
            check("隠しタブ・社", (wb[MS.META].sheet_state, meta.get("site")), ("hidden", "client-zz"))
            check("優先の高い順に並ぶ", tiers, sorted(tiers))
            check("止まっている記事の依頼が先頭", keys[0], "req:zz-a:0")
            kinds = {k.split(":")[0] for k in keys}
            check("更新・指摘の分類・語のまとまり・決まった質問がそろう", {"req", "upd", "gap", "kw", "tpl"} <= kinds, True)
            upd = [r for r, k in zip(rows, keys) if k == "upd:zz-a"]
            check("台帳の数字は前回値と時点を添えて更新を聞く（機械の一次情報 own- は聞かない）",
                  (len(upd), "前回値" in str(upd[0][2]) and "15社" in str(upd[0][len(MS.COLS) + 3]), "upd:own-x" in keys),
                  (1, True, False))
            kw = [k for k in keys if k.startswith("kw:")]
            check("語のまとまりごとに名指しで聞く（声は :voice で登録しない側へ・お客様の社は費用も聞く）",
                  (any(k.endswith(":voice") for k in kw), any(k.endswith(":price") for k in kw),
                   all("◯" in str(r[3]) or "そのまま" in str(r[3]) or "登録" in str(r[3]) for r, k in zip(rows, keys) if k.startswith("kw:"))),
                  (True, True, True))
            check("答えの欄は空（数字を作らない）", all(not r[4] for r in rows), True)
            import intake_watch as IW
            check("intake_watch が材料の質問シートと見分ける", IW.kind(p), "material")
            site, answers = MS.read(p)
            check("返ってきても --apply の読み方で読める（空なら答え0件）", (site, answers), ("client-zz", []))
            fx = [r for r in wb["一次情報（事実・数字）"].iter_rows(min_row=2, values_only=True) if r[0]]
            za = [r for r in fx if r[0] == "zz-a"][0]
            check("一次情報は使っている記事を本文から引く", "zz-a" in str(za[10]) and "zz-b" in str(za[10]), True)
            fg = [r for r in wb["図表の数字"].iter_rows(min_row=2, values_only=True) if r[0]]
            got = {(r[2], str(r[6])) for r in fg}
            check("図（数字のある項目）と表の数字を記事ごとに抜く", (("図（bars）", "30") in got, ("表", "30万円") in got), (True, True))
            check("数字の無い図は抜かない・元データの無い画像は「取れなかった」と行に書く",
                  (any(r[3] == "流れ" for r in fg), any("取れなかった" in str(r[9]) and "nazo" in str(r[3]) for r in fg)), (False, True))
            check("件数を返す", (counts["質問"] == len(rows), counts["図表の数字"] == len(fg)), (True, True))
    finally:
        _unstub(MS, old)


def test_apply_registers_update_rows_in_place():
    import material_sheet as MS
    print("\n■ 更新（upd:）の答え: 前回値と違っても食い違いにせず、同じ id の文を新しい時点で書き換える")
    ans = {"no": "3", "key": "upd:zz-a", "item": "一次情報の更新",
           "question": f"前回値：15社・20時間・6時間（時点：2026-09）。今の文: 「{FACTS[0]['claim']}」",
           "answer": "2026年10月時点で18社、月20時間から5時間", "period": "", "count": "18社", "ok": "可"}
    reply = json.dumps({"facts": [{"no": "3", "claim": "2023年4月から2026年10月までに担当した18社で、入力が月20時間から5時間になりました"}],
                        "experiences": [], "conflicts": []}, ensure_ascii=False)
    old = MS.existing
    MS.existing = lambda s: ([("zz-a", FACTS[0]["claim"])], [])
    try:
        f, e, held = MS.plan("ai-lab", [ans], ask=lambda prompt: reply)
    finally:
        MS.existing = old
    check("更新は replaces と updated が付き、保留にならない", ([(x.get("replaces"), x.get("updated")) for x in f], held),
          ([("zz-a", True)], []))
    check("答えに無い数字は更新でも止める", MS.invented("18社で3時間", ans), ["3"])
    with tempfile.TemporaryDirectory() as d:
        own_f, own_e = MS.OWN_FACTS, MS.OWN_EXPS
        MS.OWN_FACTS, MS.OWN_EXPS = Path(d) / "f.json", Path(d) / "e.json"
        MS.OWN_FACTS.write_text(json.dumps({"facts": [dict(FACTS[0])]}, ensure_ascii=False), encoding="utf-8")
        MS.OWN_EXPS.write_text(json.dumps({"experiences": []}), encoding="utf-8")
        try:
            MS.write("ai-lab", f, [], today="2026-10-31")
            got = json.loads(MS.OWN_FACTS.read_text(encoding="utf-8"))["facts"]
        finally:
            MS.OWN_FACTS, MS.OWN_EXPS = own_f, own_e
    check("同じ id を書き換え、時点を新しくする（同じ事実を2つ並べない）",
          ([g["id"] for g in got], got[0]["as_of"], "18社" in got[0]["claim"], "更新" in got[0]["source"]),
          (["zz-a"], "2026-10", True, True))
    src = (ROOT / "scripts" / "material_sheet.py").read_text(encoding="utf-8")
    check("CI では一冊を作らない（お客様の語をログ・成果物に出さない）", "MATERIAL_BOOK=skip" in src, True)
