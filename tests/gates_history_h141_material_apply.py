# -*- coding: utf-8 -*-
"""2026-10-31 運用者が記入して返した3冊（material_sheet --book）を --apply で読んだら、登録の判定が甘く・数え方が誤っていた。
1. 「記事に出してよいか」が「不可（根拠確認後に判断）」のような条件つきで書かれ、完全一致の判定を「不可（…）」がすり抜けた
2. 答え406行を claude の1回に詰め、応答が壊れて {} になったのを「登録0件」と報告した
3. 同じ集計の新しい期間の計測（期間が後・母数が大きい）を、食い違いとして保留にしていた
4. 作ったときの行（174行）に、記入した側が行を足していた（406行）。どの行がどの質問の答えかを確かめずに読んでいた
5. 今の台帳の言い直し（同じ数字の文）を新しい一次情報として登録していた
AI は差し替えて呼ばない。社は架空の client-zz・zz-*。
"""
import json
import sys
import tempfile
from pathlib import Path

from test_gates import check, ROOT

sys.path.insert(0, str(ROOT / "scripts"))

CTR = {"id": "zz-ctr", "claim": "2026-06-23〜2026-09-20に自社3サイトで計測した7015回（表示）をもとに集計した検索順位ごとのクリック率",
       "period": "2026-06-23〜2026-09-20", "as_of": "2026-09"}
HOURS = {"id": "zz-hours", "claim": "2023年4月から2026年9月までに担当した15社で、入力が月20時間から6時間になりました",
         "period": "2023-04〜2026-09", "as_of": "2026-10"}


def _ans(no, answer, ok="可", key=None, item="事例カード", question="事例を1件", period="", count="", target=""):
    return {"no": str(no), "key": key or "tpl:case1", "item": item, "question": question, "answer": answer,
            "period": period, "count": count, "ok": ok, "target": target, "bind": "", "added": False}


def _stub(MS, recs=(CTR, HOURS)):
    old = MS.existing, MS.records
    MS.existing = lambda s: ([(r["id"], r["claim"]) for r in recs], [])
    MS.records = lambda s: {r["id"]: dict(r) for r in recs}
    return old


def _unstub(MS, old):
    MS.existing, MS.records = old


def test_conditional_permission():
    import facts as F
    import material_sheet as MS
    import private_store as P
    print("\n■ 1. 「記事に出してよいか」の条件つきの書き方: 不可・要確認・仮置きは登録しない／「可（条件）」は条件を使ってよい範囲として残す")
    old_client = P.is_client
    P.is_client = lambda s: s == "client-zz"
    try:
        got = {(ok, site): MS.permission(ok, site) for ok, site in (
            ("不可（根拠確認後に判断）", "ai-lab"), ("要確認（出典を示せない）", "ai-lab"), ("不可（仮置き。実績としての内訳は未確定）", "ai-lab"),
            ("可（当社によるAI回答調査としてのみ。本文の事実は別途裏付け）", "ai-lab"), ("可（既存公開部分のみ。要確認事項は除外）", "ai-lab"),
            ("可（行った修正の事実のみ。効果は未確認）", "ai-lab"), ("可", "ai-lab"), ("", "ai-lab"),
            ("お客様の記事だけ（社名は伏せる）", "client-zz"), ("お客様の記事だけ（社名は伏せる）", "ai-lab"))}
    finally:
        P.is_client = old_client
    kinds = {k[0] if k[1] == "ai-lab" else k: (v[0], v[2].split("：")[0] if v[2] else "") for k, v in got.items()}
    check("悪い例: 不可・要確認・仮置き（完全一致でなくても）は登録しない",
          [kinds[k] for k in ("不可（根拠確認後に判断）", "要確認（出典を示せない）", "不可（仮置き。実績としての内訳は未確定）")],
          [(False, "不可")] * 3)
    check("悪い例: 「可」でも条件に未確認・要確認が入るものは保留",
          [kinds[k] for k in ("可（既存公開部分のみ。要確認事項は除外）", "可（行った修正の事実のみ。効果は未確認）")],
          [(False, "条件が未確認")] * 2)
    check("悪い例: 空欄・自社のシートの「お客様の記事だけ」は登録しない",
          (kinds[""], kinds["お客様の記事だけ（社名は伏せる）"][0]), ((False, "不可"), False))
    check("良い例: 「可（条件）」は登録し、条件を使ってよい範囲に残す／「可」だけは範囲なし／お客様の社の「お客様の記事だけ」は登録",
          (got[("可（当社によるAI回答調査としてのみ。本文の事実は別途裏付け）", "ai-lab")][:2], got[("可", "ai-lab")][:2],
           got[("お客様の記事だけ（社名は伏せる）", "client-zz")][:2]),
          ((True, "当社によるAI回答調査としてのみ。本文の事実は別途裏付け"), (True, ""), (True, "お客様の記事だけ（社名は伏せる）")))
    old = _stub(MS)
    ans = [_ans(1, "2025年3月に卸売の会社で入力が月40時間から8時間になった", ok="可（当社によるAI回答調査としてのみ）"),
           _ans(2, "2025年4月に飲食店で予約が月10件から30件", ok="不可（根拠確認後に判断）")]
    reply = json.dumps({"facts": [{"no": "1", "claim": "2025年3月に卸売の会社で、入力が月40時間から8時間になりました"},
                                  {"no": "2", "claim": "2025年4月に飲食店で予約が月10件から30件になりました"}]}, ensure_ascii=False)
    try:
        f, e, held = MS.plan("ai-lab", ans, ask=lambda p: reply)
    finally:
        _unstub(MS, old)
    check("不可（…）の答えは文にしても登録しない・可（条件）は scope が付く",
          ([(x["no"], x.get("scope")) for x in f], [n for n, w in held if "記事に出さない" in w]),
          ([("1", "当社によるAI回答調査としてのみ")], ["2"]))
    with tempfile.TemporaryDirectory() as d:
        own_f, own_e = MS.OWN_FACTS, MS.OWN_EXPS
        MS.OWN_FACTS, MS.OWN_EXPS = Path(d) / "f.json", Path(d) / "e.json"
        MS.OWN_FACTS.write_text(json.dumps({"facts": []}), encoding="utf-8")
        MS.OWN_EXPS.write_text(json.dumps({"experiences": []}), encoding="utf-8")
        try:
            MS.write("ai-lab", f, [], today="2026-10-31")
            got_f = json.loads(MS.OWN_FACTS.read_text(encoding="utf-8"))["facts"]
        finally:
            MS.OWN_FACTS, MS.OWN_EXPS = own_f, own_e
    check("台帳に使ってよい範囲（scope）が残り、書く側の材料に出る",
          (got_f[0].get("scope"), "当社によるAI回答調査としてのみ" in F.scope_note(got_f[0]), F.scope_note({"claim": "x"})),
          ("当社によるAI回答調査としてのみ", True, ""))
    src = (ROOT / "scripts" / "facts.py").read_text(encoding="utf-8") + (ROOT / "scripts" / "quality_loop.py").read_text(encoding="utf-8")
    check("書く工程（facts.py の一覧）と直す工程（quality_loop.materials）の両方が範囲を渡す", src.count("scope_note(f)}"), 2)


def test_chunks_and_broken_reply_are_not_zero():
    import material_sheet as MS
    print("\n■ 2. 答えをまとまりに分けて claude を呼ぶ。応答が空・壊れたまとまりは「分からない」と件数で知らせ、0件と言わない")
    ans = [_ans(i, f"2025年{i % 12 + 1}月に{i}社で入力が{i + 100}時間減った") for i in range(1, 66)]
    calls = []

    def ask(prompt):
        calls.append(prompt)
        nos = [n for n in range(1, 66) if f"[{n}] " in prompt]
        if nos[0] == 1:
            return json.dumps({"facts": [{"no": str(n), "claim": f"2025年{n % 12 + 1}月に{n}社で入力が{n + 100}時間減りました"} for n in nos]},
                              ensure_ascii=False)
        return "" if nos[0] <= 31 else '{"facts": [{"no": "61", "claim": "途中で切れ'
    old = _stub(MS, ())
    try:
        f, e, held = MS.plan("ai-lab", ans, ask=ask)
    finally:
        _unstub(MS, old)
    sizes = [sum(1 for n in range(1, 66) if f"[{n}] " in p) for p in calls]
    check("1回に渡す答えは CHUNK 件まで（悪い例: 65件を1回に詰める）", (max(sizes) <= MS.CHUNK, MS.CHUNK <= 30), (True, True))
    check("空・壊れた応答はまとまりごとに1回だけ呼び直す（3まとまりで5回）", len(calls), 5)
    unknown = [n for n, w in held if w.startswith("分からない")]
    check("登録できたのは応答のあったまとまりだけ・残りは「分からない」として保留", (len(f), len(unknown)), (30, 35))
    old, old_read = _stub(MS, ()), MS.read
    MS.read = lambda src: ("ai-lab", ans)
    try:
        ok, msg = MS.apply(Path("x.xlsx"), False, ask=lambda prompt: "")
    finally:
        _unstub(MS, old)
        MS.read = old_read
    check("全部の応答が空なら「登録0件」ではなく「分からない」と件数で言う", ("分からない 65件" in msg, ok), (True, False))


def test_newer_period_is_update_not_conflict():
    import material_sheet as MS
    print("\n■ 3. 同じ集計の新しい期間の計測は「更新」。古い値を消さず、期間つきの新しい値を足す。対象が曖昧なものは保留のまま")
    ans = [_ans(1, "2026-07-26〜2026-10-23の自社3サイトで表示19,798回・クリック84回", period="2026-07-26〜2026-10-23"),
           _ans(2, "2026-06-23〜2026-09-20の自社3サイトで表示7,100回", period="2026-06-23〜2026-09-20"),
           _ans(3, "支援先は約20社", period="")]
    reply = json.dumps({"updates": [
        {"no": "1", "existing": "zz-ctr", "claim": "2026-07-26〜2026-10-23に自社3サイトで計測した表示19,798回・クリック84回",
         "period": "2026-07-26〜2026-10-23"},
        {"no": "2", "existing": "zz-ctr", "claim": "2026-06-23〜2026-09-20に自社3サイトで計測した表示7,100回", "period": "2026-06-23〜2026-09-20"}],
        "conflicts": [{"no": "3", "existing": "zz-hours", "why": "「15社」と「約20社」が同じ対象か分からない"}]}, ensure_ascii=False)
    old = _stub(MS)
    try:
        f, e, held = MS.plan("ai-lab", ans, ask=lambda p: reply)
    finally:
        _unstub(MS, old)
    why = dict(held)
    check("良い例: 期間が後の同じ集計は更新（follows に元の id）", [(x["no"], x.get("follows")) for x in f], [("1", "zz-ctr")])
    check("悪い例: 期間が新しくない「更新」は食い違いとして保留・対象が曖昧なものも保留",
          (why.get("2", "").startswith("食い違い"), why.get("3", "").startswith("食い違い")), (True, True))
    with tempfile.TemporaryDirectory() as d:
        own_f, own_e = MS.OWN_FACTS, MS.OWN_EXPS
        MS.OWN_FACTS, MS.OWN_EXPS = Path(d) / "f.json", Path(d) / "e.json"
        MS.OWN_FACTS.write_text(json.dumps({"facts": [dict(CTR)]}, ensure_ascii=False), encoding="utf-8")
        MS.OWN_EXPS.write_text(json.dumps({"experiences": []}), encoding="utf-8")
        try:
            MS.write("ai-lab", f, [], today="2026-10-31")
            got = json.loads(MS.OWN_FACTS.read_text(encoding="utf-8"))["facts"]
        finally:
            MS.OWN_FACTS, MS.OWN_EXPS = own_f, own_e
    check("古い値は消さず（文はそのまま）、新しい値を期間つきで足し、古い方に新しい id を記す",
          (len(got), got[0]["claim"] == CTR["claim"], got[1].get("follows"), got[1].get("period"), got[0].get("newer") == got[1]["id"]),
          (2, True, "zz-ctr", "2026-07-26〜2026-10-23", True))
    reply2 = json.dumps({"updates": [json.loads(reply)["updates"][0]],
                         "conflicts": [{"no": "2", "existing": "zz-ctr", "why": "同じ数え方か分からない"}]}, ensure_ascii=False)
    old = _stub(MS)
    try:
        f2, _e, held2 = MS.plan("ai-lab", ans, ask=lambda p: reply2)
    finally:
        _unstub(MS, old)
    ds = dict(CTR, id="dataset-zz-ctr")
    reply3 = json.dumps({"updates": [dict(json.loads(reply)["updates"][0], existing="dataset-zz-ctr")]}, ensure_ascii=False)
    old = _stub(MS, (ds, HOURS))
    try:
        f3, _e, held3 = MS.plan("ai-lab", ans[:1], ask=lambda p: reply3)
    finally:
        _unstub(MS, old)
    check("機械で作る記録（dataset-）の新しい期間は台帳に足さず「更新の候補」として一覧（公開データの工程が数え直す）",
          (f3, dict(held3).get("1", "")[:5]), ([], "更新の候補"))
    check("悪い例: 同じ記録が別の答えで食い違いと読まれたら、その記録の更新も保留（迷うものは登録しない）",
          (f2, dict(held2).get("1", "")[:4]), ([], "食い違い"))


def test_rows_tie_to_their_questions():
    import material_sheet as MS
    from openpyxl import load_workbook
    print("\n■ 4. 行と隠しタブの印の対応を確かめる。足された行はその行の質問への答えとして読み、印の無い行・種類の合わない行は保留")
    old = MS.requests_of, MS.display_name, MS.fact_gaps
    MS.fact_gaps = lambda *a, **k: []
    MS.requests_of, MS.display_name = (lambda s: {}), (lambda s: "テストの社")
    try:
        with tempfile.TemporaryDirectory() as d:
            p, n = MS.make("ai-lab", d, today="2026-10-31")
            wb = load_workbook(p)
            ws, meta = wb[MS.TAB], wb[MS.META]
            for row in ws.iter_rows(min_row=5, max_row=6):
                row[4].value, row[7].value = "2025年3月に3社で入力が月40時間から8時間", "可"
            ws.append([n + 1, "追加質問：入力時間", "3社の平均は何時間減ったか", "", "3社で平均32時間減", "2025年", "3社", "可"])
            ws.append([n + 2, "追加質問：印なし", "どこの数字か", "", "5社", "", "", "可"])
            meta.append([n + 1, "extra:0-1"])
            for r in meta.iter_rows():
                if str(r[0].value) == "2":
                    r[1].value = "upd:zz-hours"          # 2行目は「事例カード」なのに「一次情報の更新」の印（ずれ）
            wb.save(p)
            site, answers = MS.read(p)
    finally:
        MS.requests_of, MS.display_name, MS.fact_gaps = old
    by = {a["no"]: a for a in answers}
    check("作ったときの行は印どおり（結びつく）", (by["1"]["bind"], by["1"]["added"]), ("", False))
    check("悪い例: 項目と印の種類が合わない行・印の無い行は「結びつかない行」",
          (by["2"]["bind"].startswith("結びつかない行"), by[str(n + 2)]["bind"].startswith("結びつかない行")), (True, True))
    check("良い例: 足された行は印つきなら、その行の質問への答えとして読む", (by[str(n + 1)]["bind"], by[str(n + 1)]["added"]), ("", True))
    seen = []
    old = _stub(MS, ())
    try:
        f, e, held = MS.plan(site, answers, ask=lambda pr: seen.append(pr) or json.dumps({"facts": []}))
    finally:
        _unstub(MS, old)
    check("結びつかない行は claude に渡さず保留・足された行は質問ごと渡す",
          (sorted(nn for nn, w in held if w.startswith("結びつかない行")), "3社の平均は何時間減ったか" in seen[0], "[2] " in seen[0]),
          (sorted(["2", str(n + 2)]), True, False))


def test_restatement_is_duplicate():
    import material_sheet as MS
    print("\n■ 5. 今の台帳の言い直し（同じ数字・期間の文）は重複として登録しない。新しい数字のある文は登録する")
    ans = [_ans(1, "2023年4月から2026年9月までに担当した15社で入力が月20時間から6時間［一次情報!C3］"),
           _ans(2, "2023年4月から2026年9月までの15社のうち最小は月3時間・最大は月10時間"),
           _ans(3, "変わらず", key="upd:zz-hours", item="一次情報の更新（記事3本で使用）",
                question=f"前回値：15社・20時間・6時間。今の文: 「{HOURS['claim']}」")]
    reply = json.dumps({"facts": [
        {"no": "1", "claim": "2023年4月から2026年9月までに担当した15社では、入力が月20時間から6時間に減りました"},
        {"no": "2", "claim": "2023年4月から2026年9月までに担当した15社のうち、入力時間が最も短い会社は月3時間、最も長い会社は月10時間でした"},
        {"no": "3", "claim": HOURS["claim"]}]}, ensure_ascii=False)
    old = _stub(MS)
    try:
        f, e, held = MS.plan("ai-lab", ans, ask=lambda p: reply)
    finally:
        _unstub(MS, old)
    why = dict(held)
    check("悪い例: 同じ数字・期間の言い直しと「変わらず」の更新は重複", (why.get("1", "")[:2], why.get("3", "")[:2]), ("重複", "重複"))
    check("良い例: 新しい数字（最小・最大）のある文は登録する", [x["no"] for x in f], ["2"])
    ans = [_ans(1, "AIに聞いた10件の答えのうち3件が費用に触れた［当社調べ（AIへの質問）!A7］"),
           _ans(2, "2026年10月に当社がAIへ聞いた答え10件のうち費用に触れたのは3件", ok="可（当社によるAI回答調査としてのみ）"),
           _ans(3, "2026年10月に当社がAIへ聞いた答え10件のうち費用に触れたのは3件", ok="可（当社によるAI回答調査としてのみ）")]
    survey = "2026年10月に当社がAIへ聞いた答え10件のうち、費用の金額に触れたのは3件でした"
    reply = json.dumps({"facts": [{"no": n, "claim": survey} for n in ("1", "2", "3")]}, ensure_ascii=False)
    old = _stub(MS, (CTR, HOURS, {"id": "zz-many", "claim": "2026年10月までの10件のうち3件・1件・2件・6件・8件が1か月以内に上がった"}))
    try:
        f, e, held = MS.plan("ai-lab", ans, ask=lambda p: reply)
    finally:
        _unstub(MS, old)
    why = dict(held)
    check("悪い例: 1冊の機械の台帳（当社調べ・図表など）の書き写しは重複／同じ文の2つ目は重複",
          (why.get("1", "")[:2], "当社調べ" in why.get("1", ""), why.get("3", "")[:2]), ("重複", True, "重複"))
    check("良い例: 小さな数字（10件・3件）だけの文を、数字の多い別の記録（zz-many）の言い直しにしない", [x["no"] for x in f], ["2"])
    a = dict(_ans(1, "15社で20時間から6時間は14時間減", item="追加質問：減った時間", key="extra:0-1", target="zz-hours"), added=True)
    reply = json.dumps({"facts": [{"no": "1", "claim": "担当した15社で、入力が月20時間から6時間になり、14時間減りました"}]},
                       ensure_ascii=False)
    old = _stub(MS)
    try:
        f, e, held = MS.plan("ai-lab", [a], ask=lambda p: reply)
    finally:
        _unstub(MS, old)
    check("悪い例: 足された行で記録（zz-hours）を計算し直しただけの文は重複", (f, dict(held).get("1", "")[:2]), ([], "重複"))
