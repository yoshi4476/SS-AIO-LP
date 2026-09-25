# -*- coding: utf-8 -*-
"""実績とお客様の声を、記入シート1枚から登録して掲載する。

数字は本人にしか出せない。**機械が作ることは絶対にない**（架空の実績は優良誤認）。
シートに入れた数字だけを add_fact と同じ条件（割合には母数と期間・10件未満は実数）で
検査し、通ったものを
  - 一次情報 data/first_party_facts.json   … 次に書かれる記事から使われる
  - お客様の声 data/voices.json            … 運営者情報（/about/）と LP（/lp/）に掲載
に登録する。

  python scripts/jisseki_intake.py --sheet             # 記入シートを作る（intake/実績記入シート.xlsx）
  python scripts/jisseki_intake.py <xlsx>              # 検査だけ
  python scripts/jisseki_intake.py <xlsx> --apply      # 登録して掲載

intake/ に置いたまま `intake_watch.py --apply` でも拾う（ファイル名に「実績」）。
シートは会社名・担当者名が入るためコミットしない（intake/*.xlsx は .gitignore 済み）。
"""
import argparse
import html
import json
import re
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
IN = ROOT / "intake"
SHEET = IN / "実績記入シート.xlsx"
VOICES = ROOT / "data" / "voices.json"
ABOUT = ROOT / "site" / "about" / "index.html"
LP = ROOT / "site" / "lp" / "index.html"
SITES = ("ai-lab", "corporate", "subsidy")
START, END = "<!-- voices:start -->", "<!-- voices:end -->"

RULES = [
    "この1枚に入れた数字だけが、記事と運営者情報・LPに載ります。機械は数字を作りません。",
    "割合（◯%・◯割）を載せるには「何社中何社か（母数）」と「いつからいつまで（期間）」が必要です。",
    "母数が10件未満なら割合にはせず、実数（◯社のうち◯社）で載ります。",
    "お客様の声は「掲載可否」が「可」のものだけ載ります。「可」にできるのは、本人が書いた文章か、本人が内容を確認して掲載を許可したものだけです（事業者が作った感想の表示はステルスマーケティング規制の違反）。会社名を出さない場合は表示名（例: 大阪市の歯科医院）を書いてください。",
    "数字を添える場合は「前」「後」「内容（例: 月の問い合わせ件数）」「期間」の4つをそろえてください。",
    "空欄の項目は登録しません。分からないものは空欄のままで構いません。",
]
BPO_ROWS = [
    ("導入社数", "", "例: 12（経理BPOを導入した社数）"),
    ("集計期間（開始）", "", "例: 2025-01"),
    ("集計期間（終了）", "", "例: 2026-08"),
    ("月次の締めが短くなった営業日数（平均）", "", "例: 3（導入前と比べて）"),
    ("掲載サイト", "corporate", "corporate / subsidy / ai-lab をカンマ区切り"),
]
KEIZOKU_ROWS = [
    ("契約社数", "", "例: 20（期間内に契約した社数）"),
    ("継続社数", "", "例: 18（そのうち今も継続している社数）"),
    ("期間（開始）", "", "例: 2025-04"),
    ("期間（終了）", "", "例: 2026-03"),
    ("掲載サイト", "corporate,subsidy,ai-lab", "corporate / subsidy / ai-lab をカンマ区切り"),
]
VOICE_COLS = ["掲載可否", "会社名", "表示名（会社名を出さない場合）", "業種", "お名前・役職", "一言（240字まで）",
              "数字（前）", "数字（後）", "数字の内容", "期間", "掲載サイト"]
VOICE_EXAMPLE = ["例", "○○歯科医院", "大阪市の歯科医院", "歯科", "院長", "口コミ返信を任せてから新患の予約が増えました",
                 "3", "11", "月の問い合わせ件数", "2026-01〜2026-06", "ai-lab"]


def make_sheet(path=SHEET):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", start_color="0B2447")
    ws = wb.active
    ws.title = "書き方"
    ws["A1"] = "実績・お客様の声 記入シート"; ws["A1"].font = Font(bold=True, size=14)
    for i, r in enumerate(RULES, 3):
        ws.cell(row=i, column=1, value=f"{i - 2}. {r}")
    ws.column_dimensions["A"].width = 110
    for name, rows in (("経理BPOの効果", BPO_ROWS), ("継続率", KEIZOKU_ROWS)):
        w = wb.create_sheet(name)
        w.append(["項目", "記入", "例・注意"])
        for c in w[1]:
            c.font, c.fill = head, fill
        for r in rows:
            w.append(list(r))
        w.column_dimensions["A"].width = 40; w.column_dimensions["B"].width = 28; w.column_dimensions["C"].width = 48
        for row in w.iter_rows(min_row=2, min_col=2, max_col=2):
            row[0].fill = PatternFill("solid", start_color="FFF9DB")
    w = wb.create_sheet("お客様の声")
    w.append(VOICE_COLS)
    for c in w[1]:
        c.font, c.fill = head, fill
        c.alignment = Alignment(wrap_text=True)
    w.append(VOICE_EXAMPLE)
    for c in w[2]:
        c.font = Font(color="888888", italic=True)
    for i in range(3):
        w.append(["可"] + [""] * (len(VOICE_COLS) - 1))
    for col, wd in zip("ABCDEFGHIJK", (10, 18, 24, 12, 16, 48, 10, 10, 18, 18, 14)):
        w.column_dimensions[col].width = wd
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _kv(ws):
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            out[str(row[0]).strip()] = ("" if row[1] is None else str(row[1]).strip())
    return out


def read(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    got = {"bpo": {}, "keizoku": {}, "voices": []}
    if "経理BPOの効果" in wb.sheetnames:
        got["bpo"] = _kv(wb["経理BPOの効果"])
    if "継続率" in wb.sheetnames:
        got["keizoku"] = _kv(wb["継続率"])
    if "お客様の声" in wb.sheetnames:
        for row in wb["お客様の声"].iter_rows(min_row=2, values_only=True):
            vals = ["" if v is None else str(v).strip() for v in (list(row) + [""] * 11)[:11]]
            if vals[0] in ("", "例") or not any(vals[1:]):
                continue
            got["voices"].append(dict(zip(("ok", "company", "display", "industry", "person", "quote",
                                           "before", "after", "metric", "period", "sites"), vals)))
    return got


def _sites(s, default):
    v = [x.strip() for x in re.split(r"[,、\s/]+", s or "") if x.strip()] or list(default)
    return v


def _int(s):
    try:
        return int(float(str(s).replace(",", "")))
    except Exception:
        return None


def _ym(s):
    s = str(s or "").strip().replace("/", "-").replace("年", "-").replace("月", "")
    m = re.match(r"^(\d{4})-(\d{1,2})$", s)
    return f"{m.group(1)}-{int(m.group(2)):02d}" if m else ""


PREFIXES = ("keizoku-", "keiri-bpo-effect-")


def _prefix(fid):
    """シートから作る一次情報の接頭辞。月が変わると id も変わるので、接頭辞で同じものとみなす"""
    return next((p for p in PREFIXES if str(fid).startswith(p)), "")


def review(got):
    """検査する。(登録する一次情報, 載せる声, 不備, 警告) を返す"""
    import add_fact
    facts, voices, ng, warn = [], [], [], []
    today = date.today().strftime("%Y-%m")

    b = got.get("bpo") or {}
    n, days = _int(b.get("導入社数")), b.get("月次の締めが短くなった営業日数（平均）")
    if n or days:
        s, e = _ym(b.get("集計期間（開始）")), _ym(b.get("集計期間（終了）"))
        try:
            days_f = float(str(days).replace(",", ""))
        except Exception:
            days_f = None
        if not n or not days_f or not s or not e:
            ng.append("経理BPOの効果: 導入社数・営業日数・期間（開始/終了）の4つすべてが要ります")
        else:
            if n < 5:
                warn.append(f"経理BPOの効果: 母数が{n}社と少なく、平均の説得力が弱いです（載せますが実数の併記を推奨）")
            d = f"{days_f:g}"
            facts.append({"id": f"keiri-bpo-effect-{today.replace('-', '')}",
                          "sites": _sites(b.get("掲載サイト"), ["corporate"]),
                          "topic": ["経理BPO", "月次決算", "業務効率化", "バックオフィス"],
                          "claim": f"{s}〜{e}に経理BPOを導入した{n}社では、月次決算の締めが平均{d}営業日短くなりました",
                          "source": "自社の支援実績（導入先の月次決算の締め日を導入前後で比較）",
                          "as_of": today, "denominator": n, "period": f"{s}〜{e}", "verifiable": True,
                          "_pending": "削減時間"})

    k = got.get("keizoku") or {}
    n, m = _int(k.get("契約社数")), _int(k.get("継続社数"))
    if n or m:
        s, e = _ym(k.get("期間（開始）")), _ym(k.get("期間（終了）"))
        if not n or m is None or not s or not e:
            ng.append("継続率: 契約社数・継続社数・期間（開始/終了）の4つすべてが要ります")
        elif m > n:
            ng.append(f"継続率: 継続社数（{m}）が契約社数（{n}）を超えています")
        else:
            if n >= 10:
                claim = f"{s}〜{e}に契約した{n}社のうち{m}社が継続しています（継続率{m / n * 100:.1f}%）"
            else:
                claim = f"{s}〜{e}に契約した{n}社のうち{m}社が継続しています"
                warn.append(f"継続率: 母数が{n}社のため割合にせず実数で載せます")
            facts.append({"id": f"keizoku-{today.replace('-', '')}",
                          "sites": _sites(k.get("掲載サイト"), SITES),
                          "topic": ["継続率", "支援実績", "顧客満足"], "claim": claim,
                          "source": "自社の契約実績", "as_of": today, "denominator": n,
                          "period": f"{s}〜{e}", "verifiable": True, "_pending": "継続率"})

    for i, v in enumerate(got.get("voices") or [], 1):
        if v["ok"] != "可":
            warn.append(f"お客様の声{i}: 掲載可否が「可」でないため載せません")
            continue
        who = v["display"] or v["company"]
        if not who:
            ng.append(f"お客様の声{i}: 会社名か表示名（例: 大阪市の歯科医院）のどちらかが要ります")
            continue
        if not v["industry"] or not v["quote"]:
            ng.append(f"お客様の声{i}: 業種と一言は必須です")
            continue
        if len(v["quote"]) > 240:
            ng.append(f"お客様の声{i}: 一言が{len(v['quote'])}字です（240字まで）")
            continue
        num = None
        if v["before"] or v["after"] or v["metric"]:
            if not (v["before"] and v["after"] and v["metric"] and v["period"]):
                ng.append(f"お客様の声{i}: 数字を載せるなら 前・後・内容・期間 の4つをそろえてください")
                continue
            num = {"before": v["before"], "after": v["after"], "metric": v["metric"], "period": v["period"]}
        voices.append({"who": who, "industry": v["industry"], "person": v["person"], "quote": v["quote"],
                       "number": num, "sites": _sites(v["sites"], ["ai-lab"]), "as_of": today})

    # 置き直しは差し替えになる（apply が同じ接頭辞の既存分を外す）ので、既存の id とは照合しない
    renew = {_prefix(f["id"]) for f in facts}
    ids = {f["id"] for f in add_fact.load()["facts"] if _prefix(f["id"]) not in renew}
    for f in facts:
        for p in add_fact.problems({k_: v_ for k_, v_ in f.items() if not k_.startswith("_")}, ids):
            ng.append(f"{f['id']}: {p}")
    if not facts and not voices and not ng and not got.get("voices"):
        ng.append("何も記入されていません（空欄のシートです）")
    return facts, voices, ng, warn


def render(voices, style="lp"):
    """お客様の声のHTML。数字を先頭に置き、引用は読み物として下に置く。

    以前は長い引用が先頭にあり、6件並ぶと何が書いてあるか掴めなかった。
    読む人が最初に知りたいのは「何がどう変わったか」なので、その一行を上に出す。
    """
    e = html.escape
    cards = []
    for v in voices:
        n = v.get("number") or {}
        metric = ""
        if n:
            metric = (f'<p class="voice-metric"><span class="m-label">{e(n["metric"])}</span>'
                      f'<span class="m-from">{e(n["before"])}</span>'
                      f'<span class="m-arw" aria-hidden="true">→</span>'
                      f'<span class="m-to">{e(n["after"])}</span></p>')
        # 表示名に業種が入っていることが多い。そのまま足すと「製造業（従業員約50名）（製造業）」になる
        who = v["who"]
        if v.get("industry") and v["industry"] not in who:
            who = f'{who}（{v["industry"]}）'
        if v.get("person"):
            who = f'{who}／{v["person"]}'
        term = f'<span class="voice-term">{e(n["period"])}</span>' if n.get("period") else ""
        cards.append(f'<figure class="voice">{metric}'
                     f'<blockquote>{e(v["quote"])}</blockquote>'
                     f'<figcaption>{e(who)}{term}</figcaption></figure>')
    note = ('<p class="voice-note">掲載はご本人の許可を得たものだけです。'
            '数字は各社の集計期間を添えています。個別の成果を保証するものではありません。</p>')
    grid = f'<div class="voice-grid">{"".join(cards)}</div>'
    if style == "about":
        body = f"<h2>お客様の声</h2>\n  {grid}\n  {note}"
    else:
        body = ('<section class="section" data-area="お客様の声" data-area-id="voices">\n'
                '  <div class="section-head"><span class="en">Voice</span><h2>お客様の声</h2>\n'
                '  <p class="section-lead">実際に運用させていただいている会社さまの言葉です。'
                '数字は各社が測った値で、集計期間を添えています。</p></div>\n'
                f"  {grid}\n  {note}\n</section>")
    return f"{START}\n{body}\n{END}"


def place(html, block, anchor):
    """印の間を置き換える。無ければ anchor の直前に入れる"""
    if START in html and END in html:
        return html[: html.index(START)] + block + html[html.index(END) + len(END):]
    if anchor in html:
        return html.replace(anchor, block + "\n\n" + anchor, 1)
    return html


def apply(facts, voices):
    import add_fact
    d = add_fact.load()
    # 置き直すたびに継続率・効果が1件ずつ増え、古い数字と新しい数字が並んで記事に使われていた。
    # 同じ接頭辞の既存分を外してから入れる
    renew = {_prefix(f["id"]) for f in facts} - {""}
    d["facts"] = [f for f in d["facts"] if _prefix(f.get("id")) not in renew]
    for f in facts:
        key = f.pop("_pending", "")
        d["facts"].append(f)
        if key:
            d["pending"] = [p for p in d["pending"] if key not in p.get("note", "")]
    if facts:
        add_fact.SRC.write_text(json.dumps(d, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    cur = json.loads(VOICES.read_text(encoding="utf-8")) if VOICES.is_file() else []
    # 同じ人の同じ言葉は1件。as_of（登録月）や数字の欄が変わると辞書として一致せず、
    # 置き直すたびに二重掲載になっていた
    new = {(v.get("who"), v.get("quote")) for v in voices}
    cur = [v for v in cur if (v.get("who"), v.get("quote")) not in new] + voices
    VOICES.parent.mkdir(parents=True, exist_ok=True)
    VOICES.write_text(json.dumps(cur, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    shown = [v for v in cur if "ai-lab" in v.get("sites", [])]
    touched = []
    if shown:
        for path, style, anchor in ((ABOUT, "about", '  <section class="cta">'), (LP, "lp", "<!-- 代表メッセージ -->")):
            if path.is_file():
                h = path.read_text(encoding="utf-8")
                u = place(h, render(shown, style), anchor)
                if u != h:
                    path.write_text(u, encoding="utf-8", newline="\n")
                    touched.append(path.relative_to(ROOT).as_posix())
    return touched


def one(src, write):
    """intake_watch から呼ばれる。(登録できたか, 説明)"""
    import intake_watch as W
    try:
        facts, voices, ng, warn = review(read(src))
    except Exception as e:
        return False, f"読めません（{str(e)[:80]}）"
    if ng:
        note = ("このシートは登録していません。次を直して intake/ に戻してください。\n\n"
                + "\n".join("  - " + x for x in ng)
                + ("\n\n（警告）\n" + "\n".join("  - " + x for x in warn) if warn else "") + "\n")
        if write:
            W.move(src, W.TODO, note)
        return False, f"{src.name}: 不備{len(ng)}件（intake/todo/ へ移しました）"
    if not write:
        return True, f"{src.name}: 一次情報{len(facts)}件・お客様の声{len(voices)}件を登録できます（--apply で実行）"
    touched = apply(facts, voices)
    W.move(src, W.DONE)
    return True, (f"{src.name}: 一次情報{len(facts)}件・お客様の声{len(voices)}件を登録しました"
                  + (f"（掲載: {', '.join(touched)}）" if touched else ""))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?")
    ap.add_argument("--sheet", action="store_true", help="記入シートを作る")
    ap.add_argument("--apply", action="store_true", help="登録して掲載する")
    a = ap.parse_args()
    if a.sheet:
        p = make_sheet()
        print(f"記入シートを作りました → {p.relative_to(ROOT).as_posix()}")
        print("  埋めたら: python scripts/jisseki_intake.py <ファイル> --apply（または intake_watch.py --apply）")
        return 0
    if not a.xlsx:
        raise SystemExit(__doc__.strip())
    facts, voices, ng, warn = review(read(Path(a.xlsx)))
    for f in facts:
        print(f"  ○ 一次情報 [{f['id']}] {f['claim']}")
    for v in voices:
        print(f"  ○ お客様の声 {v['who']}（{v['industry']}）「{v['quote'][:30]}…」")
    for w in warn:
        print(f"  △ {w}")
    for x in ng:
        print(f"  × {x}")
    if ng:
        print(f"\n  不備{len(ng)}件。登録しません")
        return 1
    if not a.apply:
        print("\n  --apply で登録して掲載します")
        return 0
    touched = apply(facts, voices)
    print(f"\n  登録しました（一次情報{len(facts)}件・お客様の声{len(voices)}件）"
          + (f"\n  掲載: {', '.join(touched)}" if touched else ""))
    print("  次に python scripts/build.py を通し、コミットしてください")
    return 0


if __name__ == "__main__":
    sys.exit(main())
