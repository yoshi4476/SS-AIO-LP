# -*- coding: utf-8 -*-
"""多言語のメニュー・料金ページと QR コードを、日本語のシート1枚から作る（訪日客向け）。

**なぜ要るか**: 訪日客が店の前で最初に知りたいのは「何がいくらか」。日本語だけのメニューだと
そこで帰る。シートに日本語で書けば、指示した言語（sites/<id>.json の languages）の
メニューページ（Menu の構造化データつき）と、店頭・卓上に置く QR コードができる。

  python scripts/menu_page.py --sheet                          # intake/メニュー記入シート.xlsx を作る
  （埋めて intake/ に置く → intake_watch.py --apply が拾う。直接なら --add <xlsx>）
  python scripts/menu_page.py --build --site <id>              # 訳す（数字の検算つき）→ ページと QR
  python scripts/menu_page.py --build --site <id> --push       # 配信先（別リポジトリ）へ置く

守ること: 価格・数字は訳の前後で1つも変えない（変わった訳は捨てる）。書いていない食材・効能を足さない。
置き場: data/menus/<site>.json（元）/ data/menus/<site>.<lang>.json（訳）/ automation/menu/<site>/（ページと QR）
"""
import argparse
import hashlib
import html as _h
import json
import re
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
DATA = ROOT / "data" / "menus"
OUT = ROOT / "automation" / "menu"
SHEET = ROOT / "intake" / "メニュー記入シート.xlsx"
LABEL = {"ja": ("メニュー", "税込", "円"), "en": ("Menu", "tax incl.", "JPY"),
         "zh": ("菜单", "含税", "日元"), "ko": ("메뉴", "세금 포함", "엔")}
LANG_ATTR = {"ja": "ja", "en": "en", "zh": "zh-Hans", "ko": "ko"}
# 記入例。消し忘れると実在しない品目が訳されて公開されるため、読むときにも飛ばす
EXAMPLE = [("麺類", "醤油ラーメン", 950, "鶏と魚介のスープ", "小麦・卵"), ("ご飯もの", "チャーシュー丼", 480, "", "")]


# ── シート ───────────────────────────────────────────
def make_sheet(path=SHEET):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "メニュー"
    ws["A1"] = "メニュー記入シート（訪日客向けの多言語メニューとQRコードを作ります）"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"], ws["B2"] = "サイトID", ""
    ws["A3"], ws["B3"] = "店名", ""
    ws["A4"] = "価格は税込の円で数字だけ（例 1200）。説明は分かる範囲で。書いていないことは訳でも足しません。"
    ws["A5"] = "7〜8行目（灰色）は記入例です。上書きするか削除してください（残っていても読み込みません）。"
    heads = ["区分（例: 麺類）", "品名", "価格（税込・円）", "説明", "注記（アレルゲン・辛さ等）"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=6, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1D3461")
    gray = PatternFill("solid", start_color="D9D9D9")
    for r, row in enumerate(EXAMPLE, 7):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=f"例: {v}" if c == 1 else v)
            cell.fill = gray
    for col, w in zip("ABCDE", (18, 26, 16, 40, 28)):
        ws.column_dimensions[col].width = w
    path.parent.mkdir(exist_ok=True)
    wb.save(path)
    return path


def read_sheet(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    site, shop = str(ws["B2"].value or "").strip(), str(ws["B3"].value or "").strip()
    items = []
    for r in range(7, ws.max_row + 1):
        cat, name, price, desc, note = (ws.cell(row=r, column=c).value for c in range(1, 6))
        if not name:
            continue
        p = re.sub(r"[^\d]", "", str(price or ""))
        is_example = str(cat or "").strip().startswith(("例:", "例：")) or any(
            str(name).strip() == e[1] and p == str(e[2]) for e in EXAMPLE)
        if is_example:
            continue
        items.append({"cat": str(cat or "").strip(), "name": str(name).strip(), "price": int(p) if p else None,
                      "desc": str(desc or "").strip(), "note": str(note or "").strip()})
    return site, shop, items


def one(src, write):
    """intake_watch から呼ばれる。(登録できたか, 説明)"""
    try:
        site, shop, items = read_sheet(src)
    except Exception as e:
        return False, f"読めません（{str(e)[:60]}）"
    ng = []
    import sites as S
    if not site or site not in S.load_all():
        ng.append(f"サイトID「{site}」が登録されていません")
    if not items:
        ng.append("品目が1つもありません")
    if [i for i in items if i["price"] is None]:
        ng.append("価格が数字でない品目があります")
    if ng:
        return False, " / ".join(ng)
    if write:
        DATA.mkdir(parents=True, exist_ok=True)
        (DATA / f"{site}.json").write_text(json.dumps({"site": site, "shop": shop, "items": items,
                                                       "updated": date.today().isoformat()}, ensure_ascii=False, indent=1), encoding="utf-8")
    return True, f"メニュー {len(items)}品目（{site}）"


# ── 訳 ──────────────────────────────────────────────
PROMPT = """Translate this Japanese restaurant/shop menu into {lang}. Output ONLY a JSON object {{"shop": "...", "items": [...]}}
with exactly the same number of items in the same order, each with keys cat, name, price, desc, note.
Rules: keep "price" exactly as given (the same integer). Do not add ingredients, effects, or claims that are not in the source.
Keep proper names (dish names that are Japanese words) in romaji or the Japanese word plus a short gloss, e.g. "Shoyu Ramen (soy-sauce ramen)".
Empty strings stay empty. Source:
{src}"""


def translate(menu, lang):
    import auto_rewrite as AR
    import i18n
    src = {"shop": menu["shop"], "items": menu["items"]}
    name = {"en": "English", "zh": "Simplified Chinese", "ko": "Korean"}[lang]
    r = AR.sh([AR.claude_bin(), "-p", "--max-turns", "1", *AR.model_args(), "--allowedTools", ""],
              timeout=600, stdin_text=PROMPT.format(lang=name, src=json.dumps(src, ensure_ascii=False)))
    m = re.search(r"\{.*\}", r.stdout or "", re.S)
    if not m:
        return None, "JSONが返らない"
    try:
        d = json.loads(m.group(0))
    except ValueError:
        return None, "JSONが壊れている"
    if len(d.get("items") or []) != len(src["items"]):
        return None, "品目の数が変わった"
    if [x.get("price") for x in d["items"]] != [x["price"] for x in src["items"]]:
        return None, "価格が変わった"
    ok, why = i18n._digits_ok(src, d)
    if not ok:
        ok2, _ = i18n._digits_ok(src, i18n._normalize_numwords(d, lang))
        if not ok2:
            return None, why
    return d, ""


# ── ページと QR ───────────────────────────────────────
def page(menu, lang, cfg, others):
    lab, tax, yen = LABEL[lang]
    by = {}
    for it in menu["items"]:
        by.setdefault(it.get("cat") or "", []).append(it)
    secs = []
    for cat, items in by.items():
        rows = "".join(
            f'<li><div class="n">{_h.escape(i["name"])}</div><div class="p">{i["price"]:,} {yen}</div>'
            + (f'<div class="d">{_h.escape(i["desc"])}</div>' if i.get("desc") else "")
            + (f'<div class="t">{_h.escape(i["note"])}</div>' if i.get("note") else "") + "</li>" for i in items)
        secs.append((f"<h2>{_h.escape(cat)}</h2>" if cat else "") + f"<ul>{rows}</ul>")
    base = f"https://{cfg['domain']}/menu/"
    alts = "\n".join(f'<link rel="alternate" hreflang="{LANG_ATTR[o]}" href="{base}{o}/">' for o in others)
    switch = " / ".join(f'<a href="{base}{o}/">{LABEL[o][0]}</a>' for o in others if o != lang)
    ld = {"@context": "https://schema.org", "@type": "Menu", "name": f"{menu['shop']} {lab}", "inLanguage": LANG_ATTR[lang],
          "hasMenuSection": [{"@type": "MenuSection", "name": cat or lab, "hasMenuItem": [
              {"@type": "MenuItem", "name": i["name"], "description": i.get("desc", ""),
               "offers": {"@type": "Offer", "price": i["price"], "priceCurrency": "JPY"}} for i in items]}
              for cat, items in by.items()]}
    return f"""<!DOCTYPE html>
<html lang="{LANG_ATTR[lang]}"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_h.escape(menu['shop'])} {lab}</title><link rel="canonical" href="{base}{lang}/">
{alts}
<link rel="alternate" hreflang="x-default" href="{base}ja/">
<style>body{{margin:0;background:#fdfcf9;color:#212b3d;font-family:-apple-system,"Hiragino Sans","Noto Sans CJK JP",sans-serif}}
main{{max-width:640px;margin:0 auto;padding:28px 18px 60px}}h1{{font-size:1.5rem;color:#1d3461}}h2{{font-size:1.1rem;border-bottom:2px solid #1d3461;padding-bottom:4px;margin-top:1.8em}}
ul{{list-style:none;padding:0;margin:0}}li{{display:grid;grid-template-columns:1fr auto;gap:2px 12px;padding:12px 0;border-bottom:1px solid #e7e2d4}}
.n{{font-weight:700}}.p{{font-weight:700;white-space:nowrap}}.d,.t{{grid-column:1/-1;font-size:.9rem;color:#5b6472}}.t{{font-size:.82rem}}
.sw{{font-size:.9rem}}.tax{{font-size:.85rem;color:#5b6472}}</style>
<script type="application/ld+json">{json.dumps(ld, ensure_ascii=False)}</script></head>
<body><main><p class="sw">{switch}</p><h1>{_h.escape(menu['shop'])} {lab}</h1><p class="tax">{tax}</p>
{''.join(secs)}</main></body></html>
"""


def qr(url, path):
    import qrcode
    img = qrcode.make(url, box_size=10, border=2)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def build(site, push=False, write=True):
    import sites as S
    cfg = S.load(site)
    src_p = DATA / f"{site}.json"
    if not src_p.is_file():
        print(f"   {site}: メニューがありません（--sheet → 記入 → intake/ に置く）")
        return 0
    menu = json.loads(src_p.read_text(encoding="utf-8"))
    h = hashlib.md5(json.dumps(menu["items"], ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    langs = ["ja"] + [l for l in (cfg.get("languages") or []) if l in ("en", "zh", "ko")]
    menus = {"ja": menu}
    for lg in langs[1:]:
        tp = DATA / f"{site}.{lg}.json"
        old = json.loads(tp.read_text(encoding="utf-8")) if tp.is_file() else {}
        if old.get("src_hash") == h:
            menus[lg] = old
            continue
        if not write:
            print(f"   {lg}: 訳が必要（--build で訳します）")
            continue
        d, why = translate(menu, lg)
        print(f"   {'○' if d else '×'} {lg} {why}")
        if d:
            d["src_hash"] = h
            tp.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
            menus[lg] = d
    out = OUT / site
    have = list(menus)
    for lg, m in menus.items():
        p = out / lg / "index.html"
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(page(m, lg, cfg, have), encoding="utf-8")
        qr(f"https://{cfg['domain']}/menu/{lg}/", out / f"qr-{lg}.png")
    print(f"   ページと QR: {out}（{'・'.join(have)}）")
    if push and cfg.get("type") in ("external-html", "external-md"):
        import publish
        tok = publish._push_token()
        dest = publish.ensure_clone(cfg, tok)
        shutil.copytree(out, dest / "menu", dirs_exist_ok=True, ignore=shutil.ignore_patterns("qr-*.png"))
        md = dest / "tools" / "make_dist.py"
        if md.is_file():
            t = md.read_text(encoding="utf-8")
            mm = re.search(r"PUBLIC_DIRS\s*=\s*\[([^\]]*)\]", t)
            if mm and '"menu"' not in mm.group(1):
                md.write_text(t[:mm.end(1)] + ', "menu"' + t[mm.end(1):], encoding="utf-8", newline="\n")
        subprocess.run(["git", "add", "-A"], cwd=dest, check=True)
        subprocess.run(["git", "-c", "user.name=AIO Pipeline Bot", "-c", "user.email=noreply@7senses.co.jp",
                        "commit", "-q", "-m", "多言語メニュー（/menu/）を更新"], cwd=dest)
        r = subprocess.run(["git", "push", f"https://x-access-token@github.com/{cfg['repo']}.git", f"HEAD:{cfg['branch']}"],
                           cwd=dest, env=publish.git_auth(tok), capture_output=True, text=True)
        print(f"   配信: {'OK' if r.returncode == 0 else 'NG'}")
    elif push:
        print(f"   {cfg.get('type')} への自動配置は未対応。{out} のファイルを先方のサイトに置いてください")
    return len(have)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", action="store_true")
    ap.add_argument("--add")
    ap.add_argument("--build", action="store_true")
    ap.add_argument("--site", default="")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--push", action="store_true")
    a = ap.parse_args()
    if a.sheet:
        print(f"作りました: {make_sheet()}")
        return 0
    if a.add:
        ok, why = one(Path(a.add), True)
        print(("登録: " if ok else "不備: ") + why)
        return 0
    if a.build:
        import sites as S
        for sid in ([a.site] if a.site else [p.stem for p in DATA.glob("*.json") if "." not in p.stem] if a.all else []):
            if sid in S.load_all():
                print(f"■ {sid}")
                build(sid, a.push)
        print("MENU_OK=yes")
        return 0
    ap.print_help()
    return 0


if __name__ == "__main__":
    sys.exit(main())
